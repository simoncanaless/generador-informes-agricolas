"""Operaciones de bajo nivel sobre openpyxl.

Este modulo no conoce el dominio ni el mapping: solo sabe de celdas, estilos,
formulas y filas. Existe porque ``ws.insert_rows()`` de openpyxl mueve los
valores pero NO arregla nada de lo que los acompana:

  - los rangos combinados se quedan en su sitio (la fila se mueve, el merge no),
  - las formulas no se reescriben (``=SUM(D7:D10)`` se queda apuntando a D7:D10
    aunque el bloque haya crecido, dando totales silenciosamente erroneos),
  - los altos de fila no se desplazan,
  - las imagenes flotantes no se desplazan,
  - las filas nuevas nacen sin estilo, sin bordes y sin las formulas por fila.

``insertar_filas()`` es la version que si hace todo eso.
"""

from __future__ import annotations

import datetime as dt
import re
from copy import copy
from decimal import Decimal
from enum import Enum
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import (column_index_from_string, get_column_letter,
                            range_boundaries)
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.worksheet import Worksheet

# --------------------------------------------------------------- formulas

# Trozos entrecomillados de una formula: ="Total: "&A1 no debe tocarse.
_RE_CADENA = re.compile(r'"[^"]*"')

# Referencia A1 con hoja opcional. Si la hoja viene explicita ('Flujo'!F110)
# la referencia apunta a OTRA pestana y no debe desplazarse.
_RE_REFERENCIA = re.compile(
    r"(?P<hoja>(?:'[^']*'|[A-Za-z_][\w.]*)!)?"
    r"(?P<c1>\$?[A-Za-z]{1,3}\$?\d+)"
    r"(?::(?P<c2>\$?[A-Za-z]{1,3}\$?\d+))?"
)

_RE_CELDA = re.compile(r"^(?P<dc>\$?)(?P<col>[A-Za-z]{1,3})(?P<df>\$?)(?P<fila>\d+)$")


def _desplazar_celda(ref: str, desde: int, cantidad: int) -> str:
    """Desplaza la fila de una referencia A1 si esta en la zona afectada.

    Se desplaza tambien si la fila es absoluta (``$D$15``): en Excel el ``$``
    protege frente a la copia, no frente a la insercion de filas.
    """
    m = _RE_CELDA.match(ref)
    if m is None:
        return ref
    fila = int(m.group("fila"))
    if fila < desde:
        return ref
    return "{}{}{}{}".format(
        m.group("dc"), m.group("col"), m.group("df"), fila + cantidad
    )


def _fila_de(ref: str) -> Optional[int]:
    m = _RE_CELDA.match(ref)
    return int(m.group("fila")) if m else None


def desplazar_referencias(formula: str, desde: int, cantidad: int,
                          hoja: Optional[str] = None,
                          extender_fin: Optional[int] = None) -> str:
    """Reescribe una formula tras insertar ``cantidad`` filas en ``desde``.

    Aplica la misma regla que Excel: una referencia a la fila ``r`` pasa a
    ``r + cantidad`` si ``r >= desde``. Como consecuencia, un rango que termina
    justo en la zona insertada se estira (``SUM(D7:D10)`` con desde=10 pasa a
    ``SUM(D7:D41)``), que es exactamente lo que se busca en las filas de total.

    Sin ``hoja``: toca solo las referencias de la propia hoja (las que no
    nombran pestana). Con ``hoja``: toca solo las que nombran ESA pestana
    -es la formula de otra hoja que apunta a la que crecio-.

    ``extender_fin``: un RANGO que termina justo en esa fila tambien crece.
    Es para la tabla de una sola fila (las instalaciones de Rabobank): las
    filas nuevas se insertan DEBAJO de ella, y su total (=SUM(E22:E22))
    tiene que estirarse hasta cubrirlas -insertando en la misma fila, como
    con las tablas de varias, Excel corre el rango entero y el total
    sumaria solo la ultima-.
    """
    if not formula.startswith("="):
        return formula

    # Protege los literales de texto sustituyendolos por marcadores.
    literales: List[str] = []

    def _guardar(m: "re.Match[str]") -> str:
        literales.append(m.group(0))
        return "\x00{}\x00".format(len(literales) - 1)

    protegida = _RE_CADENA.sub(_guardar, formula)

    def _sustituir(m: "re.Match[str]") -> str:
        prefijo = m.group("hoja") or ""
        if hoja is None:
            if prefijo:
                return m.group(0)
        else:
            nombre = prefijo[:-1]
            if nombre.startswith("'") and nombre.endswith("'"):
                nombre = nombre[1:-1].replace("''", "'")
            if not prefijo or nombre != hoja:
                return m.group(0)
        nueva = prefijo + _desplazar_celda(m.group("c1"), desde, cantidad)
        if m.group("c2"):
            fin = m.group("c2")
            if extender_fin is not None and _fila_de(fin) == extender_fin:
                nueva += ":" + _desplazar_celda(fin, extender_fin, cantidad)
            else:
                nueva += ":" + _desplazar_celda(fin, desde, cantidad)
        return nueva

    resultado = _RE_REFERENCIA.sub(_sustituir, protegida)

    for i, literal in enumerate(literales):
        resultado = resultado.replace("\x00{}\x00".format(i), literal)
    return resultado


# ----------------------------------------------------------------- celdas

def es_formula(valor: Any) -> bool:
    return isinstance(valor, str) and valor.startswith("=")


def celda_ancla(ws: Worksheet, coordenada: str) -> str:
    """Devuelve la celda escribible que representa a ``coordenada``.

    Si cae dentro de un rango combinado, el valor real vive en la esquina
    superior izquierda; escribir en cualquier otra lanza AttributeError porque
    openpyxl las expone como ``MergedCell`` de solo lectura.
    """
    celda = ws[coordenada]
    if not isinstance(celda, MergedCell):
        return coordenada
    for rango in ws.merged_cells.ranges:
        if (rango.min_row <= celda.row <= rango.max_row
                and rango.min_col <= celda.column <= rango.max_col):
            return "{}{}".format(get_column_letter(rango.min_col), rango.min_row)
    return coordenada


def normalizar(valor: Any) -> Any:
    """Convierte un valor de dominio en algo que openpyxl sepa escribir."""
    if valor is None:
        return None
    if isinstance(valor, Enum):
        return valor.value
    if isinstance(valor, bool):
        return "Si" if valor else "No"
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (int, float, str, dt.date, dt.datetime, dt.time)):
        return valor
    return str(valor)


def escribir(
    ws: Worksheet,
    coordenada: str,
    valor: Any,
    formato: Optional[str] = None,
) -> Tuple[bool, Optional[str]]:
    """Escribe un valor respetando la plantilla.

    Devuelve ``(escrito, aviso)``. No escribe si el valor es ``None`` (deja lo
    que traiga la plantilla) y se niega a pisar una formula preexistente.
    """
    valor = normalizar(valor)
    if valor is None:
        return False, None
    if isinstance(valor, str) and not valor:
        return False, None

    destino = celda_ancla(ws, coordenada)
    celda = ws[destino]

    if es_formula(celda.value):
        return False, "{}!{} contiene una formula ({}); no se sobrescribe".format(
            ws.title, destino, celda.value
        )

    aviso = None
    if destino != coordenada:
        # La celda pertenece a un bloque combinado que EMPIEZA MAS ARRIBA. En
        # la hoja de Inversion eso es intencional: 'A5:A9' es el nombre de un
        # proyecto que abarca sus cinco lineas de item. Escribir la fila 6
        # aterriza en A5 y pisa el nombre del proyecto de la fila 5: el informe
        # terminaba mostrando el ULTIMO proyecto como si fuera el unico.
        #
        # Asi que el ancla la escribe quien llega primero. Si el que llega
        # despues dice lo mismo, no hay nada que hacer; si dice otra cosa, se
        # avisa en vez de pisar, porque significa que dos proyectos distintos
        # estan compartiendo un bloque que solo admite uno.
        if celda.value not in (None, ""):
            if celda.value == valor:
                return False, None
            return False, (
                "{}!{} cae en el bloque combinado {}, que ya dice {!r}: no se "
                "sobrescribio con {!r}".format(
                    ws.title, coordenada, destino, celda.value, valor)
            )
        aviso = "{}!{} esta combinada; el valor se escribio en {}".format(
            ws.title, coordenada, destino
        )

    celda.value = valor
    if isinstance(valor, (dt.date, dt.datetime)) and not isinstance(valor, dt.time):
        celda.number_format = formato or "DD-MM-YYYY"
        _asegurar_ancho(ws, celda.column, len(celda.number_format) + 2)
    elif formato:
        celda.number_format = formato
    return True, aviso


def reemplazar_formula(
    ws: Worksheet, coordenada: str, formula: str
) -> Tuple[bool, Optional[str]]:
    """Sustituye una formula preexistente de la plantilla por otra.

    Es la unica puerta para pisar una formula del banco, y esta separada de
    ``escribir`` a proposito: ahi la negativa a sobrescribir formulas es una
    salvaguarda que no conviene poder desactivar con un parametro suelto.

    Solo actua si la celda YA trae una formula. Si trae un literal o esta
    vacia, no hace nada y lo dice: significa que la plantilla cambio y la
    correccion declarada ya no corresponde a lo que hay, asi que hay que
    mirarla antes que aplicarla a ciegas.
    """
    destino = celda_ancla(ws, coordenada)
    celda = ws[destino]

    if not es_formula(celda.value):
        return False, (
            "{}!{} ya no trae una formula (trae {!r}); la correccion declarada "
            "no se aplico".format(ws.title, destino, celda.value)
        )
    if celda.value == formula:
        return False, None   # ya estaba corregida: no es un cambio

    celda.value = formula
    return True, None


def reemplazar_literal(ws: Worksheet, coordenada: str, valor) -> bool:
    """Sustituye un rotulo de la plantilla, o lo borra si ``valor`` es None.

    Distinta de ``escribir`` en dos cosas, y las dos hacen falta aqui: admite
    BORRAR -una temporada que el cliente no entrego no debe quedar rotulada
    como si existiera- y no se salta el valor vacio. Sigue sin tocar formulas.
    """
    destino = celda_ancla(ws, coordenada)
    celda = ws[destino]
    if es_formula(celda.value):
        return False
    celda.value = valor
    return True


def _asegurar_ancho(ws: Worksheet, columna: int, minimo: float) -> None:
    """Ensancha una columna si el dato no cabe.

    Excel muestra '#########' cuando una fecha o un numero no entra en el
    ancho disponible. Es tecnicamente correcto y visualmente inservible: el
    tasador abre el informe y cree que no se escribio nada. Solo se ensancha,
    nunca se estrecha, para no alterar el diseno de la plantilla mas de lo
    imprescindible.
    """
    dimension = ws.column_dimensions[get_column_letter(columna)]
    if dimension.width is None or dimension.width < minimo:
        dimension.width = minimo


# ------------------------------------------------------------ filas nuevas

def _instantanea_fila(
    ws: Worksheet, fila: int, ultima_columna: int
) -> List[Tuple[int, Any, Optional[str]]]:
    """Captura (columna, estilo, formula) de una fila para poder clonarla."""
    captura = []
    for col in range(1, ultima_columna + 1):
        celda = ws.cell(fila, col)
        formula = celda.value if es_formula(celda.value) else None
        captura.append((col, copy(celda._style), formula))
    return captura


def _merges_horizontales(ws: Worksheet, fila: int) -> List[Tuple[int, int]]:
    """Rangos combinados contenidos por completo en una sola fila."""
    return [
        (r.min_col, r.max_col)
        for r in ws.merged_cells.ranges
        if r.min_row == r.max_row == fila
    ]


def clonar_fila(
    ws: Worksheet,
    origen: int,
    destino: int,
    captura: Optional[List[Tuple[int, Any, Optional[str]]]] = None,
    merges: Optional[List[Tuple[int, int]]] = None,
    ultima_columna: Optional[int] = None,
) -> None:
    """Replica estilo, alto, combinaciones y formulas por fila de ``origen``.

    Los valores estaticos NO se copian: la fila nueva nace vacia de datos pero
    con la rejilla, los bordes y las formulas de fila (``=SUM(B16:C16)``,
    ``=I5/$I$15``) ya trasladadas a su nueva posicion.
    """
    ultima_columna = ultima_columna or ws.max_column
    if captura is None:
        captura = _instantanea_fila(ws, origen, ultima_columna)
    if merges is None:
        merges = _merges_horizontales(ws, origen)

    for col, estilo, formula in captura:
        celda = ws.cell(destino, col)
        if isinstance(celda, MergedCell):
            continue
        celda._style = copy(estilo)
        if formula is not None:
            celda.value = Translator(
                formula, origin="{}{}".format(get_column_letter(col), origen)
            ).translate_formula("{}{}".format(get_column_letter(col), destino))

    alto = ws.row_dimensions[origen].height
    if alto is not None:
        ws.row_dimensions[destino].height = alto

    for min_col, max_col in merges:
        ws.merge_cells(
            start_row=destino, start_column=min_col,
            end_row=destino, end_column=max_col,
        )


def _desplazar_rango(ref: str, desde: int, cantidad: int,
                     extender_fin: Optional[int] = None) -> str:
    """'A5:O225' con filas insertadas en `desde`: lo que empieza en o bajo
    `desde` baja; si el rango la cruza -o termina justo en `extender_fin`-,
    crece."""
    rango = CellRange(ref)
    if rango.min_row >= desde:
        rango.shift(row_shift=cantidad)
    elif rango.max_row >= desde or rango.max_row == extender_fin:
        rango.max_row += cantidad
    return rango.coord


def _desplazar_rangos_de_hoja(ws: Worksheet, desde: int, cantidad: int,
                              extender_fin: Optional[int] = None) -> None:
    if ws.auto_filter.ref:
        ws.auto_filter.ref = _desplazar_rango(ws.auto_filter.ref, desde, cantidad)

    areas = ws.print_area
    if areas:
        nuevas = []
        for area in (areas if isinstance(areas, list) else str(areas).split(",")):
            ref = str(area).split("!")[-1].replace("$", "")
            nuevas.append(_desplazar_rango(ref, desde, cantidad))
        ws.print_area = nuevas

    for validacion in ws.data_validations.dataValidation:
        validacion.sqref = MultiCellRange(" ".join(
            _desplazar_rango(r.coord, desde, cantidad, extender_fin)
            for r in validacion.sqref.ranges))


def insertar_filas(
    ws: Worksheet,
    desde: int,
    cantidad: int,
    fila_modelo: Optional[int],
    extender_fin: Optional[int] = None,
    *,
    formulas: bool = True,
) -> None:
    """Inserta filas arreglando todo lo que openpyxl deja a medias.

    ``desde``        fila 1-based donde se abre el hueco (las filas existentes
                     a partir de ahi bajan ``cantidad`` posiciones).
    ``fila_modelo``  fila cuyo estilo y formulas por fila se clonan en el hueco.
    ``extender_fin`` los rangos que terminan en esa fila tambien crecen: para
                     insertar DEBAJO de una tabla de una sola fila (ver
                     ``desplazar_referencias``).
    ``formulas``     False no reescribe formulas, listas desplegables ni area
                     de impresion: para quien las vuelve a escribir todas
                     despues desde el molde (services/disposicion.py). Con
                     ``fila_modelo=None`` las filas nuevas nacen vacias.
    """
    if cantidad <= 0:
        return

    ultima_columna = ws.max_column
    merges_previos = [str(r) for r in ws.merged_cells.ranges]
    dimensiones = {idx: copy(dim) for idx, dim in ws.row_dimensions.items()}

    ws.insert_rows(desde, cantidad)

    # 1) rangos combinados: openpyxl movio las celdas pero no la contabilidad.
    rangos: List[CellRange] = []
    for coordenadas in merges_previos:
        rango = CellRange(coordenadas)
        if rango.min_row >= desde:
            rango.shift(row_shift=cantidad)
        elif rango.max_row >= desde:
            rango.max_row += cantidad  # el merge cruza el hueco: se estira
        rangos.append(rango)
    ws.merged_cells.ranges = rangos

    # 2) altos de fila
    ws.row_dimensions.clear()
    for idx, dim in dimensiones.items():
        nuevo = idx + cantidad if idx >= desde else idx
        dim.index = nuevo
        ws.row_dimensions[nuevo] = dim

    # 3) imagenes flotantes (los anchors de openpyxl son 0-based)
    for imagen in getattr(ws, "_images", []):
        ancla = getattr(imagen, "anchor", None)
        marca = getattr(ancla, "_from", None)
        if marca is None or marca.row < desde - 1:
            continue
        marca.row += cantidad
        hasta = getattr(ancla, "to", None)
        if hasta is not None:
            hasta.row += cantidad

    if not formulas:
        return

    # 4) formulas de toda la hoja
    for fila in ws.iter_rows():
        for celda in fila:
            if es_formula(celda.value):
                celda.value = desplazar_referencias(
                    celda.value, desde, cantidad, extender_fin=extender_fin)

    # 4a) y las de las OTRAS hojas que apuntan a esta. Excel las corre; sin
    # esto, la portada de Rabobank (=CONSTR!S23, el total de instalaciones)
    # quedaba leyendo una fila de datos en cuanto habia dos instalaciones.
    libro = getattr(ws, "parent", None)
    for otra in (libro.worksheets if libro is not None else []):
        if otra is ws:
            continue
        for fila in otra.iter_rows():
            for celda in fila:
                if es_formula(celda.value) and "!" in celda.value:
                    celda.value = desplazar_referencias(
                        celda.value, desde, cantidad, hoja=ws.title,
                        extender_fin=extender_fin)

    # 4b) lo que tambien vive en rangos de filas: el autofiltro, el area de
    # impresion y las listas desplegables. Sin esto, en una hoja con el
    # bloque de estadistica al pie (REF. Canal, Rabobank) el area de
    # impresion termina antes que la hoja y el pie no se imprime.
    _desplazar_rangos_de_hoja(ws, desde, cantidad, extender_fin)

    # 5) estilo y formulas de las filas nuevas
    if fila_modelo is None:
        return
    modelo = fila_modelo + cantidad if fila_modelo >= desde else fila_modelo
    captura = _instantanea_fila(ws, modelo, ultima_columna)
    merges = _merges_horizontales(ws, modelo)
    for destino in range(desde, desde + cantidad):
        clonar_fila(ws, modelo, destino, captura, merges, ultima_columna)


# --------------------------------------------------- borrado (lo excluible)
#
# Espejo de insertar_filas: existe porque el tasador puede pedir que una
# tabla o una hoja entera NO aparezcan en el .xlsx que descarga
# (config/mapping_santander.json:unidades_excluibles). openpyxl tiene la
# misma trampa que con insert_rows -delete_rows mueve valores, no arregla
# merges/formulas/altos/imagenes-, asi que hace falta la misma cirugia, en
# reversa.


def _fila_tras_borrar(fila: int, desde: int, hasta: int, cantidad: int) -> int:
    """A donde va a parar una fila una vez borrado el rango [desde, hasta].

    Una fila DENTRO del rango borrado no tiene a donde ir: se ancla al
    comienzo del hueco, y quien use esto se encarga de descartarla si hace
    falta (un merge que quedo con min==max ahi significa que colapso).
    """
    if fila < desde:
        return fila
    if fila > hasta:
        return fila - cantidad
    return desde


def eliminar_filas(ws: Worksheet, desde: int, cantidad: int,
                   *, formulas: bool = True) -> None:
    """Borra filas arreglando lo que openpyxl deja a medias -ver insertar_filas.

    ``desde``    primera fila 1-based a borrar.
    ``cantidad`` cuantas filas se borran (``desde`` .. ``desde+cantidad-1``).

    No es responsabilidad de esta funcion decidir SI es seguro borrar -eso lo
    resuelve la cascada de dependencias antes de llamarla (services/
    exclusiones.py): si algo mas en la hoja todavia lee una fila de las que
    se van a borrar, el problema es de ahi, no de esta funcion mecanica.
    """
    if cantidad <= 0:
        return
    hasta = desde + cantidad - 1

    merges_previos = [str(r) for r in ws.merged_cells.ranges]
    dimensiones = {idx: copy(dim) for idx, dim in ws.row_dimensions.items()}

    # 1) imagenes flotantes: las ancladas DENTRO del hueco se pierden junto
    #    con su fila (nunca es el caso de las unidades excluibles de hoy,
    #    pero una tabla futura podria tener fotos). Las de mas abajo suben.
    conservar = []
    for imagen in list(getattr(ws, "_images", [])):
        ancla = getattr(imagen, "anchor", None)
        marca = getattr(ancla, "_from", None)
        if marca is None:
            conservar.append(imagen)
            continue
        fila_1based = marca.row + 1
        if desde <= fila_1based <= hasta:
            continue
        if fila_1based > hasta:
            marca.row -= cantidad
            hasta_ancla = getattr(ancla, "to", None)
            if hasta_ancla is not None:
                hasta_ancla.row -= cantidad
        conservar.append(imagen)
    ws._images = conservar

    ws.delete_rows(desde, cantidad)

    # 2) rangos combinados: los que caian ENTERO dentro del hueco desaparecen;
    #    los de mas abajo suben; los que cruzaban el borde se achican -la
    #    misma formula sirve para los tres casos.
    rangos: List[CellRange] = []
    for coordenadas in merges_previos:
        rango = CellRange(coordenadas)
        if rango.min_row >= desde and rango.max_row <= hasta:
            continue
        nuevo_min = _fila_tras_borrar(rango.min_row, desde, hasta, cantidad)
        nuevo_max = _fila_tras_borrar(rango.max_row, desde, hasta, cantidad)
        rango.min_row = nuevo_min
        rango.max_row = max(nuevo_min, nuevo_max)
        rangos.append(rango)
    ws.merged_cells.ranges = rangos

    # 3) altos de fila
    ws.row_dimensions.clear()
    for idx, dim in dimensiones.items():
        if desde <= idx <= hasta:
            continue
        nuevo = idx - cantidad if idx > hasta else idx
        dim.index = nuevo
        ws.row_dimensions[nuevo] = dim

    if not formulas:
        return

    # 4) formulas de toda la hoja: las que sobreviven y quedaban MAS ABAJO
    #    del hueco se desplazan hacia arriba. `desplazar_referencias` solo
    #    toca referencias de la MISMA hoja (las que cruzan de pestana se
    #    devuelven intactas), que es justo lo que hace falta: si algo de
    #    OTRA hoja todavia lee una fila borrada, eso lo debio impedir la
    #    cascada de dependencias antes de llegar aqui.
    for fila in ws.iter_rows():
        for celda in fila:
            if es_formula(celda.value):
                celda.value = desplazar_referencias(celda.value, hasta + 1, -cantidad)


def vaciar_columnas(
    ws: Worksheet, fila_ini: int, fila_fin: int, col_ini: str, col_fin: str
) -> None:
    """Borra el CONTENIDO de un rango de columnas, sin tocar filas.

    Para las unidades excluibles que comparten fila con otra tabla -en
    'Dotacion de recursos', Maquinaria (A-C) e Infraestructura (F-M) ocupan
    las MISMAS filas 42-45 en columnas distintas-. Borrar filas ahi para
    sacar una destruiria la otra, asi que "excluir" solo puede significar
    vaciar el dato: se deja el marco de la tabla -bordes, encabezado- y se
    borra lo que el tasador declaro, igual que una tabla que nunca se llena.
    """
    c1 = column_index_from_string(col_ini)
    c2 = column_index_from_string(col_fin)
    for fila in range(fila_ini, fila_fin + 1):
        for col in range(c1, c2 + 1):
            celda = ws.cell(fila, col)
            if not isinstance(celda, MergedCell):
                celda.value = None


def eliminar_hoja(wb: Workbook, nombre: str) -> None:
    """Saca una hoja entera del libro."""
    if nombre in wb.sheetnames:
        wb.remove(wb[nombre])


# ---------------------------------------------------- preparado de imagenes

#: Formatos que openpyxl sabe empaquetar dentro del .xlsx sin ayuda.
FORMATOS_SEGUROS = {"png", "jpeg", "gif", "bmp"}

#: Lado mayor al que se reduce una foto antes de incrustarla.
#: Se leen de services/imagenes.py para que el criterio viva en un solo sitio:
#: el peso del informe lo pidio el agrónomo y se ajusta ahi, no aqui.
from services.imagenes import CALIDAD_INFORME, LADO_INFORME

LADO_MAXIMO_INCRUSTADO = LADO_INFORME
CALIDAD_JPEG_INCRUSTADO = CALIDAD_INFORME


def preparar_imagen(ruta: str, lado_maximo: int = LADO_MAXIMO_INCRUSTADO):
    """Deja una fotografia lista para incrustar: formato conocido y tamano sano.

    Dos problemas reales que resuelve, ambos detectados con fotos de una visita
    de verdad:

    1. **Formato**: las camaras de los drones DJI guardan en MPO (Multi Picture
       Object), un JPEG con varias imagenes dentro. openpyxl no lo reconoce y
       revienta al escribir el archivo con ``KeyError: '.mpo'``.
    2. **Peso**: catorce fotos de 4 MB producen un Excel de 50 MB que ningun
       ejecutivo de banco va a poder abrir de un correo.

    Devuelve un buffer en memoria; si algo falla se devuelve la ruta original,
    porque perder la foto es peor que incrustarla pesada.
    """
    try:
        from io import BytesIO

        from PIL import Image

        with Image.open(ruta) as imagen:
            formato = (imagen.format or "").lower()
            necesita_ajuste = (
                formato not in FORMATOS_SEGUROS or max(imagen.size) > lado_maximo
            )
            if not necesita_ajuste:
                return ruta

            copia = imagen.convert("RGB")
            copia.thumbnail((lado_maximo, lado_maximo), Image.LANCZOS)
            buffer = BytesIO()
            copia.save(buffer, format="JPEG",
                       quality=CALIDAD_JPEG_INCRUSTADO, optimize=True)
            buffer.seek(0)
            return buffer
    except Exception:  # noqa: BLE001 - antes pesada que perdida
        return ruta


# ------------------------------------------------------- medidas para fotos

#: Alto de fila por defecto en puntos cuando la hoja no declara uno.
ALTO_FILA_POR_DEFECTO = 15.0
#: Ancho de columna por defecto en unidades de caracter.
ANCHO_COLUMNA_POR_DEFECTO = 8.43
#: Ancho maximo de digito (MDW) en pixeles, para la formula de Excel de
#: abajo. Con Calibri 11 -la fuente que asume la aproximacion "7w+5" que
#: esto reemplaza- vale 7. Medido con Excel real (COM, no aproximado) contra
#: la plantilla Bsan el 2026-09-22 -ver `_ancho_columna_px`- **tambien da 7**
#: para Aptos Narrow 11, que es la fuente 'Normal' real de ese archivo: el
#: MDW no cambio, lo que estaba mal era el "+5".
_MDW = 7


def _ancho_columna_px(unidades: float) -> int:
    """Ancho de UNA columna en pixeles, con la formula real de Excel.

    La aproximacion vieja (``w*7+5``) asume Calibri 11 y ya se sabia
    imprecisa para anchos grandes; lo que no se sabia es que para anchos
    ANGOSTOS -como el 2,66 de la rejilla de fotos- el error es enorme: 23,6px
    calculados contra 19px reales, un 24% de mas. Eso hacia que
    ``_encajar`` (services/excel_engine.py) pensara que el hueco de una foto
    era mas grande de lo que es, y escalara la imagen de mas: la foto
    quedaba **desbordando el hueco declarado**, invadiendo la columna
    siguiente sin ningun aviso (medido con `Shape.Width` real via COM: una
    foto declarada para caber en B4:P13 salia 9pt mas ancha que el propio
    hueco).

    La formula real (documentada en el estandar OOXML, verificada aqui
    contra 7 anchos distintos -1; 2,66; 5; 8,43; 10; 15; 20- leyendo
    ``Range.Width`` con Excel real sobre la plantilla Bsan: las 7 dieron
    exacto, ninguna aproximada)::

        px = TRUNC(((256*w + TRUNC(128/MDW)) / 256) * MDW)
    """
    return int(((256 * unidades + int(128 / _MDW)) / 256) * _MDW)


def dimensiones_rango_px(ws: Worksheet, rango: str) -> Tuple[int, int]:
    """Mide un rango en pixeles, para encajar una imagen sin deformarla.

    El ancho pasa por `_ancho_columna_px`, la formula real de Excel -ver ahi
    por que la aproximacion vieja no servia. El alto tiene la MISMA trampa,
    aunque no involucra ninguna fuente: Excel no guarda un alto de fila
    fraccionario tal cual, lo cuantiza al pixel -medido con Excel real
    (COM) sobre 6 altos declarados el 2026-09-22: 14pt se guarda como
    13,5pt, 20pt como 19,5pt, siempre para ABAJO, nunca para arriba-. Sumar
    los puntos SIN truncar cada fila antes (``14*10 = 140pt``) da un numero
    que Excel jamas muestra (el real es 135pt); hay que truncar fila por
    fila, no al final.
    """
    min_col, min_fila, max_col, max_fila = range_boundaries(rango)

    ancho = 0.0
    for col in range(min_col, max_col + 1):
        dim = ws.column_dimensions.get(get_column_letter(col))
        unidades = dim.width if dim is not None and dim.width else ANCHO_COLUMNA_POR_DEFECTO
        ancho += _ancho_columna_px(unidades)

    alto = 0
    for fila in range(min_fila, max_fila + 1):
        dim = ws.row_dimensions.get(fila)
        puntos = dim.height if dim is not None and dim.height else ALTO_FILA_POR_DEFECTO
        alto += int(puntos * 96.0 / 72.0)

    return int(ancho), alto
