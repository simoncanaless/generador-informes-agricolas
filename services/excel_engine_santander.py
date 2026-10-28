"""Motor del formato final de Santander (2026-10-13).

Escribe un ``InformeFitogreen`` sobre ``templates/Formato Fitogreen Santander
2026.xlsx`` guiado por ``config/mapping_santander_2026.json``. Es paralelo a
``services/excel_engine.py:MotorExcel`` -que llena el formato anterior, el
"Bsan"- por la misma razon que el de Rabobank: el formato nuevo no se parece
al viejo (un bloque de flujo POR PLANTACION, una hoja de Supuestos que mueve
todo, cuatro anexos), y forzar un solo motor a servir a los dos los habria
enredado a ambos.

Lo nuevo es que el molde trae UNA vez cada parte que se repite y este motor
las repite segun el informe: eso lo resuelve ``services/disposicion.py``, que
no sabe nada de Santander. Aqui se arma, para cada parte, la lista de sus
copias y los datos de cada una, y se escriben.

El orden importa:

1. La disposicion: cuantas copias de cada parte y a que fila va cada una.
2. Armar las filas (estilo, alto, rotulos, combinadas) hoja por hoja.
3. TODAS las formulas del molde, reescritas para la disposicion final.
4. Los datos.
5. Listas desplegables, formatos condicionales y area de impresion, estiradas
   a las copias.
6. Guardar, y despues reponer lo que openpyxl pierde: los dibujos (con fotos y
   textos) y el marcado de formula dinamica de UNIQUE/FILTER.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
import unicodedata
import warnings
import zipfile
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple, Union

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from schemas.informe import CategoriaFoto, InformeFitogreen
from services import dibujos_xlsx, excel_ops
from services.disposicion import Disposicion, FaltaContexto, Parte, ParteDeColumnas
from services.excel_engine import ResultadoGeneracion

__all__ = ["MotorSantander", "temporada_real_por_defecto", "temporadas_siguientes"]


# ------------------------------------------------------------- temporadas

def temporada_real_por_defecto(fecha: Optional[date] = None) -> str:
    """La ultima temporada CERRADA a la fecha del informe. La temporada
    fruticola va de julio a junio: en octubre de 2026 corre la 2026/27, asi
    que la ultima cerrada es la 2025/26."""
    fecha = fecha or date.today()
    inicio = fecha.year - 1 if fecha.month >= 7 else fecha.year - 2
    return "{}/{:02d}".format(inicio, (inicio + 1) % 100)


def temporadas_siguientes(real: str, cuantas: int = 8) -> List[str]:
    """Las ``cuantas`` temporadas que siguen a ``real`` ('2025/26' ->
    ['2026/27', ...]). Son las columnas F..M del flujo."""
    anio = int(str(real)[:4])
    return ["{}/{:02d}".format(a, (a + 1) % 100) for a in range(anio + 1, anio + 1 + cuantas)]


# ------------------------------------------------------------- utilidades

def _normal(texto: Any) -> str:
    texto = unicodedata.normalize("NFD", str(texto or ""))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return " ".join(texto.lower().split())


def _valor(valor: Any) -> Any:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if hasattr(valor, "value"):
        return valor.value
    return valor


def _resolver(raiz: Any, ruta: str) -> Any:
    actual = raiz
    for tramo in ruta.split("."):
        if actual is None:
            return None
        actual = actual.get(tramo) if isinstance(actual, dict) else getattr(actual, tramo, None)
    return actual


_FUNCIONES_DINAMICAS = ("_xlfn.UNIQUE", "_xlfn._xlws.FILTER", "_xlfn._xlws.SORT",
                        "_xlfn.SEQUENCE", "_xlfn.SORTBY")


# ------------------------------------------------------------- los datos

class _Datos:
    """Lo que el informe dice, ordenado como lo pide el formato: las
    plantaciones agrupadas por especie (en el orden en que aparece cada una),
    y para cada parte que se repite, la lista de sus copias."""

    def __init__(self, informe: InformeFitogreen, reglas=None,
                 fecha: Optional[date] = None):
        self.informe = informe
        self.avisos: List[str] = []
        crudas = list(informe.estructura_productiva.plantaciones or [])

        self.especies: List[str] = []
        por_especie: Dict[str, list] = {}
        nombres: Dict[str, str] = {}
        for p in crudas:
            clave = _normal(p.especie)
            if clave not in por_especie:
                self.especies.append(clave)
                nombres[clave] = p.especie
                por_especie[clave] = []
            por_especie[clave].append(p)
        self.nombre_especie = nombres
        self.plantaciones = [p for clave in self.especies for p in por_especie[clave]]
        self.indice_especie = {clave: i for i, clave in enumerate(self.especies)}
        self.especie_de = [self.indice_especie[_normal(p.especie)] for p in self.plantaciones]

        dot = informe.dotacion_recursos
        self.lotes = list(dot.lotes)
        self.unidades = list(dot.unidades_prediales)
        self.fuentes = list(dot.derechos_agua)

        # Riego: un bloque por unidad predial con plantaciones -las del
        # cuadro de suelos primero, en su orden; despues las que solo
        # aparecen en el catastro-, y en cada uno sus especies, en el orden
        # en que las lista el UNIQUE de la plantilla.
        predios: List[str] = []
        for u in self.unidades:
            if any(_normal(p.predio) == _normal(u.nombre) for p in self.plantaciones):
                predios.append(u.nombre)
        for p in self.plantaciones:
            if p.predio and _normal(p.predio) not in {_normal(x) for x in predios}:
                predios.append(p.predio)
        self.riego = predios
        self.riego_especies: List[List[str]] = []
        for predio in predios:
            vistas: List[str] = []
            for p in self.plantaciones:
                if _normal(p.predio) == _normal(predio) and _normal(p.especie) not in vistas:
                    vistas.append(_normal(p.especie))
            self.riego_especies.append(vistas)
        sin_predio = [p for p in self.plantaciones if not p.predio]
        if sin_predio:
            self.avisos.append(
                "{} plantación(es) sin predio: no entran al requerimiento de riego"
                .format(len(sin_predio)))

        self.reglas = reglas
        self.fecha = fecha or informe.antecedentes.fecha_emision_informe or date.today()

    # -- copias de cada parte (para la disposicion) --------------------------

    def instancias(self) -> Dict[str, Any]:
        """Una lista vacia deja la parte sin filas; las tablas que se llenan a
        mano conservan su fila en blanco aunque no haya datos."""

        def al_menos_una(n, dimension):
            return lambda ctx: [{dimension: i} for i in range(max(1, n))]

        def especies(ctx):
            return [{"especie": i} for i in range(max(1, len(self.especies)))]

        def plantaciones(ctx):
            if not self.plantaciones:
                return [{"plantacion": 0, "especie": 0}]
            return [{"plantacion": i, "especie": self.especie_de[i]}
                    for i in range(len(self.plantaciones))]

        def de_la_especie(ctx):
            s = ctx.get("especie", 0)
            if not self.plantaciones:
                return [{"plantacion": 0, "especie": s}]
            return [{"plantacion": i, "especie": s}
                    for i in range(len(self.plantaciones)) if self.especie_de[i] == s]

        def riego(ctx):
            return [{"riego": i} for i in range(len(self.riego))]

        def riego_especies(ctx):
            n = len(self.riego_especies[ctx["riego"]]) if self.riego else 0
            return [{"riego_especie": j} for j in range(max(1, n))]

        return {
            "especies_portada": especies, "resumen_superficie": especies,
            "sistemas_produccion": especies, "comercializacion": especies,
            "flujo_superficie": especies, "flujo_produccion": especies,
            "flujo_rendimiento": especies, "flujo_precio": especies,
            "flujo_costo_ha": especies, "flujo_costo_kg": especies,
            "flujo_margen": especies, "flujo_ingresos": especies, "flujo_costos": especies,
            "secciones_supuestos": especies, "especies_anexo1": especies,
            "especies_conciliacion": especies,
            "lotes": al_menos_una(len(self.lotes), "lote"),
            "unidades": al_menos_una(len(self.unidades), "unidad"),
            "fuentes": al_menos_una(len(self.fuentes), "fuente"),
            "riego": riego, "riego_especies": riego_especies,
            "plantaciones": plantaciones, "bloques": plantaciones,
            "filas_supuestos": de_la_especie,
        }

    # -- los datos de cada copia ------------------------------------------

    def registro(self, dimension: str, ctx: Dict[str, int]) -> Dict[str, Any]:
        i = ctx.get(dimension)
        if dimension == "plantacion":
            return self._plantacion(i)
        if dimension == "especie":
            return self._especie(i)
        if dimension == "lote":
            return self._modelo(self.lotes, i, {"conservador": "conservador"})
        if dimension == "unidad":
            return self._modelo(self.unidades, i)
        if dimension == "fuente":
            datos = self._modelo(self.fuentes, i)
            if datos:
                datos["unidad_predial"] = datos.get("predio")
            return datos
        if dimension == "riego":
            if i is None or i >= len(self.riego):
                return {}
            nombre = self.riego[i]
            unidad = next((u for u in self.unidades if _normal(u.nombre) == _normal(nombre)), None)
            return {"unidad_predial": nombre,
                    "capacidad_tranques_m3": unidad.capacidad_tranques_m3 if unidad else None}
        if dimension == "riego_especie":
            u = ctx.get("riego")
            if u is None or u >= len(self.riego) or i >= len(self.riego_especies[u]):
                return {}
            especie = self.riego_especies[u][i]
            sistemas = [p.riego for p in self.plantaciones
                        if _normal(p.predio) == _normal(self.riego[u])
                        and _normal(p.especie) == especie and p.riego]
            sistema = Counter(sistemas).most_common(1)[0][0] if sistemas else None
            return {"sistema_riego": _sistema_de_riego(sistema)}
        return {}

    @staticmethod
    def _modelo(lista, i, _alias=None) -> Dict[str, Any]:
        if i is None or i >= len(lista):
            return {}
        return lista[i].model_dump()

    def _plantacion(self, i: Optional[int]) -> Dict[str, Any]:
        if i is None or i < 0 or i >= len(self.plantaciones):
            return {}
        p = self.plantaciones[i]
        datos = {
            "especie": p.especie, "variedad": p.variedad, "anio": _anio(p.anio_plantacion),
            "royalty": p.royalty if p.royalty is not None else self._royalty(p),
            "superficie_ha": p.superficie_ha,
            "conduccion": p.sistema_conduccion, "predio": p.predio, "riego": p.riego,
            "estatus": p.estatus_productivo or self._estatus(p), "densidad": p.densidad_plantacion,
            "dominio": p.dominio,
        }
        s = p.supuestos
        if s is not None:
            datos.update(s.model_dump(exclude={"rendimientos_kg_ha", "curva"}))
            for k, rend in enumerate(s.rendimientos_kg_ha[:8], start=1):
                datos["rendimiento_{}".format(k)] = rend
        return datos

    def _royalty(self, p) -> Optional[bool]:
        """Si/No segun la lista de variedades con royalty de la visita; sin
        lista, queda en blanco para el tasador."""
        texto = self.informe.estructura_productiva.variedades_con_royalty
        if not texto:
            return None
        variedades = {_normal(v) for v in re.split(r"[,;\n]", texto) if v.strip()}
        return _normal(p.variedad) in variedades

    def _estatus(self, p) -> Optional[str]:
        from services.catastro_import import anio_inicial
        inicio = anio_inicial(p.anio_plantacion)
        if inicio is None or self.reglas is None:
            return None
        edad = self.fecha.year - inicio
        formacion, crecientes = self.reglas.umbrales(p.especie)
        if edad <= formacion:
            return "Formación"
        return "Creciente" if edad <= crecientes else "Plena"

    def _especie(self, i: Optional[int]) -> Dict[str, Any]:
        if i is None or i >= len(self.especies):
            return {}
        clave = self.especies[i]
        nombre = self.nombre_especie[clave]
        inf = self.informe
        datos: Dict[str, Any] = {"nombre": nombre, "nombre_mayusculas": nombre.upper()}

        sistema = next((x for x in inf.estructura_productiva.sistemas_produccion
                        if _normal(x.especie) == clave), None)
        if sistema is not None:
            datos.update({
                "sistema_produccion": sistema.sistema_produccion,
                "densidad": sistema.densidad_plantacion,
                "sistema_riego": sistema.sistema_riego,
                "techos": sistema.considera_techos,
                "heladas": sistema.sistema_contra_heladas,
                "nivel_tecnologico": sistema.nivel_tecnologico,
                "homogeneidad": sistema.nivel_homogeneidad,
                "estado_fitosanitario": sistema.estado_fitosanitario,
                "nivel_productivo": sistema.nivel_productivo_promedio,
            })
        venta = next((x for x in inf.estructura_productiva.sistemas_comercializacion
                      if _normal(x.especie) == clave), None)
        if venta is not None:
            datos.update({
                "destino": venta.mercado, "empresas": venta.empresas_comercializadoras,
                "sistema_venta": venta.comercializacion, "certificaciones": venta.certificaciones,
                "packing": venta.proceso, "proceso_exportacion": venta.gestion_exportacion,
                "periodo_cosecha": venta.periodo_cosecha,
                "periodo_liquidaciones": venta.periodo_liquidaciones,
            })
        fila = next((x for x in inf.anexo.resultado_real.filas if _normal(x.especie) == clave), None)
        if fila is not None:
            exp = fila.produccion_exportacion_kg or 0.0
            interno = fila.produccion_mercado_interno_kg or 0.0
            total = exp + interno
            datos.update({
                "produccion_total_kg": total or None,
                "pct_exportacion": round(exp / total, 4) if total else None,
                "ventas_exportacion_pesos": fila.ventas_exportacion_pesos,
                "ventas_mercado_interno_pesos": fila.ventas_mercado_interno_pesos,
                "costos_directos_pesos": fila.costos_directos_pesos,
            })
        sens = inf.supuestos.sensibilizacion.get(nombre) or next(
            (v for k, v in inf.supuestos.sensibilizacion.items() if _normal(k) == clave), None)
        datos.update({
            "sens_precio_exportacion": sens.precio_exportacion if sens else 0,
            "sens_precio_mercado_interno": sens.precio_mercado_interno if sens else 0,
            "sens_rendimiento": sens.rendimiento if sens else 0,
            "sens_costos": sens.costos if sens else 0,
        })
        nota = inf.supuestos.notas.get(nombre) or next(
            (v for k, v in inf.supuestos.notas.items() if _normal(k) == clave), None)
        datos["nota_supuestos"] = nota
        return datos

    def declarados_sin_catastro(self) -> List[str]:
        """Especies con resultados declarados que no estan en el catastro: el
        formato pone el Anexo 1 en el mismo orden que el resumen del flujo,
        asi que no tienen donde ir."""
        return [f.especie for f in self.informe.anexo.resultado_real.filas
                if f.especie and _normal(f.especie) not in self.indice_especie]

    # -- celdas sueltas ----------------------------------------------------

    def campo(self, ruta: str) -> Any:
        if ruta.startswith("@"):
            return self._calculado(ruta[1:])
        return _resolver(self.informe, ruta)

    def _calculado(self, nombre: str) -> Any:
        flujo = self.informe.flujo_proyectado
        cambio = flujo.tipo_cambio_pesos_usd or {}
        real = self.temporada_real()
        if nombre == "temporada_real":
            return real
        if nombre == "tc_real":
            return cambio.get("real") or cambio.get(real)
        if nombre == "tc_proyeccion":
            siguiente = temporadas_siguientes(real, 1)[0]
            return cambio.get("proyeccion") or cambio.get(siguiente) or cambio.get("preliminar")
        if nombre == "tc_anterior":
            anterior = "{}/{:02d}".format(int(real[:4]) - 1, int(real[:4]) % 100)
            return cambio.get("anterior") or cambio.get(anterior)
        return None

    def temporada_real(self) -> str:
        declarada = (self.informe.anexo.resultado_real.temporada or "").strip()
        return declarada or temporada_real_por_defecto(self.fecha)


def _anio(texto: Any) -> Any:
    if texto is None:
        return None
    try:
        return int(str(texto).strip())
    except ValueError:
        return texto


_SISTEMAS_RIEGO = ["Goteo", "Microaspersión", "Aspersión o Pivote", "Surco", "Californiano"]


def _sistema_de_riego(texto: Optional[str]) -> Optional[str]:
    """El nombre de la lista de la hoja Opciones -de ahi sale la eficiencia-,
    sin importar tildes ni mayusculas. Si no calza, el texto tal cual."""
    if not texto:
        return None
    buscado = _normal(texto)
    for nombre in _SISTEMAS_RIEGO:
        if _normal(nombre) == buscado or buscado.startswith(_normal(nombre).split()[0]):
            return nombre
    return texto


# ------------------------------------------------------------- el motor

class MotorSantander:
    """Uso::

        motor = MotorSantander(Path("templates/Formato Fitogreen Santander 2026.xlsx"),
                               Path("config/mapping_santander_2026.json"))
        resultado = motor.generar(informe, Path("output/informe.xlsx"))
    """

    def __init__(self, plantilla: Union[str, Path], mapping: Union[str, Path, Dict[str, Any]],
                 reglas=None):
        self.plantilla = Path(plantilla)
        if not self.plantilla.exists():
            raise FileNotFoundError("no existe la plantilla: {}".format(self.plantilla))
        if isinstance(mapping, dict):
            self.mapa = mapping
        else:
            self.mapa = json.loads(Path(mapping).read_text(encoding="utf-8"))
        self.hojas: Dict[str, str] = self.mapa["hojas"]
        self.reglas = reglas

    # -- API publica -------------------------------------------------------

    def generar(self, informe: InformeFitogreen, destino: Union[str, Path],
                crear_galeria_fotos: bool = False,
                excluir: Optional[Sequence[str]] = None) -> ResultadoGeneracion:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        with warnings.catch_warnings():
            # "Data Validation extension is not supported": las listas que
            # leen de otra hoja (x14); se reponen a mano en _listas_x14.
            warnings.simplefilter("ignore")
            molde = load_workbook(self.plantilla, rich_text=True)
            wb = load_workbook(self.plantilla, rich_text=True)
        self._listas_x14(molde)
        self._listas_x14(wb)

        datos = _Datos(informe, self.reglas)
        resultado = ResultadoGeneracion(ruta=destino)
        resultado.avisos.extend(datos.avisos)
        if excluir:
            resultado.avisos.append(
                "sacar hojas o tablas todavía no está disponible en el formato "
                "nuevo de Santander: el informe sale completo")
        for especie in datos.declarados_sin_catastro():
            resultado.avisos.append(
                "Anexo 1: '{}' tiene resultados declarados pero no está en el "
                "catastro; no se escribió".format(especie))

        disp = self._disposicion(molde, datos)
        corrimientos: Dict[str, List[Tuple[int, int]]] = {}
        for nombre in disp.filas:
            corr = disp.materializar(wb[nombre], molde[nombre])
            corrimientos[nombre] = corr
            resultado.filas_insertadas += max(0, disp.total_filas(nombre) - molde[nombre].max_row)
        self._copiar_columnas(wb, molde, datos)

        sin_formula = self._celdas_de_dato(disp, datos)
        self._formulas(wb, molde, disp, sin_formula, resultado)
        self._datos(wb, disp, datos, resultado)
        self._inversiones(wb, informe, disp, resultado)
        self._listas_y_formatos(wb, molde, disp)
        imagenes, textos, recortes = self._dibujos(informe, datos, wb, resultado)
        for hoja, corr in recortes.items():
            corrimientos.setdefault(hoja, []).extend(corr)

        wb.save(destino)
        filas_de = {hoja: (lambda fila, hoja=hoja: disp.filas_de(hoja, fila))
                    for hoja in disp.filas}
        for aviso in dibujos_xlsx.injertar_dibujos(
                self.plantilla, destino, textos=textos, imagenes=imagenes,
                filas_insertadas=corrimientos, filas_de=filas_de):
            resultado.avisos.append(aviso)
        self._matrices_dinamicas(destino)
        return resultado

    async def generar_async(self, informe: InformeFitogreen, destino: Union[str, Path],
                            crear_galeria_fotos: bool = False,
                            excluir: Optional[Sequence[str]] = None) -> ResultadoGeneracion:
        return await asyncio.to_thread(self.generar, informe, destino,
                                       crear_galeria_fotos, excluir)

    # -- 1. disposicion ----------------------------------------------------

    def _hoja(self, alias: str) -> str:
        return self.hojas.get(alias, alias)

    def _disposicion(self, molde, datos: _Datos) -> Disposicion:
        partes = [
            Parte(id=p["id"], hoja=self._hoja(p["hoja"]), desde=p["filas"][0],
                  hasta=p["filas"][1], dimension=p["dimension"], dentro_de=p.get("dentro_de"),
                  columnas_de_afuera=tuple(p.get("columnas_de_afuera", ())),
                  agrega=p.get("agrega"))
            for p in self.mapa["partes"]
        ]
        columnas = [
            ParteDeColumnas(id=c["id"], hoja=self._hoja(c["hoja"]), columna=c["columna"],
                            desde=c["filas"][0], hasta=c["filas"][1], dimension=c["dimension"])
            for c in self.mapa.get("columnas_repetidas", [])
        ]
        hojas = {p.hoja for p in partes} | {c.hoja for c in columnas}
        return Disposicion(partes, columnas, datos.instancias(),
                           {h: molde[h].max_row for h in hojas})

    def _copiar_columnas(self, wb, molde, datos: _Datos) -> None:
        """La columna de sensibilizacion de la primera especie, copiada a la
        derecha una vez por especie mas."""
        for c in self.mapa.get("columnas_repetidas", []):
            ws, origen = wb[self._hoja(c["hoja"])], molde[self._hoja(c["hoja"])]
            base = excel_ops.column_index_from_string(c["columna"])
            for i in range(1, len(datos.especies)):
                for fila in range(c["filas"][0], c["filas"][1] + 1):
                    fuente = origen.cell(fila, base)
                    ws.cell(fila, base + i)._style = copy(fuente._style)
                ancho = origen.column_dimensions[c["columna"]].width
                if ancho:
                    ws.column_dimensions[get_column_letter(base + i)].width = ancho

    # -- 3. formulas -------------------------------------------------------

    def _celdas_de_dato(self, disp: Disposicion, datos: _Datos) -> set:
        """Celdas del molde que traen formula pero donde manda el dato, si
        existe: (hoja, fila_final, columna)."""
        salida = set()
        for p in self.mapa["partes"]:
            hoja = self._hoja(p["hoja"])
            celdas = dict(self._escrituras(p))
            for clave in p.get("pisa_formula", []):
                letra, numero = re.match(r"([A-Z]+)(\d*)", clave).groups()
                fila = int(numero) if numero else p["filas"][0]
                campo = celdas.get((fila, letra))
                col = excel_ops.column_index_from_string(letra)
                for final, ctx in disp.ocurrencias(hoja, fila):
                    if datos.registro(p["dimension"], ctx).get(campo) is not None:
                        salida.add((hoja, final, col))
        return salida

    @staticmethod
    def _escrituras(p: Dict[str, Any]) -> List[Tuple[Tuple[int, str], str]]:
        """[((fila del molde, columna), campo)] de una parte: sus 'columnas'
        (si es de una sola fila) y sus 'celdas'."""
        salida = []
        desde, hasta = p["filas"]
        if desde == hasta:
            salida.extend(((desde, letra), campo) for letra, campo in p.get("columnas", {}).items())
        for celda, campo in p.get("celdas", {}).items():
            letra, numero = re.match(r"([A-Z]+)(\d+)", celda).groups()
            salida.append(((int(numero), letra), campo))
        return salida

    def _formulas(self, wb, molde, disp: Disposicion, sin_formula: set,
                  resultado: ResultadoGeneracion) -> None:
        for origen in molde.worksheets:
            ws = wb[origen.title]
            hoja = origen.title
            tiene_partes = hoja in disp.filas
            for fila in origen.iter_rows():
                for celda in fila:
                    valor = celda.value
                    dinamica = isinstance(valor, ArrayFormula)
                    texto = valor.text if dinamica else valor
                    if not excel_ops.es_formula(texto):
                        continue
                    destinos = disp.ocurrencias(hoja, celda.row) if tiene_partes \
                        else [(celda.row, {})]
                    for final, ctx in destinos:
                        col = celda.column
                        if (hoja, final, col) in sin_formula:
                            continue
                        if tiene_partes and disp.es_columna_de_afuera_repetida(hoja, celda.row, col, ctx):
                            continue
                        try:
                            nueva = disp.remapear(texto, hoja, celda.row, ctx, col)
                        except FaltaContexto as error:
                            resultado.avisos.append(
                                "fórmula sin resolver en {}!{}{}: {}".format(
                                    hoja.strip(), celda.column_letter, celda.row, error))
                            continue
                        destino = ws.cell(final, col)
                        if dinamica:
                            destino.value = ArrayFormula(
                                "{}{}".format(get_column_letter(col), final), nueva)
                        else:
                            destino.value = nueva

    # -- 4. datos ----------------------------------------------------------

    def _escribir(self, ws, fila: int, col: int, valor: Any, resultado) -> None:
        valor = _valor(valor)
        if valor is None or (isinstance(valor, str) and not valor.strip()):
            return
        coordenada = excel_ops.celda_ancla(ws, "{}{}".format(get_column_letter(col), fila))
        celda = ws[coordenada]
        celda.value = valor
        if isinstance(valor, date) and "y" not in (celda.number_format or "").lower():
            celda.number_format = "DD-MM-YYYY"
        resultado.celdas_escritas += 1

    def _datos(self, wb, disp: Disposicion, datos: _Datos, resultado) -> None:
        for p in self.mapa["partes"]:
            hoja = self._hoja(p["hoja"])
            ws = wb[hoja]
            for (fila, letra), campo in self._escrituras(p):
                col = excel_ops.column_index_from_string(letra)
                for final, ctx in disp.ocurrencias(hoja, fila):
                    if disp.es_columna_de_afuera_repetida(hoja, fila, col, ctx):
                        continue
                    registro = datos.registro(p["dimension"], ctx)
                    self._escribir(ws, final, col, registro.get(campo), resultado)

        for c in self.mapa.get("columnas_repetidas", []):
            ws = wb[self._hoja(c["hoja"])]
            base = excel_ops.column_index_from_string(c["columna"])
            hoja = self._hoja(c["hoja"])
            for i in range(len(datos.especies)):
                registro = datos.registro(c["dimension"], {c["dimension"]: i})
                for fila, campo in c["celdas"].items():
                    final = disp.fila(hoja, int(fila), {}) if hoja in disp.filas else int(fila)
                    ws.cell(final, base + i).value = _valor(registro.get(campo))
                    resultado.celdas_escritas += 1

        for referencia, ruta in self.mapa.get("campos", {}).items():
            hoja, celda = self._celda(referencia)
            letra, numero = re.match(r"([A-Z]+)(\d+)", celda).groups()
            fila = disp.fila(hoja, int(numero), {}) if hoja in disp.filas else int(numero)
            self._escribir(wb[hoja], fila, excel_ops.column_index_from_string(letra),
                           datos.campo(ruta), resultado)

        for referencia, ruta in self.mapa.get("listas", {}).items():
            hoja, rango = self._celda(referencia)
            valores = list(_resolver(datos.informe, ruta) or [])
            r = CellRange(rango)
            for i, fila in enumerate(range(r.min_row, r.max_row + 1)):
                if i < len(valores):
                    self._escribir(wb[hoja], fila, r.min_col, valores[i], resultado)

    def _celda(self, referencia: str) -> Tuple[str, str]:
        alias, celda = referencia.split("!", 1)
        return self._hoja(alias), celda

    def _inversiones(self, wb, informe: InformeFitogreen, disp: Disposicion, resultado) -> None:
        """Dos proyectos de hasta cinco items, como trae el formato: el nombre,
        el año y la superficie combinados a lo alto de sus items."""
        spec = self.mapa.get("inversion")
        if not spec:
            return
        ws = wb[self._hoja(spec["hoja"])]
        proyectos: Dict[str, list] = {}
        for linea in informe.inversiones.lineas:
            proyectos.setdefault(linea.proyecto, []).append(linea)
        hoja_tc, celda_tc = self._celda(spec["tipo_cambio"])
        letra, numero = re.match(r"([A-Z]+)(\d+)", celda_tc).groups()
        fila_tc = disp.fila(hoja_tc, int(numero), {}) if hoja_tc in disp.filas else int(numero)
        tc = "'{}'!${}${}".format(hoja_tc, letra, fila_tc)
        bloques = spec["bloques"]
        if len(proyectos) > len(bloques):
            resultado.avisos.append(
                "Inversión: el formato admite {} proyectos y vienen {}; los demás no "
                "se escribieron".format(len(bloques), len(proyectos)))
        for (desde, hasta), (nombre, lineas) in zip(bloques, proyectos.items()):
            primera = lineas[0]
            for letra, campo in spec["columnas_proyecto"].items():
                self._escribir(ws, desde, excel_ops.column_index_from_string(letra),
                               getattr(primera, campo, None), resultado)
            if len(lineas) > hasta - desde + 1:
                resultado.avisos.append(
                    "Inversión: '{}' trae {} ítems y el formato admite {}".format(
                        nombre, len(lineas), hasta - desde + 1))
            for fila, linea in zip(range(desde, hasta + 1), lineas):
                for letra, campo in spec["columnas_item"].items():
                    self._escribir(ws, fila, excel_ops.column_index_from_string(letra),
                                   getattr(linea, campo, None), resultado)
                for letra, formula in spec["formulas_item"].items():
                    ws["{}{}".format(letra, fila)].value = formula.format(
                        fila=fila, superficie="$C${}".format(desde), tc=tc,
                        total=spec["fila_total"])
                for letra, mascara in spec.get("formatos", {}).items():
                    ws["{}{}".format(letra, fila)].number_format = mascara

    # -- 5. listas, formatos condicionales, impresion ------------------------

    def _listas_x14(self, molde) -> None:
        """Las listas desplegables que leen de otra hoja vienen como extension
        x14, que openpyxl descarta al leer. Se reponen como listas comunes
        (Excel 2010 en adelante acepta la referencia a otra hoja ahi)."""
        partes = dibujos_xlsx._leer_paquete(self.plantilla)
        for nombre, parte in dibujos_xlsx._hojas(partes).items():
            xml = partes[parte].decode("utf-8")
            for bloque in re.findall(r"<x14:dataValidation .*?</x14:dataValidation>", xml, re.S):
                formula = re.search(r"<xm:f>([^<]*)</xm:f>", bloque)
                rangos = re.search(r"<xm:sqref>([^<]*)</xm:sqref>", bloque)
                if not formula or not rangos:
                    continue
                tipo = re.search(r'type="([^"]+)"', bloque)
                lista = DataValidation(type=tipo.group(1) if tipo else "list",
                                       formula1=formula.group(1).replace("&amp;", "&"),
                                       allow_blank=True)
                lista.sqref = MultiCellRange(rangos.group(1))
                molde[nombre].data_validations.append(lista)

    def _listas_y_formatos(self, wb, molde, disp: Disposicion) -> None:
        for nombre in disp.filas:
            ws, origen = wb[nombre], molde[nombre]

            def estirar(sqref) -> Tuple[str, Optional[Tuple[int, int]]]:
                rangos = []
                for r in sqref.ranges:
                    rangos.extend(disp.expandir_rango(nombre, r.coord))
                if not rangos:
                    return "", None
                primero = CellRange(rangos[0])
                return " ".join(rangos), (primero.min_row, primero.min_col)

            ws.data_validations.dataValidation = []
            for lista in origen.data_validations.dataValidation:
                nuevo, esquina = estirar(lista.sqref)
                if not nuevo:
                    continue
                copia = copy(lista)
                copia.sqref = MultiCellRange(nuevo)
                previa = _primero(lista.sqref)
                if esquina and (previa.min_row, previa.min_col) != esquina and copia.formula1 \
                        and not copia.formula1.startswith('"'):
                    copia.formula1 = Translator(
                        "=" + copia.formula1,
                        origin="{}{}".format(get_column_letter(previa.min_col), previa.min_row),
                    ).translate_formula(
                        "{}{}".format(get_column_letter(esquina[1]), esquina[0]))[1:]
                ws.data_validations.append(copia)

            formatos = ConditionalFormattingList()
            for regla_cf in origen.conditional_formatting:
                nuevo, esquina = estirar(regla_cf.sqref)
                if not nuevo:
                    continue
                previa = _primero(regla_cf.sqref)
                for regla in regla_cf.rules:
                    regla = copy(regla)
                    if esquina and (previa.min_row, previa.min_col) != esquina and regla.formula:
                        regla.formula = [
                            Translator("=" + f, origin="{}{}".format(
                                get_column_letter(previa.min_col), previa.min_row)
                            ).translate_formula("{}{}".format(
                                get_column_letter(esquina[1]), esquina[0]))[1:]
                            for f in regla.formula]
                    formatos.add(nuevo, regla)
            ws.conditional_formatting = formatos

            # Los saltos de pagina manuales (el del flujo, antes del resumen)
            # van a donde quedo su fila.
            saltos = [b.id for b in origen.row_breaks.brk]
            if saltos:
                from openpyxl.worksheet.pagebreak import Break, RowBreak
                ws.row_breaks = RowBreak()
                for fila in saltos:
                    finales = disp.filas_de(nombre, fila)
                    if finales:
                        ws.row_breaks.append(Break(id=max(finales)))

            area = origen.print_area
            if area:
                rango = str(area).split("!")[-1].replace("$", "")
                partes = disp.expandir_rango(nombre, rango)
                if partes:
                    inicio, fin = CellRange(partes[0]), CellRange(partes[-1])
                    ws.print_area = "{}{}:{}{}".format(
                        get_column_letter(inicio.min_col), inicio.min_row,
                        get_column_letter(fin.max_col), fin.max_row)

    # -- 6. dibujos y formulas dinamicas -------------------------------------

    def _dibujos(self, informe: InformeFitogreen, datos: _Datos, wb,
                 resultado: ResultadoGeneracion):
        """Lo que se pone en los dibujos despues de guardar -fotos y cuadros
        de texto-; los titulos de las fotos van en celdas, asi que se
        escriben ahora."""
        textos = {}
        recortes: Dict[str, List[Tuple[int, int]]] = {}
        for referencia, ruta in self.mapa.get("cuadros_de_texto", {}).items():
            hoja, forma = self._celda(referencia)
            textos[(hoja, forma)] = datos.campo(ruta) or ""

        imagenes: Dict[Tuple[str, str], Optional[bytes]] = {}
        spec = self.mapa.get("imagenes", {})
        fotos = list(informe.fotografias or [])
        emplazamiento = next((f for f in fotos if f.categoria == CategoriaFoto.EMPLAZAMIENTO), None)
        if "emplazamiento" in spec:
            hueco = (self._hoja(spec["emplazamiento"]["hoja"]), spec["emplazamiento"]["forma"])
            imagenes[hueco] = _leer(emplazamiento.ruta) if emplazamiento else None
            if emplazamiento is None:
                resultado.avisos.append(
                    "sin imagen de emplazamiento: la portada sale sin ella")
            else:
                resultado.imagenes_insertadas += 1

        galeria = spec.get("galeria")
        if galeria:
            hoja = self._hoja(galeria["hoja"])
            excluidas = {CategoriaFoto.EMPLAZAMIENTO, CategoriaFoto.LEGAL}
            propias = [f for f in fotos if f.categoria not in excluidas]
            if len(propias) > len(galeria["formas"]):
                resultado.avisos.append(
                    "el Anexo 4 admite {} fotografías y vienen {}: van las primeras".format(
                        len(galeria["formas"]), len(propias)))
            ws = wb[hoja]
            for i, forma in enumerate(galeria["formas"]):
                foto = propias[i] if i < len(propias) else None
                imagenes[(hoja, forma)] = _leer(foto.ruta) if foto else None
                if foto is not None:
                    resultado.imagenes_insertadas += 1
                    if foto.titulo:
                        ws[galeria["titulos"][i]].value = foto.titulo
            recortes[hoja] = self._recortar_galeria(ws, galeria, len(propias))
        return imagenes, textos, recortes

    @staticmethod
    def _recortar_galeria(ws, galeria: Dict[str, Any], fotos: int) -> List[Tuple[int, int]]:
        """Las filas de fotos que no se usan se sacan: una grilla de quince
        marcos con tres fotos se ve como un informe a medio hacer. Queda
        siempre una fila (los marcos vacios de la ultima fila usada quedan)."""
        desde, hasta = galeria["filas"]
        por_fila, alto = galeria["fotos_por_fila"], galeria["filas_por_fila"]
        usadas = max(1, -(-fotos // por_fila))
        primera_sobrante = desde + usadas * alto
        if primera_sobrante > hasta:
            return []
        cantidad = hasta - primera_sobrante + 1
        excel_ops.eliminar_filas(ws, primera_sobrante, cantidad)
        area = ws.print_area
        if area:
            rango = CellRange(str(area).split("!")[-1].replace("$", ""))
            ws.print_area = "{}{}:{}{}".format(
                get_column_letter(rango.min_col), rango.min_row,
                get_column_letter(rango.max_col), rango.max_row - cantidad)
        return [(hasta + 1, -cantidad)]

    def _matrices_dinamicas(self, destino: Path) -> None:
        """openpyxl guarda UNIQUE/FILTER como formula matricial clasica (la de
        Ctrl+Mayus+Enter). Se les devuelve el marcado de formula dinamica
        -cm="1" y la parte xl/metadata.xml- que traia la plantilla."""
        molde = dibujos_xlsx._leer_paquete(self.plantilla)
        if "xl/metadata.xml" not in molde:
            return
        partes = dibujos_xlsx._leer_paquete(destino)
        cambio = False
        for nombre in [n for n in partes if re.match(r"xl/worksheets/sheet\d+\.xml$", n)]:
            xml = partes[nombre].decode("utf-8")

            def marcar(m):
                celda, cuerpo = m.group(1), m.group(2)
                if ' cm="' in celda or not any(f in cuerpo for f in _FUNCIONES_DINAMICAS):
                    return m.group(0)
                return celda[:-1] + ' cm="1">' + cuerpo

            nuevo = re.sub(r'(<c [^>]*>)(<f t="array"[^>]*>.*?</f>)', marcar, xml, flags=re.S)
            if nuevo != xml:
                partes[nombre] = nuevo.encode("utf-8")
                cambio = True
        if not cambio:
            return
        partes["xl/metadata.xml"] = molde["xl/metadata.xml"]
        rels = partes["xl/_rels/workbook.xml.rels"].decode("utf-8")
        if "sheetMetadata" not in rels:
            rels = rels.replace(
                "</Relationships>",
                '<Relationship Id="rIdFgMeta" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/sheetMetadata" Target="metadata.xml"/>'
                "</Relationships>")
            partes["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
        tipos = partes["[Content_Types].xml"].decode("utf-8")
        if "/xl/metadata.xml" not in tipos:
            tipos = tipos.replace(
                "</Types>",
                '<Override PartName="/xl/metadata.xml" ContentType="application/'
                'vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
            partes["[Content_Types].xml"] = tipos.encode("utf-8")
        dibujos_xlsx._escribir_paquete(partes, destino)


def _leer(ruta: Optional[str]) -> Optional[bytes]:
    if not ruta:
        return None
    try:
        return Path(ruta).read_bytes()
    except OSError:
        return None


def _primero(sqref) -> CellRange:
    """El rango de arriba a la izquierda de un sqref (openpyxl los guarda en un
    conjunto, sin orden): es la celda a la que Excel refiere las formulas
    relativas de una lista o de un formato condicional."""
    return sorted(sqref.ranges, key=lambda r: (r.min_row, r.min_col))[0]
