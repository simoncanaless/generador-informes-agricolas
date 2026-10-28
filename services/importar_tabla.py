"""Llena una seccion de la visita desde la planilla que manda el cliente.

El cuello de botella del informe no son las visitas: es tipear lo que el
cliente ya tiene escrito (§3). Maquinaria, resultados de temporada,
comercializacion, inversiones — todo eso llega por correo en un Excel y hoy se
copia a mano celda por celda.

**Esto no reemplaza al tasador, le ahorra el tecleo.** Lo importado entra al
formulario marcado como SIN REVISAR: se ve distinto, se cuenta en el aviso de
lo que falta, y deja de estarlo en cuanto alguien toca el campo. La regla es la
de siempre: el dato del cliente no es un dato verificado, y quien firma el
informe tiene que haberlo mirado.

No hay un segundo motor: se reusa el del catastro. `_emparejar` ya sabe
resolver 'Sup. (ha)' contra 'superficie_ha' con cuatro pasadas de tolerancia
creciente, y esa logica costo bastante como para tener dos copias. Lo unico
propio de aqui es que los campos esperados salen de `formulario_visita.json`
en vez de estar fijos.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from services.catastro_import import (
    CatastroIlegible, _emparejar, limpiar_numero, limpiar_texto, normalizar,
)

__all__ = [
    "TablaIlegible", "ColumnaReconocida", "TablaImportada",
    "importar_seccion",
]

#: Con menos de esto no se da por encontrada la tabla. Dos es poco exigente a
#: proposito: hay secciones de tres columnas (maquinaria) y pedir tres dejaria
#: fuera una planilla que solo traiga la maquina y la cantidad.
MINIMO_COLUMNAS = 2

#: Hasta donde se busca la fila de encabezados. Las planillas de agricultor
#: suelen traer titulo, logo y un par de lineas en blanco antes de la tabla.
MAXIMO_FILAS_ENCABEZADO = 25


class TablaIlegible(RuntimeError):
    """No se encontro una tabla que se parezca a la seccion pedida."""


@dataclass
class ColumnaReconocida:
    campo: str
    etiqueta_esperada: str
    encabezado_en_el_archivo: str
    columna: str


@dataclass
class TablaImportada:
    seccion: str
    hoja: str
    fila_encabezado: int
    filas: List[Dict[str, Any]] = field(default_factory=list)
    reconocidas: List[ColumnaReconocida] = field(default_factory=list)
    #: Encabezados del archivo que no se supo a que campo llevar. Se devuelven
    #: para poder decirlo: una columna ignorada en silencio es un dato perdido.
    sin_reconocer: Dict[str, str] = field(default_factory=dict)
    #: Campos de la seccion que la planilla no traia.
    faltantes: List[str] = field(default_factory=list)
    avisos: List[str] = field(default_factory=list)

    @property
    def cuantas(self) -> int:
        return len(self.filas)


def _sinonimos_de(seccion: Dict[str, Any]) -> Dict[str, List[str]]:
    """Lo que la definicion del formulario declara para cada campo."""
    return {c["id"]: list(c.get("sinonimos") or [])
            for c in seccion.get("campos", []) if c.get("sinonimos")}


def _detectar(ws, sinonimos: Dict[str, List[str]]):
    """La fila con mas encabezados reconocidos, y a que columna va cada campo.

    Se mira fila por fila en vez de asumir la primera: las planillas reales
    traen titulo y subtitulo arriba, y a veces una fila de unidades debajo.
    """
    mejor = (0, 0, {}, {}, {})
    tope = min(MAXIMO_FILAS_ENCABEZADO, ws.max_row or 0)

    for fila in range(1, tope + 1):
        columnas: Dict[str, int] = {}
        textos: Dict[str, str] = {}
        sueltas: Dict[str, str] = {}
        for columna in range(1, (ws.max_column or 0) + 1):
            etiqueta = limpiar_texto(ws.cell(fila, columna).value)
            if not etiqueta:
                continue
            campo = _emparejar(etiqueta, sinonimos)
            if campo and campo not in columnas:
                columnas[campo] = columna
                textos[campo] = etiqueta
            elif campo is None:
                sueltas[get_column_letter(columna)] = etiqueta
        if len(columnas) > mejor[1]:
            mejor = (fila, len(columnas), columnas, textos, sueltas)

    return mejor[0], mejor[2], mejor[3], mejor[4]


def _valor(celda: Any, tipo: str) -> Any:
    if tipo == "numero":
        return limpiar_numero(celda)
    return limpiar_texto(celda)


def importar_seccion(
    origen: Union[str, Path, bytes],
    seccion: Dict[str, Any],
    maximo_filas: Optional[int] = None,
) -> TablaImportada:
    """Lee la planilla y devuelve las filas que calzan con esa seccion.

    ``seccion`` es la entrada de `formulario_visita.json`, tal cual: de ahi
    salen los campos, sus tipos y sus sinonimos.
    """
    if seccion.get("tipo") != "tabla":
        raise TablaIlegible(
            "la sección «{}» no es una tabla: solo se pueden importar las que "
            "tienen filas".format(seccion.get("id")))

    sinonimos = _sinonimos_de(seccion)
    if not sinonimos:
        raise TablaIlegible(
            "la sección «{}» no declara sinónimos, así que no hay forma de "
            "saber qué columna es cada cosa".format(seccion.get("id")))

    try:
        libro = load_workbook(origen, data_only=True, read_only=False)
    except Exception as error:                       # noqa: BLE001
        raise TablaIlegible(
            "no se pudo abrir la planilla: {}".format(error)) from error

    # Se prueban TODAS las hojas y gana la que mas columnas reconoce. Una
    # planilla de cliente suele traer varias -"Maquinaria", "Riego", "Hoja1"-
    # y pedirle que diga cual es seria pedirle que entienda nuestro modelo.
    mejor_hoja, mejor = None, (0, 0, {}, {}, {})
    for ws in libro.worksheets:
        fila, columnas, textos, sueltas = _detectar(ws, sinonimos)
        if len(columnas) > mejor[1]:
            mejor_hoja = ws
            mejor = (fila, len(columnas), columnas, textos, sueltas)

    fila_encabezado, cuantas, columnas, textos, sueltas = mejor
    if mejor_hoja is None or cuantas < MINIMO_COLUMNAS:
        raise TablaIlegible(
            "no se reconoció una tabla de «{}» en esta planilla: solo {} "
            "columna(s) calzaron. Revise que sea el archivo correcto.".format(
                seccion.get("titulo", seccion.get("id")), cuantas))

    tipos = {c["id"]: c.get("tipo", "texto") for c in seccion["campos"]}
    resultado = TablaImportada(
        seccion=seccion["id"], hoja=mejor_hoja.title,
        fila_encabezado=fila_encabezado,
        reconocidas=[
            ColumnaReconocida(
                campo=campo,
                etiqueta_esperada=next(
                    (c["etiqueta"] for c in seccion["campos"] if c["id"] == campo),
                    campo),
                encabezado_en_el_archivo=textos[campo],
                columna=get_column_letter(columnas[campo]))
            for campo in columnas],
        sin_reconocer=sueltas,
        faltantes=[c["etiqueta"] for c in seccion["campos"]
                   if c["id"] not in columnas],
    )

    tope = maximo_filas or seccion.get("filas") or 0
    for numero in range(fila_encabezado + 1, (mejor_hoja.max_row or 0) + 1):
        fila = {}
        for campo, columna in columnas.items():
            valor = _valor(mejor_hoja.cell(numero, columna).value,
                           tipos.get(campo, "texto"))
            if valor is not None and valor != "":
                fila[campo] = valor
        if not fila:
            continue
        # Una fila de totales no es un item del inventario. Se descarta por el
        # rotulo y no por la posicion: no siempre esta al final.
        primera = str(next(iter(fila.values()), "")).strip()
        if normalizar(primera) in {"total", "totales", "suma", "subtotal"}:
            continue
        resultado.filas.append(fila)

    if tope and len(resultado.filas) > tope:
        resultado.avisos.append(
            "La planilla trae {} filas y la sección admite {}. Se cargaron las "
            "primeras {}; el resto hay que agregarlo a mano o ampliar la "
            "sección.".format(len(resultado.filas), tope, tope))
        resultado.filas = resultado.filas[:tope]

    if resultado.sin_reconocer:
        resultado.avisos.append(
            "Estas columnas del archivo no se usaron: {}.".format(
                ", ".join("{} ({})".format(t, c)
                          for c, t in sorted(resultado.sin_reconocer.items()))))
    if resultado.faltantes:
        resultado.avisos.append(
            "La planilla no traía: {}. Esos campos quedan en blanco.".format(
                ", ".join(resultado.faltantes)))
    if not resultado.filas:
        resultado.avisos.append(
            "Se reconoció la tabla pero no había ninguna fila con datos "
            "debajo del encabezado.")

    return resultado


def claves_sin_revisar(seccion_id: str, filas: Sequence[Dict[str, Any]]
                       ) -> List[str]:
    """Las claves `seccion|fila|campo` de todo lo que entro desde el archivo.

    Es el mismo formato con que la visita ya identifica un campo al guardar
    (§16.a.vicies.nonies), asi que tocar el campo lo saca de la lista sin
    ningun mecanismo nuevo.
    """
    return ["{}|{}|{}".format(seccion_id, numero, campo)
            for numero, fila in enumerate(filas) for campo in fila]
