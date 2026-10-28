"""Disposicion de un formato con partes que se repiten.

El formato final de Santander (templates/Formato Fitogreen Santander 2026.xlsx)
trae UNA vez cada parte que se repite -un bloque de flujo por plantacion, una
fila por especie, un bloque de riego por unidad predial- y el generador las
repite segun el informe (CLAUDE.md §16.a.centies.sexiesdecies). Este modulo
hace lo que eso exige, sin saber nada del dominio ni de Santander:

1. **A que fila va cada fila del molde.** La hoja final es una lista de filas;
   cada una dice de que fila del molde sale y su CONTEXTO -que plantacion, que
   especie- que la distingue de las otras copias de la misma fila.
2. **Reescribir las formulas del molde para cada copia**, con tres reglas:
   - una referencia a la fila modelo de otra parte que se repite apunta a la
     copia de ESE contexto: el bloque de la plantacion 5 lee la fila 5 de
     Supuestos, y la fila del cerezo en el resumen lee la del cerezo en el
     Anexo 1;
   - un rango sobre una parte que se repite, visto desde donde no hay
     contexto para elegir una copia, se estira hasta cubrirlas todas: el total
     suma todas las especies;
   - lo que no se repite se corre lo que se corrio la hoja.
   Y una cuarta para los totales por especie del flujo: una fila que suma una
   dimension que no tiene (la especie suma sus plantaciones) se escribe como
   la suma de sus copias -=F9+F19+F39-, generada, no escrita a mano.
3. **Armar las filas en la hoja**: borrar la parte del molde, abrir el hueco
   del tamano final y copiar cada fila de su fila modelo -estilo, alto,
   rotulos, celdas combinadas-. Las formulas no se copian: se escriben todas
   despues, ya con la disposicion final conocida.

Por que no se usa la copia de Excel (insertar filas y dejar que las formulas
se corran solas): Excel corre las referencias relativas y deja quietas las
absolutas, y aqui hace falta lo contrario muchas veces -el bloque de una
plantacion lee Supuestos!$F$24 y tiene que leer la fila de SU plantacion-.
el agrónomo lo hacia a mano en el informe de origen, bloque por bloque.
"""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from typing import Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from services import excel_ops

__all__ = ["Parte", "ParteDeColumnas", "Disposicion", "FaltaContexto"]

Contexto = Dict[str, int]


@dataclass
class Parte:
    """Un tramo de filas del molde que se repite.

    ``dimension``: lo que distingue a cada copia ('plantacion', 'especie'...).
    ``dentro_de``: id de la parte que la contiene (las filas de plantacion de
    cada tabla de especie en Supuestos).
    ``columnas_de_afuera``: columnas de una parte de adentro que pertenecen a
    la de afuera -en el riego, la demanda del predio ocupa una celda
    combinada a lo alto de todas sus especies-.
    ``agrega``: la dimension que sus formulas suman cuando apuntan a una parte
    que tiene una dimension mas: el resumen por especie del flujo suma los
    bloques de sus plantaciones.
    """

    id: str
    hoja: str
    desde: int
    hasta: int
    dimension: str
    dentro_de: Optional[str] = None
    columnas_de_afuera: Tuple[str, ...] = ()
    agrega: Optional[str] = None


@dataclass
class ParteDeColumnas:
    """Una columna del molde que se repite hacia la derecha (una por especie
    en la sensibilizacion de Supuestos). No se insertan columnas -romperian las
    tablas de abajo-: se usa el espacio libre a la derecha."""

    id: str
    hoja: str
    columna: str
    desde: int
    hasta: int
    dimension: str


class FaltaContexto(Exception):
    """Una referencia a una fila que se repite, desde una celda que no sabe
    cual de sus copias le corresponde."""

    def __init__(self, dimension: str, referencia: str):
        super().__init__("{} necesita saber '{}'".format(referencia, dimension))
        self.dimension = dimension
        self.referencia = referencia


# ------------------------------------------------------------- formulas

_RE_CADENA = re.compile(r'"[^"]*"')
_RE_REF = re.compile(
    r"(?<![A-Za-z0-9_.$'!\]])"
    r"(?P<hoja>(?:'(?:[^']|'')+'|[A-Za-z_À-ɏ][\w.À-ɏ]*)!)?"
    r"(?P<c1>\$?[A-Z]{1,3}\$?\d+)"
    r"(?::(?P<c2>\$?[A-Z]{1,3}\$?\d+))?"
    r"(?![\w(!])"
)
_RE_CELDA = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")


@dataclass
class _Celda:
    col_abs: bool
    col: int
    fila_abs: bool
    fila: int

    @classmethod
    def leer(cls, texto: str) -> "_Celda":
        m = _RE_CELDA.match(texto)
        return cls(bool(m.group(1)), column_index_from_string(m.group(2)),
                   bool(m.group(3)), int(m.group(4)))

    def texto(self) -> str:
        return "{}{}{}{}".format("$" if self.col_abs else "", get_column_letter(self.col),
                                 "$" if self.fila_abs else "", self.fila)


def _nombre_de_hoja(prefijo: str) -> str:
    nombre = prefijo[:-1]
    if nombre.startswith("'") and nombre.endswith("'"):
        nombre = nombre[1:-1].replace("''", "'")
    return nombre


# --------------------------------------------------------------- disposicion

class Disposicion:
    """La hoja final de cada hoja del molde, fila por fila.

    ``instancias[id_parte](contexto_de_afuera)`` devuelve la lista de
    contextos de sus copias (cada uno con su propia dimension, y si sirve
    alguna mas: un bloque de plantacion tambien dice su especie, para que su
    formula de sensibilizacion encuentre la columna de esa especie).
    """

    def __init__(
        self,
        partes: Sequence[Parte],
        columnas: Sequence[ParteDeColumnas],
        instancias: Dict[str, Callable[[Contexto], List[Contexto]]],
        filas_por_hoja: Dict[str, int],
    ):
        self.partes = list(partes)
        self.columnas = list(columnas)
        self.instancias = instancias
        self._por_id = {p.id: p for p in self.partes}
        # dimension -> {valor: contexto}: la misma plantacion aparece en
        # varias partes (su fila del catastro, su bloque, su fila de
        # Supuestos) y se cuenta una vez.
        self._contextos_por_dimension: Dict[str, Dict[int, Contexto]] = {}
        self._max: Dict[str, int] = dict(filas_por_hoja)
        # hoja -> lista de (fila_molde, contexto); el indice + 1 es la fila final
        self.filas: Dict[str, List[Tuple[int, Contexto]]] = {}
        self._indice: Dict[str, Dict[Tuple[int, Tuple], int]] = {}
        self._por_fila: Dict[str, Dict[int, List[Tuple[int, Contexto]]]] = {}
        for hoja, total in filas_por_hoja.items():
            filas = self._expandir(hoja, 1, total, {}, None)
            self.filas[hoja] = filas
            indice, por_fila = {}, {}
            for i, (fila, ctx) in enumerate(filas, start=1):
                indice[(fila, self._clave(hoja, fila, ctx))] = i
                por_fila.setdefault(fila, []).append((i, ctx))
            self._indice[hoja] = indice
            self._por_fila[hoja] = por_fila

    # -- armado de la lista ----------------------------------------------

    def _hijas(self, hoja: str, padre: Optional[str]) -> List[Parte]:
        return sorted((p for p in self.partes if p.hoja == hoja and p.dentro_de == padre),
                      key=lambda p: p.desde)

    def _expandir(self, hoja: str, desde: int, hasta: int, ctx: Contexto,
                  padre: Optional[str]) -> List[Tuple[int, Contexto]]:
        salida: List[Tuple[int, Contexto]] = []
        partes = {p.desde: p for p in self._hijas(hoja, padre) if desde <= p.desde <= hasta}
        fila = desde
        while fila <= hasta:
            parte = partes.get(fila)
            if parte is None:
                salida.append((fila, ctx))
                fila += 1
                continue
            for propia in self.instancias[parte.id](ctx):
                completo = dict(ctx)
                completo.update(propia)
                vistos = self._contextos_por_dimension.setdefault(parte.dimension, {})
                vistos.setdefault(completo[parte.dimension], {}).update(completo)
                salida.extend(self._expandir(hoja, parte.desde, parte.hasta, completo, parte.id))
            fila = parte.hasta + 1
        return salida

    def cadena(self, hoja: str, fila: int) -> List[Parte]:
        """Las partes que contienen esa fila, de afuera hacia adentro."""
        cadena = [p for p in self.partes if p.hoja == hoja and p.desde <= fila <= p.hasta]
        orden = []
        actual = None
        while True:
            siguiente = [p for p in cadena if p.dentro_de == actual]
            if not siguiente:
                return orden
            orden.append(siguiente[0])
            actual = siguiente[0].id

    def _clave(self, hoja: str, fila: int, ctx: Contexto) -> Tuple:
        return tuple((p.dimension, ctx.get(p.dimension)) for p in self.cadena(hoja, fila))

    # -- consultas --------------------------------------------------------

    def fila(self, hoja: str, fila: int, ctx: Contexto) -> int:
        """La fila final de esa fila del molde en ese contexto."""
        if fila > self._max[hoja]:
            # Mas abajo de lo que ocupa el molde: se corre lo que crecio la hoja.
            return fila + len(self.filas[hoja]) - self._max[hoja]
        faltan = [p.dimension for p in self.cadena(hoja, fila) if p.dimension not in ctx]
        if faltan:
            raise FaltaContexto(faltan[-1], "{}!{}".format(hoja.strip(), fila))
        clave = (fila, self._clave(hoja, fila, ctx))
        if clave not in self._indice[hoja]:
            raise FaltaContexto("(sin copias)", "{}!{}".format(hoja.strip(), fila))
        return self._indice[hoja][clave]

    def filas_de(self, hoja: str, fila: int, ctx: Optional[Contexto] = None) -> List[int]:
        """Todas las copias de esa fila que concuerdan con el contexto dado
        (en las dimensiones que el contexto trae)."""
        ctx = ctx or {}
        return [i for i, c in self._por_fila.get(hoja, {}).get(fila, [])
                if all(c.get(k) == v for k, v in ctx.items() if k in c)]

    def ocurrencias(self, hoja: str, fila: int) -> List[Tuple[int, Contexto]]:
        return list(self._por_fila.get(hoja, {}).get(fila, []))

    def contextos(self, dimension: str) -> List[Contexto]:
        vistos = self._contextos_por_dimension.get(dimension, {})
        return [dict(vistos[k]) for k in sorted(vistos)]

    def total_filas(self, hoja: str) -> int:
        return len(self.filas[hoja])

    def columna(self, hoja: str, col: int, fila: int, ctx: Contexto) -> int:
        for parte in self.columnas:
            base = column_index_from_string(parte.columna)
            if parte.hoja == hoja and col == base and parte.desde <= fila <= parte.hasta \
                    and parte.dimension in ctx:
                return base + ctx[parte.dimension]
        return col

    # -- formulas ---------------------------------------------------------

    def remapear(self, formula: str, hoja: str, fila: int, ctx: Contexto,
                 columna: Optional[int] = None) -> str:
        """La formula de la celda (hoja, fila, columna) del molde, escrita para
        la copia de contexto ``ctx``."""
        if not formula.startswith("="):
            return formula
        self._de_la_parte_de_afuera = None
        propia = self.cadena(hoja, fila)
        if columna is not None and propia and                 get_column_letter(columna) in propia[-1].columnas_de_afuera:
            # Una celda de la parte de afuera en una fila de adentro (la
            # demanda del predio, combinada a lo alto de sus especies): sus
            # rangos sobre las filas de adentro cubren TODAS las del grupo.
            self._de_la_parte_de_afuera = propia[-1]
        try:
            return self._remapear(formula, hoja, fila, ctx)
        except FaltaContexto as falta:
            propia = self.cadena(hoja, fila)
            agrega = propia[-1].agrega if propia else None
            if agrega != falta.dimension:
                raise
            # La fila suma una dimension que no tiene: una copia del cuerpo
            # por cada contexto compatible, sumadas.
            cuerpo = formula[1:].lstrip("+")
            sumandos = [
                self._remapear("=" + cuerpo, hoja, fila, dict(ctx, **otro))[1:]
                for otro in self.contextos(agrega)
                if all(otro.get(k) == v for k, v in ctx.items() if k in otro)
            ]
            return "=" + ("+".join(sumandos) if sumandos else "0")

    def _remapear(self, formula: str, hoja: str, fila: int, ctx: Contexto) -> str:
        literales: List[str] = []

        def guardar(m):
            literales.append(m.group(0))
            return "\x00{}\x00".format(len(literales) - 1)

        protegida = _RE_CADENA.sub(guardar, formula)
        propia = self.cadena(hoja, fila)

        def sustituir(m):
            prefijo = m.group("hoja") or ""
            destino = _nombre_de_hoja(prefijo) if prefijo else hoja
            if destino not in self.filas:
                return m.group(0)
            c1 = _Celda.leer(m.group("c1"))
            if not m.group("c2"):
                return prefijo + self._celda(destino, c1, ctx).texto()
            c2 = _Celda.leer(m.group("c2"))
            inicio, fin = self._rango(destino, c1, c2, ctx, hoja, propia)
            return prefijo + inicio.texto() + ":" + fin.texto()

        resultado = _RE_REF.sub(sustituir, protegida)
        for i, literal in enumerate(literales):
            resultado = resultado.replace("\x00{}\x00".format(i), literal)
        return resultado

    def _celda(self, hoja: str, c: _Celda, ctx: Contexto) -> _Celda:
        nueva = copy(c)
        nueva.fila = self.fila(hoja, c.fila, ctx)
        nueva.col = self.columna(hoja, c.col, c.fila, ctx)
        return nueva

    def _de_afuera(self, hoja: str, fila: int, ctx: Contexto) -> Contexto:
        """El contexto que fija las partes de AFUERA de la que contiene la
        fila: un rango sobre las filas de plantacion de Supuestos, visto desde
        el total de una especie, cubre las de ESA especie; visto desde fuera de
        toda tabla, todas. Las dimensiones de mas que traiga el contexto (la
        especie de una fila del resumen) no recortan una tabla que no esta
        adentro de nada: la lista de plantaciones se recorre entera."""
        cadena = self.cadena(hoja, fila)
        return {p.dimension: ctx[p.dimension] for p in cadena[:-1] if p.dimension in ctx}

    def _extremo(self, hoja: str, c: _Celda, ctx: Contexto, minimo: bool) -> _Celda:
        afuera = getattr(self, "_de_la_parte_de_afuera", None)
        if afuera is not None and afuera in self.cadena(hoja, c.fila):
            ctx = {k: v for k, v in ctx.items() if k != afuera.dimension}
        try:
            return self._celda(hoja, c, ctx)
        except FaltaContexto:
            filas = self.filas_de(hoja, c.fila, self._de_afuera(hoja, c.fila, ctx))
            if not filas:
                raise
            nueva = copy(c)
            nueva.fila = min(filas) if minimo else max(filas)
            return nueva

    def _rango(self, hoja: str, c1: _Celda, c2: _Celda, ctx: Contexto,
               hoja_propia: str, propia: List[Parte]) -> Tuple[_Celda, _Celda]:
        # =ROWS($A$35:A35): el comienzo fijo es la PRIMERA copia de la parte
        # (dentro de la de afuera), el final la copia actual.
        cadena = self.cadena(hoja, c1.fila)
        if (hoja == hoja_propia and c1.fila == c2.fila and c1.fila_abs and not c2.fila_abs
                and cadena and propia and cadena[-1] is propia[-1]):
            inicio = copy(c1)
            inicio.fila = min(self.filas_de(hoja, c1.fila, self._de_afuera(hoja, c1.fila, ctx)))
            inicio.col = self.columna(hoja, c1.col, c1.fila, ctx)
            return inicio, self._celda(hoja, c2, ctx)
        return self._extremo(hoja, c1, ctx, True), self._extremo(hoja, c2, ctx, False)

    # -- rangos sin contexto: listas desplegables, formatos, impresion -----

    def expandir_rango(self, hoja: str, rango: str) -> List[str]:
        """Un rango del molde, en la hoja final: todas las copias de sus filas,
        en tramos continuos."""
        r = CellRange(rango)
        finales = sorted({f for fila in range(r.min_row, r.max_row + 1)
                          for f in self.filas_de(hoja, fila)})
        tramos: List[Tuple[int, int]] = []
        for f in finales:
            if tramos and f == tramos[-1][1] + 1:
                tramos[-1] = (tramos[-1][0], f)
            else:
                tramos.append((f, f))
        c1, c2 = get_column_letter(r.min_col), get_column_letter(r.max_col)
        return ["{}{}:{}{}".format(c1, a, c2, b) if (a, c1) != (b, c2) else "{}{}".format(c1, a)
                for a, b in tramos]

    # -- armar la hoja ----------------------------------------------------

    def materializar(self, ws: Worksheet, molde: Worksheet) -> List[Tuple[int, int]]:
        """Deja ``ws`` (una copia de ``molde``) con la disposicion final.

        Devuelve los corrimientos aplicados, [(fila, cantidad)] en el orden en
        que se aplicaron, para correr los dibujos igual
        (dibujos_xlsx.injertar_dibujos(filas_insertadas=...)).
        """
        hoja = molde.title
        lista = self.filas[hoja]
        corrimientos: List[Tuple[int, int]] = []
        ultima_col = molde.max_column
        for parte in sorted(self._hijas(hoja, None), key=lambda p: -p.desde):
            alto = parte.hasta - parte.desde + 1
            # La parte ocupa en la lista final un tramo continuo que empieza
            # donde empezaba en el molde: lo de arriba todavia no se movio.
            inicio = self._inicio_en_lista(hoja, parte)
            tramo = lista[inicio - 1:inicio - 1 + self._largo(hoja, parte)]
            excel_ops.eliminar_filas(ws, parte.desde, alto, formulas=False)
            if tramo:
                excel_ops.insertar_filas(ws, parte.desde, len(tramo), None, formulas=False)
            corrimientos.append((parte.hasta + 1, len(tramo) - alto))
            for i, (fila_molde, _) in enumerate(tramo):
                _copiar_fila(molde, fila_molde, ws, parte.desde + i, ultima_col)
            self._combinar(ws, molde, parte, tramo, parte.desde)
        return corrimientos

    def _largo(self, hoja: str, parte: Parte) -> int:
        cuenta = 0
        dentro = False
        for fila, _ in self.filas[hoja]:
            if parte.desde <= fila <= parte.hasta:
                dentro = True
                cuenta += 1
            elif dentro:
                break
        return cuenta

    def _inicio_en_lista(self, hoja: str, parte: Parte) -> int:
        for i, (fila, _) in enumerate(self.filas[hoja], start=1):
            if parte.desde <= fila <= parte.hasta:
                return i
        return len(self.filas[hoja]) + 1

    def _combinar(self, ws: Worksheet, molde: Worksheet, parte: Parte,
                  tramo: List[Tuple[int, Contexto]], base: int) -> None:
        """Las celdas combinadas del molde dentro de la parte, en cada copia; y
        las columnas de la parte de afuera, a lo alto de las filas de adentro."""
        posicion = {(f, self._clave(molde.title, f, c)): base + i
                    for i, (f, c) in enumerate(tramo)}
        for rango in list(molde.merged_cells.ranges):
            if not (parte.desde <= rango.min_row and rango.max_row <= parte.hasta):
                continue
            for i, (fila, ctx) in enumerate(tramo):
                if fila != rango.min_row:
                    continue
                fin = posicion.get((rango.max_row, self._clave(molde.title, rango.max_row, ctx)))
                if fin is None:
                    continue
                ws.merge_cells(start_row=base + i, start_column=rango.min_col,
                               end_row=fin, end_column=rango.max_col)
        for hija in (p for p in self.partes if p.hoja == molde.title):
            if not hija.columnas_de_afuera or not (parte.desde <= hija.desde <= parte.hasta):
                continue
            grupos: Dict[Tuple, List[int]] = {}
            for i, (fila, ctx) in enumerate(tramo):
                if hija.desde <= fila <= hija.hasta:
                    afuera = tuple(sorted((k, v) for k, v in ctx.items()
                                          if k != hija.dimension))
                    grupos.setdefault(afuera, []).append(base + i)
            for filas in grupos.values():
                if len(filas) < 2:
                    continue
                for letra in hija.columnas_de_afuera:
                    col = column_index_from_string(letra)
                    for f in filas[1:]:
                        celda = ws.cell(f, col)
                        if not isinstance(celda, MergedCell):
                            celda.value = None
                    ws.merge_cells(start_row=filas[0], start_column=col,
                                   end_row=filas[-1], end_column=col)

    def es_columna_de_afuera_repetida(self, hoja: str, fila: int, col: int,
                                      ctx: Contexto) -> bool:
        """Una celda de la parte de afuera que cae en una fila de adentro que
        no es la primera de su grupo: no se escribe (queda combinada)."""
        cadena = self.cadena(hoja, fila)
        if not cadena:
            return False
        hija = cadena[-1]
        if get_column_letter(col) not in hija.columnas_de_afuera:
            return False
        afuera = {k: v for k, v in ctx.items() if k != hija.dimension}
        return self.fila(hoja, fila, ctx) != min(self.filas_de(hoja, fila, afuera))


def _copiar_fila(molde: Worksheet, origen: int, ws: Worksheet, destino: int,
                 ultima_col: int) -> None:
    """Estilo, alto y rotulos de una fila del molde. Sin formulas: esas se
    escriben despues, ya remapeadas."""
    for col in range(1, ultima_col + 1):
        fuente = molde.cell(origen, col)
        celda = ws.cell(destino, col)
        if isinstance(celda, MergedCell):
            continue
        if fuente.has_style:
            celda._style = copy(fuente._style)
        valor = fuente.value
        if isinstance(valor, ArrayFormula) or excel_ops.es_formula(valor):
            valor = None
        celda.value = valor
    dim = molde.row_dimensions.get(origen)
    if dim is not None and dim.height is not None:
        ws.row_dimensions[destino].height = dim.height
    if dim is not None and dim.hidden:
        ws.row_dimensions[destino].hidden = True
