"""Precios de mercado interno, del boletin de fruta de ODEPA.

Llena un hueco que no tenia fuente: la fila **«Precio mercado interno
(USD/kg)»** del flujo proyectado (§7, offset +3 de cada bloque). Hasta ahora se
tecleaba de memoria.

Se baja del mismo boletin mensual que el FOB, archivo **«Precios mayoristas y
consumidor»** (ver la cabecera de `importar_fob_odepa.py` para el detalle de
donde y por que se baja a mano).

    python scripts/importar_mayoristas_odepa.py "ruta/Precios-mayoristas...xlsx"

**La trampa, y es la misma clase que la del FOB:** el archivo trae CUATRO
puntos de venta y sus precios no se parecen. El arandano, 2026:

    Mercado Mayorista        2.926 $/kg
    Feria libre              3.225 $/kg
    Supermercado            17.666 $/kg   <- 6 veces el mayorista
    Supermercado en Linea   19.114 $/kg

**Para el flujo sirve el MERCADO MAYORISTA**, que es lo mas cercano a lo que
recibe el productor cuando vende al mercado interno. Los otros tres llevan
margen de retail y ponerlos en la fila del flujo inflaria el ingreso del
cliente varias veces. Se guardan los cuatro -sirven de contexto- pero solo el
mayorista queda marcado como utilizable.

**Ojo con la estacionalidad.** El precio es mensual y la fruta se paga muy
distinto segun el mes. Se guarda el promedio anual Y la serie mensual: el
promedio sirve de referencia gruesa, pero para una especie de ventana corta
-la cereza- lo que importa son sus meses de cosecha.
"""

import argparse
import io
import json
import os
import sys
from collections import defaultdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook                                # noqa: E402

from config.settings import ajustes                               # noqa: E402

ANIO, MES, PUNTO, PRODUCTO, UNIDAD, PRECIO = 0, 1, 2, 3, 4, 5

#: El unico punto de venta que se parece a lo que recibe el productor.
PUNTO_UTIL = "Mercado Mayorista"


def _clave(texto):
    import unicodedata
    sin = unicodedata.normalize("NFKD", str(texto or ""))
    sin = "".join(c for c in sin if not unicodedata.combining(c))
    return " ".join(sin.lower().split())


def leer(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    filas = libro.worksheets[0].iter_rows(values_only=True)
    encabezado = next(filas)

    # (producto, punto, anio) -> {mes: precio}
    series = defaultdict(dict)
    unidades = set()
    for f in filas:
        if not f or f[ANIO] is None or f[PRECIO] is None:
            continue
        unidades.add(str(f[UNIDAD] or "").strip())
        clave = (str(f[PRODUCTO]).strip(), str(f[PUNTO]).strip(), int(f[ANIO]))
        series[clave][int(f[MES])] = float(f[PRECIO])
    libro.close()
    return series, sorted(unidades), list(encabezado)


def armar(series, unidades, origen, encabezado):
    anios = sorted({a for (_p, _pv, a) in series})
    # El ultimo anio del archivo suele traer pocos meses: el boletin se publica
    # a mitad de anio. El promedio de tres meses de verano no es el precio del
    # anio, asi que la referencia sale del ultimo anio COMPLETO que haya.
    completos = sorted({a for (_p, _pv, a), meses in series.items()
                        if len(meses) >= 10})
    referencia = completos[-1] if completos else anios[-1]

    productos = {}
    for (producto, punto, anio), meses in series.items():
        clave = _clave(producto)
        entrada = productos.setdefault(clave, {
            "nombre": producto,
            "puntos_de_venta": {},
            "tipo": "precio mercado interno",
            "unidad": "CLP/kg",
            "fuente": "ODEPA - Boletin de la fruta (registros diarios de precios)",
        })
        valores = list(meses.values())
        registro = entrada["puntos_de_venta"].setdefault(punto, {})
        registro[str(anio)] = {
            "promedio_clp_kg": round(sum(valores) / len(valores), 2),
            "meses": len(valores),
            "por_mes": {str(m): round(v, 2) for m, v in sorted(meses.items())},
        }

    # El valor de referencia de cada producto: mayorista, ultimo anio completo.
    for clave, entrada in productos.items():
        mayorista = entrada["puntos_de_venta"].get(PUNTO_UTIL, {})
        elegido = mayorista.get(str(referencia))
        if elegido is None and mayorista:
            ultimo = sorted(mayorista)[-1]
            elegido = mayorista[ultimo]
            entrada["anio_de_referencia"] = ultimo
            entrada["confianza"] = "media"
            entrada["advertencia"] = (
                "No hay un anio completo de mercado mayorista: se usa {} con "
                "{} mes(es).".format(ultimo, elegido["meses"]))
        else:
            entrada["anio_de_referencia"] = str(referencia)
            entrada["confianza"] = "alta" if elegido else "baja"
        if elegido:
            entrada["promedio_clp_kg"] = elegido["promedio_clp_kg"]
            entrada["meses_del_promedio"] = elegido["meses"]
        else:
            entrada["promedio_clp_kg"] = None
            entrada["advertencia"] = (
                "Este producto no tiene precio de mercado mayorista: solo "
                "feria, supermercado o ambos, que llevan margen de retail y "
                "no sirven para la fila del flujo.")
            entrada["confianza"] = "baja"

    return {
        "version": "1.0",
        "generado": datetime.now().isoformat(timespec="seconds"),
        "origen": os.path.basename(origen),
        "encabezado_del_archivo": encabezado,
        "fuente_declarada": (
            "ODEPA, Boletin de la fruta. Precios capturados diariamente en "
            "mercados mayoristas, ferias libres y supermercados."),
        "tipo": "precio mercado interno",
        "unidades_del_archivo": unidades,
        "anios": [str(a) for a in anios],
        "anio_de_referencia": str(referencia),
        "punto_de_venta_utilizable": PUNTO_UTIL,
        "metodo": (
            "Promedio simple de los meses disponibles del anio de referencia, "
            "en el punto de venta «{}». Los otros puntos se guardan como "
            "contexto.".format(PUNTO_UTIL)),
        "advertencia": (
            "SOLO EL MERCADO MAYORISTA sirve para la fila de precio del flujo. "
            "Feria libre y supermercado llevan margen de retail: el arandano "
            "vale 2.926 $/kg mayorista y 17.666 en supermercado. Poner el de "
            "retail multiplicaria el ingreso del cliente."),
        "advertencia_estacional": (
            "El precio es mensual y la fruta se paga muy distinto segun el "
            "mes. Para una especie de ventana corta conviene mirar «por_mes» "
            "de sus meses de cosecha, no el promedio del anio."),
        "productos": dict(sorted(productos.items())),
    }


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("archivo", help="el «Precios mayoristas y consumidor» de ODEPA")
    p.add_argument("--salida", default=None)
    args = p.parse_args()

    if not os.path.exists(args.archivo):
        raise SystemExit("  No existe el archivo: {}".format(args.archivo))

    print("=" * 74)
    print("PRECIOS DE MERCADO INTERNO - ODEPA")
    print("=" * 74)
    series, unidades, encabezado = leer(args.archivo)
    datos = armar(series, unidades, args.archivo, encabezado)

    salida = args.salida or str(
        ajustes.raiz / "config" / "precios_mercado_interno.json")
    with io.open(salida, "w", encoding="utf-8") as fh:
        json.dump(datos, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    utiles = [e for e in datos["productos"].values() if e["promedio_clp_kg"]]
    print("  {} productos, {} con precio mayorista utilizable".format(
        len(datos["productos"]), len(utiles)))
    print("  anio de referencia: {}".format(datos["anio_de_referencia"]))
    print()
    print("  {:<26} {:>12} {:>8}".format("producto", "CLP/kg", "meses"))
    print("  " + "-" * 50)
    for e in sorted(utiles, key=lambda x: -x["promedio_clp_kg"])[:14]:
        print("  {:<26} {:>12,.0f} {:>8}".format(
            e["nombre"][:26], e["promedio_clp_kg"], e["meses_del_promedio"]))

    sin = [e["nombre"] for e in datos["productos"].values()
           if not e["promedio_clp_kg"]]
    if sin:
        print()
        print("  Sin precio mayorista ({}): {}".format(
            len(sin), ", ".join(sin[:8]) + ("…" if len(sin) > 8 else "")))

    print()
    print("  Escrito en {}".format(salida))
    return 0


if __name__ == "__main__":
    sys.exit(main())
