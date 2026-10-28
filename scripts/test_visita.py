"""Prueba de la pantalla de captura de la visita a terreno.

Lo que se comprueba aqui es la cadena completa de los datos de ORIGEN A: lo
que el tasador escribe en el formulario tiene que llegar al Excel del banco,
convertido a los modelos del dominio y sin que nadie lo transcriba dos veces.

Desde el 2026-10-14 la visita es la del formato final de Santander: una
temporada real en pesos (Anexo 1), ventas del F29 mes a mes, lotes y
unidades prediales, sensibilizacion por especie; sin maquinaria ni
infraestructura, que el formato ya no trae.

Uso: python scripts/test_visita.py
"""

import os
import sys
import tempfile
import warnings
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config.settings import ajustes as config  # noqa: E402
from services.visita import (  # noqa: E402
    aplicar_visita, cargar_formulario, resumen_de_avance,
)

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)


def formulario():
    return cargar_formulario(config.ruta_formulario_visita)


def informe_vacio(cliente="Prueba Visita"):
    from schemas.informe import Antecedentes, InformeFitogreen
    return InformeFitogreen(antecedentes=Antecedentes(cliente=cliente, rut="76.543.210-3"))


CAPTURA = {
    "temporadas_anexo": {"temporada_real": "2025/26"},
    "resultado_real": [
        {"especie": "Palto", "produccion_exportacion_kg": 200000,
         "produccion_mercado_interno_kg": 50000,
         "ventas_exportacion_pesos": 450_000_000,
         "ventas_mercado_interno_pesos": 60_000_000,
         "costos_directos_pesos": 300_000_000},
        {"especie": "Cerezo", "produccion_exportacion_kg": 80000,
         "ventas_exportacion_pesos": 240_000_000,
         "costos_directos_pesos": 120_000_000},
        {}, {},
    ],
    "gastos_indirectos": {"real_administracion_pesos": 30_000_000,
                          "real_gastos_generales_pesos": 20_000_000,
                          "administracion_usd_anio": 45000},
    "ventas_f29": {"f29_01": 10_000_000, "f29_06": 300_000_000, "f29_12": 5_000_000,
                   "nota_f29": "Agosto sin base imponible."},
    "lotes": [{"unidad_predial": "Fundo El Manzano", "nombre": "Hijuela 1", "rol": "123-45",
               "region": "VI", "comuna": "Requínoa", "superficie_legal_ha": 20,
               "superficie_real_ha": 21.5}, {}],
    "unidades_prediales": [
        {"nombre": "Fundo El Manzano", "i_ii_riego_ha": 12.5, "iii_iv_riego_ha": 8.0,
         "topografia": "Plana", "capacidad_tranques_m3": 4000},
        {}, {},
    ],
    "derechos_agua": [{"predio": "Fundo El Manzano", "tipo_fuente": "Subterránea",
                       "ejercicio": "Permanente y  continuo", "cantidad": 20,
                       "caudal_real_ls": 18}, {}],
    "sistemas_produccion": [{"especie": "Palto", "densidad_plantacion": "Alta",
                             "estado_fitosanitario": "Bueno"}],
    "sensibilizacion": [{"especie": "Palto", "precio_exportacion": -5, "costos": 10}, {}],
    "sistemas_comercializacion": [
        {"especie": "Palto", "empresas_comercializadoras": "Exportadora Andes",
         "mercado": "Exportación", "certificaciones": "GlobalG.A.P.",
         "periodo_cosecha": "Sep-Ene"}],
    "tipo_cambio": {"tc_ultima_temporada": 940, "tc_proyeccion": 950,
                    "uf_del_informe": 39500},
    "comentarios": {"notas_suelos": "Sin estudio agrológico reciente.",
                    "variedades_con_royalty": "Hass Plus, Lapins"},
    "opinion": {"gestion_texto": "Administracion profesionalizada.",
                "gestion_calificacion": "Muy Bueno"},
}


def probar_definicion():
    print("\n[1] El formulario se declara en datos, no en el codigo")
    f = formulario()
    pintable = f.para_navegador()
    ids = [s["id"] for s in pintable["secciones"]]
    comprobar(len(ids) >= 12, "trae las secciones del formato final", str(ids))
    comprobar(not {"maquinaria", "infraestructura", "resultado_preliminar"} & set(ids),
              "sin las secciones que el formato final ya no tiene")
    comprobar({"ventas_f29", "lotes", "unidades_prediales", "sensibilizacion"} <= set(ids),
              "con las que agrega: F29, lotes, unidades prediales, sensibilizacion")

    destinos = [s.get("destino") for s in f.secciones if s.get("tipo") == "tabla"]
    comprobar(all(destinos), "cada tabla declara a que parte del informe va", str(destinos))

    con_opciones = [c for s in pintable["secciones"] for c in s["campos"] if c.get("opciones")]
    comprobar(len(con_opciones) >= 15,
              "las escalas llegan expandidas al navegador, que no las conoce",
              str(len(con_opciones)))
    fitosanitario = next(c["opciones"] for c in con_opciones
                         if c["etiqueta"] == "Estado fitosanitario")
    comprobar(fitosanitario == ["Bueno", "Regular", "Malo"],
              "y son las listas del formato, no unas inventadas", str(fitosanitario))
    calificacion = next(c["opciones"] for c in con_opciones if c["etiqueta"] == "Calificación")
    comprobar(calificacion == ["Muy Bueno", "Bueno", "Regular", "Malo"],
              "la calificacion es la escala de cuatro del formato final", str(calificacion))
    comprobar(all("escala" not in c for s in pintable["secciones"] for c in s["campos"]),
              "no se filtra el nombre interno de la escala al navegador")


def probar_traduccion():
    print("\n[2] De lo capturado al modelo del dominio")
    informe = informe_vacio()
    avisos = aplicar_visita(informe, CAPTURA, formulario())

    comprobar(not avisos, "ninguna ruta del formulario apunta a un lugar inexistente",
              str(avisos))
    comprobar(informe.anexo.resultado_real.temporada == "2025/26",
              "la temporada real")
    comprobar(len(informe.anexo.resultado_real.filas) == 2,
              "las dos filas con datos entran; las dos vacias no")
    comprobar(informe.anexo.resultado_real.administracion_pesos == 30_000_000
              and informe.supuestos.administracion_usd_anio == 45000,
              "la administracion real (pesos) y la proyectada (USD/año)")
    f29 = informe.anexo.ventas_f29_pesos
    comprobar(len(f29) == 12 and f29[0] == 10_000_000 and f29[5] == 300_000_000
              and f29[11] == 5_000_000 and f29[1] is None,
              "las ventas del F29 caen en su mes, y un mes vacio queda vacio", str(f29))
    lotes = informe.dotacion_recursos.lotes
    comprobar(len(lotes) == 1 and lotes[0].rol == "123-45" and lotes[0].superficie_real_ha == 21.5,
              "los lotes con sus datos legales y la superficie real")
    unidades = informe.dotacion_recursos.unidades_prediales
    comprobar(len(unidades) == 1 and unidades[0].i_ii_riego_ha == 12.5
              and unidades[0].capacidad_tranques_m3 == 4000,
              "la unidad predial con su clasificacion y sus tranques")
    comprobar(informe.dotacion_recursos.derechos_agua[0].ejercicio == "Permanente y  continuo",
              "el ejercicio del derecho de agua")
    sens = informe.supuestos.sensibilizacion.get("Palto")
    comprobar(sens is not None and sens.precio_exportacion == -0.05 and sens.costos == 0.1,
              "la sensibilizacion se pide en % y se guarda como fraccion", str(sens))
    comprobar(informe.flujo_proyectado.tipo_cambio_pesos_usd.get("proyeccion") == 950,
              "el tipo de cambio de la proyeccion")
    comprobar(informe.estructura_productiva.sistemas_produccion[0].estado_fitosanitario == "Bueno",
              "el estado fitosanitario")
    comprobar(informe.estructura_productiva.sistemas_comercializacion[0].certificaciones
              == "GlobalG.A.P.", "las certificaciones")
    comprobar(informe.dotacion_recursos.notas_suelos.startswith("Sin estudio"),
              "los comentarios de cada hoja")
    comprobar(informe.opinion.gestion_administrativa_productiva.calificacion is not None,
              "la calificacion de la opinion")


def probar_llenado_parcial():
    print("\n[3] Se guarda a medias, que es como se trabaja de verdad")
    informe = informe_vacio()
    aplicar_visita(informe, {"lotes": [{"unidad_predial": "Fundo"}]}, formulario())
    comprobar(len(informe.dotacion_recursos.lotes) == 1,
              "una seccion sola se aplica sin exigir las demas")
    comprobar(informe.dotacion_recursos.lotes[0].rol is None,
              "y un campo sin llenar queda en None, no en cero")

    con_borrador = informe_vacio()
    con_borrador.opinion.conclusion_final.texto = "Borrador redactado por la IA."
    aplicar_visita(con_borrador, {"opinion": {"gestion_texto": "Lo mio."}}, formulario())
    comprobar(con_borrador.opinion.conclusion_final.texto == "Borrador redactado por la IA.",
              "lo que el tasador no toco se queda como estaba")
    comprobar(con_borrador.opinion.gestion_administrativa_productiva.texto == "Lo mio.",
              "y lo que si escribio, manda")
    comprobar(not aplicar_visita(informe_vacio(), {}, formulario()),
              "una visita sin nada no produce avisos ni escribe nada")


def probar_avance():
    print("\n[4] El avance por seccion, para saber que falta")
    avance = {a["id"]: a for a in resumen_de_avance(CAPTURA, formulario())}
    comprobar(avance["resultado_real"]["llenos"] == 2,
              "cuenta las filas con datos en resultados", str(avance["resultado_real"]["llenos"]))
    comprobar(avance["inversiones"]["vacio"], "marca como vacia la seccion que no se toco")
    comprobar(avance["gastos_indirectos"]["llenos"] == 3,
              "y en las secciones de campos cuenta campos, no filas",
              str(avance["gastos_indirectos"]["llenos"]))


def probar_ajustes_de_terreno():
    print("\n[4b] Ajustes decididos en terreno")
    from services.visita import ajustes_por_curva

    informe = informe_vacio()
    avisos = aplicar_visita(informe, {"ajustes_de_curva": [
        {"especie": "Palto", "ajuste_pct": -2, "motivo": "riego mas o menos, presion despareja"},
        {"especie": "Cerezo", "ajuste_pct": 5, "motivo": "huerto muy bien manejado"},
        {}, {},
    ]}, formulario())
    comprobar(not avisos, "el ajuste de terreno entra sin avisos", str(avisos))
    por_curva = ajustes_por_curva(informe)
    comprobar(set(por_curva) == {"Palto", "Cerezo"} and por_curva["Palto"][0].factor_pct == -2,
              "se indexan por la especie que escribio el tasador, con su porcentaje")
    solo_cero = informe_vacio()
    aplicar_visita(solo_cero, {"ajustes_de_curva": [{"especie": "Palto", "ajuste_pct": 0}]},
                   formulario())
    comprobar(ajustes_por_curva(solo_cero) == {},
              "un 0% no genera ajuste: no cambia nada y ensuciaria el informe")


def probar_hasta_el_excel():
    print("\n[5] Lo capturado llega al Excel del formato final de Santander")
    from openpyxl import load_workbook

    from schemas.informe import Plantacion
    from services.excel_engine_santander import MotorSantander

    plantilla, mapping = config.rutas_de_plantilla("santander")
    informe = informe_vacio("Demo Captura")
    informe.estructura_productiva.plantaciones = [
        Plantacion(especie="Palto", variedad="Hass Plus", anio_plantacion="2015",
                   superficie_ha=10, predio="Fundo El Manzano"),
        Plantacion(especie="Cerezo", variedad="Regina", anio_plantacion="2018",
                   superficie_ha=5, predio="Fundo El Manzano"),
    ]
    aplicar_visita(informe, CAPTURA, formulario())
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "test_visita.xlsx"
        resultado = MotorSantander(plantilla, mapping).generar(informe, destino)
        # Sin fotos el motor avisa que la portada sale sin emplazamiento: es
        # lo esperado aca, no un problema de la visita.
        otros = [a for a in resultado.avisos if "emplazamiento" not in a]
        comprobar(not otros, "el motor escribe la visita sin avisos", str(otros))
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            libro = load_workbook(destino)
    a1 = libro["Anexo 1 Resultado"]
    comprobar(a1["A5"].value == "Palto" and a1["E5"].value == 450_000_000,
              "la especie declarada y sus ventas en el Anexo 1", str(a1["A5"].value))
    comprobar(a1["B5"].value == 250000 and a1["D5"].value == 0.8,
              "la produccion total y el % exportado salen de los kilos declarados")
    comprobar(str(a1["G5"].value).startswith("="), "sobre esos datos, las formulas siguen vivas")
    a2 = libro["Anexo 2 Ventas F29"]
    comprobar(a2["B5"].value == 10_000_000 and a2["B10"].value == 300_000_000
              and a2["B6"].value is None, "las ventas del F29 mes a mes en el Anexo 2")
    rec = libro["Recursos"]
    comprobar(rec["E8"].value == "123-45" and rec["F8"].value == "VI" and rec["M8"].value == 21.5,
              "el lote en DATOS LEGALES TERRENOS")
    unidad = next(c.row for c in rec["A"] if c.value == "Fundo El Manzano" and c.row > 10)
    comprobar(rec.cell(unidad, 2).value == 12.5 and rec.cell(unidad, 11).value == "Plana",
              "la unidad predial en la clasificacion de suelos")
    plant = libro["Plantaciones"]
    comprobar(plant["D7"].value == "Sí" and plant["D8"].value == "No",
              "el royalty sale de la lista de variedades con royalty")
    sup = libro["Supuestos"]
    comprobar(sup["B4"].value == "Palto" and sup["B5"].value == -0.05 and sup["B8"].value == 0.1
              and sup["B10"].value == 45000,
              "la sensibilizacion y la administracion en Supuestos")
    comprobar(libro["Portada"]["H14"].value == 39500, "la UF en la portada")
    opinion = libro["Comentarios generales Opinion"]
    comprobar(opinion["H4"].value == "Muy Bueno", "la calificacion en la casilla del banco")


def main():
    print("=" * 72)
    print("PRUEBA DE LA CAPTURA DE VISITA A TERRENO")
    print("=" * 72)
    probar_definicion()
    probar_traduccion()
    probar_llenado_parcial()
    probar_avance()
    probar_ajustes_de_terreno()
    probar_hasta_el_excel()
    print("\n" + "=" * 72)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
