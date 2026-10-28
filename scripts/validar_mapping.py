"""Valida cada mapping registrado en config/plantillas.json contra su
plantilla y su esquema Pydantic.

Comprueba tres cosas:
  1. que cada hoja referenciada exista en la plantilla,
  2. que cada ruta de atributo exista en el esquema del formato (ver
     MODELO_POR_ALIAS -Rabobank y Santander NO comparten dominio),
  3. que ninguna celda de destino pise una formula o una etiqueta de la plantilla.
"""

import json
import os
import re
import sys
from typing import get_args, get_origin

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from schemas.informe import InformeFitogreen  # noqa: E402
from schemas.retasacion import RetasacionRabobank  # noqa: E402

from config.settings import ajustes  # noqa: E402

# Sin argumentos valida todos los formatos registrados; con un alias, solo ese.
ALIAS = sys.argv[1] if len(sys.argv) > 1 else None
MAPPING = None
PLANTILLA = None

#: Cada formato tiene su propio dominio -Rabobank mide garantia (hipoteca/
#: prenda), no capacidad de pago, asi que NO comparte esquema con Santander
#: ("no mezclar los motores"). InformeFitogreen es el default: sirve tanto
#: para 'santander' como para 'fitogreen' (catastro), que reusa sus
#: submodelos.
MODELO_POR_ALIAS = {
    "rabobank": RetasacionRabobank,
}
MODELO = InformeFitogreen

errores = []
avisos = []


def resolver_atributo(ruta):
    """Recorre 'a.b.c' por los modelos anidados. Devuelve True si existe."""
    modelo = MODELO
    for parte in ruta.split("."):
        if modelo is None:
            return False
        campos = getattr(modelo, "model_fields", None)
        if campos is None or parte not in campos:
            return False
        ann = campos[parte].annotation
        # desenvuelve Optional[...] y List[...]
        for _ in range(3):
            args = [a for a in get_args(ann) if a is not type(None)]
            if get_origin(ann) is not None and args:
                ann = args[0] if get_origin(ann) is not dict else None
            else:
                break
        modelo = ann if hasattr(ann, "model_fields") else None
    return True


def validar_santander2026(m, wb, hojas):
    """El mapping del formato final de Santander no declara tablas fijas sino
    partes que se repiten (services/disposicion.py): se verifica que cada
    parte caiga dentro de su hoja, que cada celda suelta exista y no pise una
    formula del molde, que cada ruta exista en el esquema, y que los huecos de
    imagen y los cuadros de texto existan con su nombre."""
    from services.dibujos_xlsx import nombres_de_formas

    for parte in m["partes"]:
        hoja = hojas.get(parte["hoja"])
        if hoja not in wb.sheetnames:
            errores.append("parte '{}': hoja inexistente {!r}".format(parte["id"], parte["hoja"]))
            continue
        desde, hasta = parte["filas"]
        if not (1 <= desde <= hasta <= wb[hoja].max_row):
            errores.append("parte '{}': filas {}-{} fuera de la hoja {!r}".format(
                parte["id"], desde, hasta, hoja))
        for letra in parte.get("columnas", {}):
            if not re.match(r"^[A-Z]{1,3}$", letra):
                errores.append("parte '{}': columna invalida {!r}".format(parte["id"], letra))
        for celda in list(parte.get("celdas", {})) + [c for c in parte.get("pisa_formula", []) if re.search(r"\d", c)]:
            fila = int(re.search(r"\d+", celda).group())
            if not desde <= fila <= hasta:
                errores.append("parte '{}': la celda {} no esta en sus filas".format(parte["id"], celda))
    print("partes verificadas: {}".format(len(m["partes"])))

    for referencia, ruta in m.get("campos", {}).items():
        alias, celda = referencia.split("!")
        ws = wb[hojas[alias]]
        valor = ws[celda].value
        if isinstance(valor, str) and valor.startswith("="):
            errores.append("{} apunta a {} que contiene una FORMULA".format(ruta, referencia))
        elif valor not in (None, ""):
            avisos.append("{} apunta a {} que ya tiene contenido: {!r}".format(ruta, referencia, valor))
        if not ruta.startswith("@") and not resolver_atributo(ruta):
            errores.append("atributo inexistente en el schema: {}".format(ruta))
    for referencia, ruta in m.get("listas", {}).items():
        if not resolver_atributo(ruta):
            errores.append("atributo inexistente en el schema: {}".format(ruta))
    print("campos verificados: {}".format(len(m.get("campos", {}))))

    formas = nombres_de_formas(PLANTILLA)
    n = 0
    imagenes = m.get("imagenes", {})
    pedidos = []
    if "emplazamiento" in imagenes:
        pedidos.append((imagenes["emplazamiento"]["hoja"], imagenes["emplazamiento"]["forma"]))
    if "galeria" in imagenes:
        pedidos += [(imagenes["galeria"]["hoja"], f) for f in imagenes["galeria"]["formas"]]
    pedidos += [tuple(k.split("!")) for k in m.get("cuadros_de_texto", {})]
    for alias, forma in pedidos:
        n += 1
        if forma not in [x for _, x in formas.get(hojas[alias], [])]:
            errores.append("la forma '{}' no existe en la hoja {!r}".format(forma, hojas[alias]))
    print("formas de dibujo verificadas: {}".format(n))


def main():
    with open(str(MAPPING), encoding="utf-8") as fh:
        m = json.load(fh)
    wb = load_workbook(PLANTILLA)

    hojas = m["hojas"]
    for alias, nombre in hojas.items():
        if nombre not in wb.sheetnames:
            errores.append("hoja inexistente: alias '{}' -> '{}'".format(alias, nombre))
    print("hojas verificadas: {}/{}".format(len(hojas) - len(errores), len(hojas)))
    if "partes" in m:
        validar_santander2026(m, wb, hojas)
        return _cerrar()

    # --- campos simples
    n_campos = 0
    for ruta, spec in m["campos"].items():
        n_campos += 1
        ws = wb[hojas[spec["hoja"]]]
        celda = ws[spec["celda"]]
        if isinstance(celda.value, str) and celda.value.startswith("="):
            errores.append("{} apunta a {}!{} que contiene una FORMULA"
                           .format(ruta, ws.title, spec["celda"]))
        elif celda.value is not None:
            avisos.append("{} apunta a {}!{} que ya tiene contenido: {!r}"
                          .format(ruta, ws.title, spec["celda"], celda.value))
        # la raiz 'flujo_proyectado.tipo_cambio...' usa sufijos libres
        base = ruta.rsplit(".", 1)[0] if ruta.endswith((".historico", ".proyeccion")) else ruta
        if not resolver_atributo(base):
            errores.append("atributo inexistente en el schema: {}".format(base))
    print("campos simples verificados: {}".format(n_campos))

    # --- tablas
    n_tablas = 0
    for nombre, t in m["tablas"].items():
        n_tablas += 1
        ws = wb[hojas[t["hoja"]]]
        if "atributo" in t and not resolver_atributo(t["atributo"]):
            errores.append("tabla '{}': atributo inexistente {}".format(nombre, t["atributo"]))
        cols = t.get("columnas", {})
        if isinstance(cols, dict):
            for campo, col in cols.items():
                if not re.match(r"^[A-Z]{1,3}$", col):
                    errores.append("tabla '{}': columna invalida {!r}".format(nombre, col))
                    continue
                if "primera_fila" in t:
                    c = ws.cell(t["primera_fila"], column_index_from_string(col))
                    if isinstance(c.value, str) and c.value.startswith("="):
                        errores.append(
                            "tabla '{}': {} -> {}!{}{} es FORMULA".format(
                                nombre, campo, ws.title, col, t["primera_fila"]))
        # el encabezado declarado debe tener texto
        fe = t.get("fila_encabezado")
        for fila in (fe if isinstance(fe, list) else [fe] if fe else []):
            vals = [ws.cell(fila, c).value for c in range(1, 15)]
            if not any(isinstance(v, str) and v.strip() for v in vals):
                avisos.append("tabla '{}': fila_encabezado {} sin texto en {}"
                              .format(nombre, fila, ws.title))
    print("tablas verificadas: {}".format(n_tablas))

    # --- fotografias (no todos los formatos reservan hueco para imagenes)
    huecos = m.get("fotografias", {}).get("soportado_por_plantilla", {})
    for nombre, foto in huecos.items():
        ws = wb[hojas[foto["hoja"]]]
        if foto.get("titulo_en") and ws[foto["titulo_en"]].value != foto.get("titulo_texto"):
            errores.append("el titulo de imagen '{}' no coincide: {}!{} = {!r}".format(
                nombre, ws.title, foto["titulo_en"], ws[foto["titulo_en"]].value))
        if ws[foto["celda_ancla"]].value is not None:
            errores.append("la celda ancla de foto '{}' {}!{} no esta vacia".format(
                nombre, ws.title, foto["celda_ancla"]))
    print("huecos de fotografia verificados: {}".format(len(huecos)))

    # --- dibujos (Rabobank): cada cuadro de texto y cada hueco de imagen que
    # declara el mapping tiene que existir en la plantilla con ese nombre. Si
    # el banco manda una plantilla nueva y una forma cambia de nombre, el
    # motor solo avisaria al generar; aqui se ve antes.
    from services.dibujos_xlsx import nombres_de_formas
    formas = None
    n_formas = 0
    for seccion in ("cuadros_de_texto", "huecos_de_imagen"):
        for clave, spec in (m.get(seccion) or {}).items():
            if not isinstance(spec, dict):
                continue
            if formas is None:
                formas = nombres_de_formas(PLANTILLA)
            n_formas += 1
            nombre_hoja = hojas.get(spec["hoja"])
            if spec["forma"] not in [n for _, n in formas.get(nombre_hoja, [])]:
                errores.append("{} '{}': la forma '{}' no existe en la hoja {!r} de la "
                               "plantilla".format(seccion, clave, spec["forma"], nombre_hoja))
    print("formas de dibujo verificadas: {}".format(n_formas))
    return _cerrar()


def _cerrar():
    print()
    for a in avisos:
        print("AVISO  ", a)
    for e in errores:
        print("ERROR  ", e)
    print()
    print("{} error(es), {} aviso(s)".format(len(errores), len(avisos)))
    return 1 if errores else 0


def validar_todos():
    """Valida cada formato registrado en config/plantillas.json."""
    global MAPPING, PLANTILLA, MODELO, errores, avisos

    alias_a_validar = [ALIAS] if ALIAS else sorted(ajustes.registro_plantillas)
    total_errores = 0

    for alias in alias_a_validar:
        PLANTILLA, MAPPING = ajustes.rutas_de_plantilla(alias)
        MODELO = MODELO_POR_ALIAS.get(alias, InformeFitogreen)
        errores.clear()
        avisos.clear()
        print("=" * 70)
        print("FORMATO '{}'  ->  {}".format(alias, PLANTILLA.name))
        print("=" * 70)
        total_errores += main()
        print()

    return 1 if total_errores else 0


if __name__ == "__main__":
    sys.exit(validar_todos())
