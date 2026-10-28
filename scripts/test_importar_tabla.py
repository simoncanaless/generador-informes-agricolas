"""Llenar una seccion de la visita desde la planilla del cliente.

Lo que se comprueba no es solo que lea: es que reconozca una planilla ESCRITA
POR OTRO -abreviaturas, titulo arriba, varias hojas, fila de totales- y que lo
importado quede marcado como sin revisar.

Uso: python scripts/test_importar_tabla.py
"""

import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook                                    # noqa: E402

from config.settings import ajustes                              # noqa: E402
from services.importar_tabla import (                            # noqa: E402
    TablaIlegible, claves_sin_revisar, importar_seccion,
)

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)


def planilla_de_cliente(ruta):
    """Como llega de verdad: titulo, abreviaturas, varias hojas y un total."""
    libro = Workbook()
    portada = libro.active
    portada.title = "Portada"
    portada["A1"] = "AGRICOLA EL PEUMO LTDA."

    ws = libro.create_sheet("Predios")
    ws["B2"] = "1. TERRENOS"
    ws["B4"], ws["C4"], ws["D4"] = "Nombre del Predio", "Rol (es)", "Superficie Física"
    ws["B5"], ws["C5"], ws["D5"] = "Fundo El Peumo", "123-45", 20
    ws["B6"], ws["C6"], ws["D6"] = "Hijuela Norte", "123-46", 12.5
    ws["B7"], ws["C7"], ws["D7"] = "Parcela 7", "123-47", 8
    ws["B8"], ws["D8"] = "TOTAL", 40.5

    ws2 = libro.create_sheet("Resultados")
    ws2["A1"] = "Temporada 2024/25"
    for columna, rotulo in zip("ABCDEFGH", (
            "Cultivo", "Sup. (ha)", "Kg exportacion", "Kg mercado interno",
            "Ventas exportacion", "Ventas mercado interno", "Costos directos",
            "Comentarios")):
        ws2["{}3".format(columna)] = rotulo
    datos = [
        ("Palto Hass", 69.5, 765000, 92000, 1_680_000_000, 74_000_000,
         690_000_000, "buen anio"),
        ("Cerezo", 12.0, 84000, 9000, 410_000_000, 11_000_000, 180_000_000, ""),
        ("Nogal", 43.0, 210000, 31000, 305_000_000, 22_000_000, 160_000_000, ""),
    ]
    for numero, valores in enumerate(datos, start=4):
        for columna, valor in zip("ABCDEFGH", valores):
            ws2["{}{}".format(columna, numero)] = valor

    libro.save(ruta)


def main():
    print("=" * 72)
    print("IMPORTAR UNA SECCION DE LA VISITA DESDE LA PLANILLA DEL CLIENTE")
    print("=" * 72)

    formulario = json.load(
        io.open(ajustes.ruta_formulario_visita, encoding="utf-8"))
    secciones = {s["id"]: s for s in formulario["secciones"]}

    carpeta = ajustes.ruta_salida / "pruebas"
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / "planilla_cliente.xlsx"
    planilla_de_cliente(ruta)

    print("\n[1] Todas las tablas declaran sinonimos")
    tablas = [s for s in formulario["secciones"] if s.get("tipo") == "tabla"]
    sin = [s["id"] for s in tablas
           if not any(c.get("sinonimos") for c in s["campos"])]
    comprobar(not sin, "las {} secciones tabulares se pueden importar".format(
        len(tablas)), ", ".join(sin))

    print("\n[2] Lotes: elige la hoja y reconoce los encabezados del cliente")
    r = importar_seccion(ruta, secciones["lotes"])
    comprobar(r.hoja == "Predios", "elige la hoja correcta entre tres", r.hoja)
    comprobar(r.fila_encabezado == 4,
              "encuentra el encabezado en la fila 4, no en la 1",
              str(r.fila_encabezado))
    reconocidas = {c.campo: c.encabezado_en_el_archivo for c in r.reconocidas}
    comprobar(reconocidas.get("unidad_predial") == "Nombre del Predio",
              "«Nombre del Predio» -> Unidad predial")
    comprobar(reconocidas.get("rol") == "Rol (es)", "«Rol (es)» -> Rol")
    comprobar(reconocidas.get("superficie_legal_ha") == "Superficie Física",
              "«Superficie Física» -> Superficie legal")
    comprobar(r.cuantas == 3, "lee 3 filas y descarta la de TOTAL",
              str(r.cuantas))
    comprobar(r.filas[1]["superficie_legal_ha"] == 12.5, "los números llegan como números",
              str(r.filas[1].get("superficie_legal_ha")))

    print("\n[3] Resultados: 6 columnas, y avisa de las que no usó")
    r2 = importar_seccion(ruta, secciones["resultado_real"])
    comprobar(len(r2.reconocidas) == 6, "reconoce las 6 columnas del Anexo 1",
              str(len(r2.reconocidas)))
    comprobar(r2.filas[0]["ventas_exportacion_pesos"] == 1_680_000_000,
              "y las ventas sin perder ceros")
    comprobar(any("Comentarios" in a for a in r2.avisos),
              "avisa de la columna que no supo dónde poner",
              " | ".join(r2.avisos))

    print("\n[4] Lo importado queda marcado")
    claves = claves_sin_revisar("lotes", r.filas)
    comprobar(len(claves) == 9, "9 claves para 3 filas de 3 campos",
              str(len(claves)))
    comprobar("lotes|0|rol" in claves,
              "con el mismo formato que usa la visita para fusionar")

    print("\n[5] Una planilla que no corresponde se rechaza")
    otro = carpeta / "planilla_ajena.xlsx"
    libro = Workbook()
    libro.active["A1"] = "Lista de compras"
    libro.active["A2"] = "Pan"
    libro.save(otro)
    try:
        importar_seccion(otro, secciones["lotes"])
        comprobar(False, "rechaza una planilla que no tiene nada que ver")
    except TablaIlegible as error:
        comprobar("no se reconoció" in str(error),
                  "rechaza una planilla que no tiene nada que ver, y lo explica",
                  str(error)[:60])

    print("\n[6] Una sección de campos sueltos no se puede importar")
    try:
        importar_seccion(ruta, secciones["tipo_cambio"])
        comprobar(False, "se niega con una sección que no es tabla")
    except TablaIlegible:
        comprobar(True, "se niega con una sección que no es tabla")

    print("\n" + "=" * 72)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
