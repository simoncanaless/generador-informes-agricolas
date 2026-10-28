"""Extrae los precios FOB de exportacion del boletin de fruta de ODEPA.

Cierra el dato que §10.2 tenia pendiente: *"de donde saca el FOB"*. La
respuesta es el **boletin de la fruta** de ODEPA, que publica todos los meses
los datos crudos del Servicio Nacional de Aduanas.

**De donde se baja.** La pagina interactiva
(https://apps.odepa.gob.cl/powerBI/boletin_fruta.html) es un Power BI embebido
y **no** es la via: raspar su API interna es fragil y no esta documentada. Cada
boletin mensual publica los mismos datos como Excel:

    https://www.odepa.gob.cl/publicaciones/boletines
      -> el boletin del mes -> «Exportaciones por Temporada»

Ese archivo (`Exportaciones-por-temporadaMMAAAA.xlsx`, ~4,5 MB) es la entrada
de este script. El enlace lleva un identificador distinto cada mes, asi que se
baja a mano y se le pasa la ruta: es una vez al mes como mucho, y un
identificador fijo escrito en el codigo apuntaria al boletin de un mes viejo
para siempre sin que nadie lo note.

    python scripts/importar_fob_odepa.py "ruta/Exportaciones-por-temporada042026.xlsx"

**Dos trampas del dato, y las dos cambian el numero:**

1. **El FOB va por CODIGO ARANCELARIO, no por grupo.** El grupo «Avellana»
   mezcla dos productos con precios que difieren 2,4 veces: con cascara
   (08022100) a 4,17 USD/kg y sin cascara (08022200) a 9,99, temporada
   2024/25. El promedio del grupo -9,13- no describe a ninguno de los dos, y
   cual corresponde depende de que vende el cliente: el productor cosecha CON
   cascara; el descascarado es proceso industrial. Por eso se guarda el
   desglose completo y el grupo queda marcado como mezcla cuando lo es.
2. **La ultima temporada del archivo esta INCOMPLETA.** El archivo trae, a la
   vez, la temporada cerrada («Sep 2024 / Ago 2025») y la en curso («Sep 2025 /
   Mar 2026»). Tomar la parcial sesga el precio hacia la fruta de principio de
   temporada: en el avellano da 11,60 contra 9,99. Solo las cerradas -las que
   terminan en agosto- valen como referencia; las parciales se guardan aparte
   y rotuladas.

**Esto NO entra directo al flujo.** Un FOB no es lo que recibe el productor:
falta restarle la cadena de packing, frio, flete y comision (§10.2). El
archivo se marca entero como `FOB exportacion`, asi que
`PrecioReferencia.listo_para_el_flujo` da False y el motor no lo puede usar
por error.
"""

import argparse
import io
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook                                # noqa: E402

from config.settings import ajustes                               # noqa: E402

#: Columnas del archivo de ODEPA, por posicion. Se leen por posicion y no por
#: nombre a proposito: el encabezado cambio de «Grupo de producto» a «Grupo de
#: productos» entre versiones del boletin, y un KeyError por eso es una caida
#: tonta. El orden lleva estable desde 2016.
TEMPORADA, PAIS, GRUPO, CODIGO, PRODUCTO = 0, 1, 2, 3, 4
VOLUMEN, VALOR_FOB, PROCESAMIENTO = 5, 6, 7

#: Una temporada cerrada va de septiembre a agosto. Las que terminan en otro
#: mes son la temporada en curso, cortada donde alcanzo el boletin.
TEMPORADA_CERRADA = re.compile(r"/\s*Ago\s", re.IGNORECASE)

#: Debajo de esto el precio no es representativo: son embarques sueltos, y una
#: sola operacion rara mueve el promedio de un producto entero.
MINIMO_TONELADAS = 50.0

#: Lo que le interesa a una tasacion agricola. Las conservas, los jugos y los
#: aceites son otro negocio -y otro precio- que el predio no vende.
PROCESAMIENTOS = ("Fresca", "Frutos secos")


def _clave(texto):
    """Minusculas y sin tildes, igual que en el resto del proyecto."""
    import unicodedata
    sin = unicodedata.normalize("NFKD", str(texto or ""))
    sin = "".join(c for c in sin if not unicodedata.combining(c))
    return " ".join(sin.lower().split())


def _numero(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _orden_de_temporada(texto):
    """Para ordenar «Sep 2016 / Ago 2017» cronologicamente."""
    anios = re.findall(r"\d{4}", texto or "")
    return (anios[0] if anios else "0000", texto or "")


def leer(ruta):
    """Suma volumen y valor por temporada, producto y procesamiento."""
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hoja = libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas)

    por_producto = defaultdict(lambda: [0.0, 0.0])
    por_grupo = defaultdict(lambda: [0.0, 0.0])
    temporadas = set()
    leidas = 0

    for fila in filas:
        if not fila or fila[TEMPORADA] is None:
            continue
        procesamiento = str(fila[PROCESAMIENTO] or "").strip()
        if procesamiento not in PROCESAMIENTOS:
            continue
        temporada = str(fila[TEMPORADA]).strip()
        grupo = str(fila[GRUPO] or "").strip()
        codigo = str(fila[CODIGO] or "").strip()
        producto = str(fila[PRODUCTO] or "").strip()
        volumen, valor = _numero(fila[VOLUMEN]), _numero(fila[VALOR_FOB])
        if volumen <= 0:
            continue

        temporadas.add(temporada)
        leidas += 1
        clave = (temporada, grupo, procesamiento, codigo, producto)
        por_producto[clave][0] += volumen
        por_producto[clave][1] += valor
        por_grupo[(temporada, grupo, procesamiento)][0] += volumen
        por_grupo[(temporada, grupo, procesamiento)][1] += valor

    libro.close()
    return por_producto, por_grupo, sorted(temporadas, key=_orden_de_temporada), \
        leidas, list(encabezado)


def _usd_kg(volumen_ton, valor_usd):
    """El volumen viene en toneladas y el valor en dolares."""
    return round(valor_usd / (volumen_ton * 1000.0), 4) if volumen_ton else None


def armar(por_producto, por_grupo, temporadas, origen, encabezado):
    cerradas = [t for t in temporadas if TEMPORADA_CERRADA.search(t)]
    parciales = [t for t in temporadas if t not in cerradas]
    if not cerradas:
        raise SystemExit(
            "  El archivo no trae ninguna temporada cerrada (Sep a Ago).\n"
            "  Sin una temporada completa no hay precio de referencia.")
    ultima = cerradas[-1]

    grupos = {}
    for (temporada, grupo, procesamiento), (volumen, valor) in por_grupo.items():
        if temporada != ultima or volumen < MINIMO_TONELADAS:
            continue
        # El desglose por codigo arancelario del MISMO grupo y temporada.
        codigos = sorted(
            ({"codigo": c, "producto": p,
              "toneladas": round(v[0], 2), "usd_kg": _usd_kg(v[0], v[1])}
             for (t, g, pr, c, p), v in por_producto.items()
             if t == temporada and g == grupo and pr == procesamiento
             and v[0] >= MINIMO_TONELADAS),
            key=lambda x: -x["toneladas"])

        precios = [c["usd_kg"] for c in codigos if c["usd_kg"]]
        # «Mezcla» cuando los codigos del grupo difieren tanto que su promedio
        # no describe a ninguno. El caso del avellano: 4,17 y 9,99.
        es_mezcla = bool(precios) and max(precios) > 1.5 * min(precios)

        # Sin tildes: la busqueda normaliza, y una clave con tilde queda
        # inalcanzable. Paso con «arandanos y mirtilos», «limon», «castana» y
        # «nectarin»: cuatro grupos que existian y nadie podia encontrar.
        clave = _clave(grupo)
        grupos[clave] = {
            "nombre": grupo,
            "procesamiento": procesamiento,
            "temporada": temporada,
            "toneladas": round(volumen, 2),
            "usd_kg": _usd_kg(volumen, valor),
            "codigos": codigos,
            "es_mezcla": es_mezcla,
            "tipo": "FOB exportacion",
            "fuente": "ODEPA - Boletin de la fruta (Servicio Nacional de Aduanas)",
            "confianza": "media" if es_mezcla else "alta",
        }
        if es_mezcla:
            grupos[clave]["advertencia"] = (
                "El grupo mezcla {} productos con precios de {} a {} USD/kg. "
                "El promedio no describe a ninguno: elija el codigo que "
                "corresponde a lo que vende el cliente.".format(
                    len(codigos), min(precios), max(precios)))

    # La serie historica por grupo, para poder mostrar la tendencia.
    #
    # `_clave` y no `strip().lower()`: la clave del grupo se normaliza SIN
    # tildes, asi que con la tilde puesta el `if clave in grupos` de abajo
    # fallaba y cuatro grupos -arandano, castana, limon y nectarin- quedaban
    # con la serie VACIA. No se veia como un error: simplemente no salia el
    # rango de precios de esas especies. Es la tercera copia de esta misma
    # regla que muerde (§16.a.quadragies.quinquies).
    series = defaultdict(dict)
    for (temporada, grupo, _proc), (volumen, valor) in por_grupo.items():
        if temporada in cerradas and volumen >= MINIMO_TONELADAS:
            series[_clave(grupo)][temporada] = _usd_kg(volumen, valor)
    for clave, serie in series.items():
        if clave in grupos:
            grupos[clave]["serie_usd_kg"] = dict(
                sorted(serie.items(), key=lambda kv: _orden_de_temporada(kv[0])))

    return {
        "version": "1.0",
        "generado": datetime.now().isoformat(timespec="seconds"),
        "origen": os.path.basename(origen),
        "encabezado_del_archivo": encabezado,
        "fuente_declarada": (
            "ODEPA, Boletin de la fruta. Datos de comercio exterior del "
            "Servicio Nacional de Aduanas."),
        "tipo": "FOB exportacion",
        "temporada_de_referencia": ultima,
        "temporadas_cerradas": cerradas,
        "temporadas_parciales": parciales,
        "minimo_toneladas": MINIMO_TONELADAS,
        "procesamientos": list(PROCESAMIENTOS),
        "metodo": (
            "USD/kg = suma de «Valor (USD FOB)» / (suma de «Volumen (ton)» x "
            "1000), agregando todos los paises de destino. Es un promedio "
            "ponderado por volumen, que es lo que corresponde: un embarque de "
            "1.000 t pesa mas que uno de 10."),
        "advertencia": (
            "UN FOB NO ES LO QUE RECIBE EL PRODUCTOR. Falta restarle packing, "
            "frio, flete y comision de la exportadora (§10.2). Ponerlo tal "
            "cual en la fila de precio del flujo duplicaria el ingreso del "
            "cliente."),
        "advertencia_temporada": (
            "Solo se usan temporadas CERRADAS (septiembre a agosto). La "
            "temporada en curso que trae el archivo esta cortada donde "
            "alcanzo el boletin y sesga el precio hacia la fruta temprana."),
        "grupos": dict(sorted(grupos.items())),
    }


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("archivo", help="el «Exportaciones por Temporada» de ODEPA")
    p.add_argument("--salida", default=None,
                   help="por omision, config/precios_fob_odepa.json")
    args = p.parse_args()

    if not os.path.exists(args.archivo):
        raise SystemExit("  No existe el archivo: {}".format(args.archivo))

    print("=" * 74)
    print("PRECIOS FOB DE EXPORTACION - ODEPA / Aduanas")
    print("=" * 74)
    print("  leyendo {} ...".format(os.path.basename(args.archivo)))

    por_producto, por_grupo, temporadas, leidas, encabezado = leer(args.archivo)
    print("  {:,} filas utiles, {} temporadas".format(leidas, len(temporadas)))

    datos = armar(por_producto, por_grupo, temporadas, args.archivo, encabezado)

    salida = args.salida or str(ajustes.raiz / "config" / "precios_fob_odepa.json")
    with io.open(salida, "w", encoding="utf-8") as fh:
        json.dump(datos, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print()
    print("  temporada de referencia: {}".format(datos["temporada_de_referencia"]))
    if datos["temporadas_parciales"]:
        print("  descartadas por incompletas: {}".format(
            ", ".join(datos["temporadas_parciales"])))
    print("  {} grupos con precio".format(len(datos["grupos"])))
    print()
    print("  {:<24} {:>11} {:>10}   {}".format("grupo", "ton", "USD/kg", ""))
    print("  " + "-" * 70)
    for clave, g in sorted(datos["grupos"].items(), key=lambda kv: -kv[1]["toneladas"]):
        marca = "  <- MEZCLA: ver desglose" if g["es_mezcla"] else ""
        print("  {:<24} {:>11,.0f} {:>10.2f}{}".format(
            g["nombre"][:24], g["toneladas"], g["usd_kg"], marca))

    mezclas = [g for g in datos["grupos"].values() if g["es_mezcla"]]
    if mezclas:
        print()
        print("  Grupos que MEZCLAN productos de precio muy distinto:")
        for g in mezclas:
            print("    {}".format(g["nombre"]))
            for c in g["codigos"]:
                print("      {:<10} {:<40} {:>9,.0f} t {:>8.2f}".format(
                    c["codigo"], c["producto"][:40], c["toneladas"], c["usd_kg"]))

    print()
    print("  Escrito en {}".format(salida))
    print("  OJO: es FOB. No va directo a la fila de precio del flujo (§10.2).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
