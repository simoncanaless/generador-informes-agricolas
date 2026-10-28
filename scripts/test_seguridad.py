"""Seguridad del sistema: lo que se reviso y arreglo el 2026-10-05, para que
no se vuelva a perder.

el usuario pidio "trabajar en la seguridad del sistema". La revision encontro, en
orden de gravedad:

1. Tres rutas de expedientes que NO preguntaban de quien era el informe
   (guardar el catastro, sus curvas faltantes, su resumen): cualquiera con
   sesion podia leer o reemplazar el catastro de un informe ajeno.
2. `escapar()` de la pagina no escapaba comillas: un texto con «"» en un dato
   compartido salia del atributo y pasaba a ser HTML de la pagina de otro.
3. Contrasenas sin limite de intentos, y una respuesta mas rapida para un RUT
   que no existe (se podia saber que RUT tienen cuenta midiendo el tiempo).
4. El servidor aceptaba conexiones de todo el wifi de la casa, no solo de
   este computador y de Tailscale.
5. Sin cabeceras de seguridad ni politica de contenido.
6. Un texto que empezaba con "=" llegaba al Excel como formula viva.
7. Dos subidas sin tope de tamano; enlaces de descarga adivinables.

Ninguna prueba de aca necesita la base de datos ni Excel: los repositorios
se reemplazan por dobles en memoria y la app se llama con el TestClient.

Uso: python scripts/test_seguridad.py
"""

import asyncio
import base64
import hashlib
import inspect
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
from pathlib import Path
from types import SimpleNamespace

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)
    return condicion


# ------------------------------------------------------------- utilidades

def rutas_de_la_api():
    """(metodo, ruta, APIRoute) de todas las rutas, entrando a los routers
    incluidos (FastAPI 0.14x los guarda como _IncludedRouter)."""
    import main
    from fastapi.routing import APIRoute

    def recorrer(rutas):
        for r in rutas:
            if type(r).__name__ == "_IncludedRouter":
                yield from recorrer(r.original_router.routes)
            elif isinstance(r, APIRoute):
                yield r

    return list(recorrer(main.app.routes))


def dependencias_de(dependant, acumulado=None):
    acumulado = set() if acumulado is None else acumulado
    for d in dependant.dependencies:
        acumulado.add(d.call)
        dependencias_de(d, acumulado)
    return acumulado


class RepoUsuariosFalso:
    """Las cuentas, en memoria. Misma interfaz que RepositorioUsuarios."""

    def __init__(self):
        from services.usuarios import cifrar_contrasena
        self.filas = {}
        self.agregar("11.111.111-1", "Persona Prueba", "clave-de-prueba-1", cifrar_contrasena)

    def agregar(self, rut, nombre, clave, cifrar):
        from services.usuarios import normalizar_rut
        fila = SimpleNamespace(id=uuid.uuid4(), rut=normalizar_rut(rut), nombre=nombre,
                               clave=cifrar(clave), activo=True)
        self.filas[str(fila.id)] = fila
        return fila

    async def buscar_por_rut(self, rut):
        return next((f for f in self.filas.values() if f.rut == rut), None)

    async def obtener(self, id_usuario):
        return self.filas.get(str(id_usuario))

    async def listar(self):
        return list(self.filas.values())

    async def cambiar_clave(self, id_usuario, clave):
        self.filas[str(id_usuario)].clave = clave

    async def marcar_entrada(self, id_usuario):
        return None


# --------------------------------------------------------------- pruebas

def probar_rutas_protegidas():
    print("\n[1] Toda ruta pide sesion, salvo las que deben estar abiertas")
    from api.dependencias import obtener_usuario_actual

    abiertas_permitidas = {
        ("POST", "/acceso/entrar"), ("POST", "/acceso/salir"),
        ("GET", "/salud"), ("GET", "/salud/modelos"), ("GET", "/"),
    }
    rutas = rutas_de_la_api()
    comprobar(len(rutas) > 50, "se encontraron las rutas de verdad ({})".format(len(rutas)))
    abiertas = set()
    for ruta in rutas:
        if obtener_usuario_actual not in dependencias_de(ruta.dependant):
            for metodo in ruta.methods:
                abiertas.add((metodo, ruta.path))
    sobrantes = abiertas - abiertas_permitidas
    comprobar(not sobrantes,
              "ninguna ruta nueva quedo abierta sin estar en la lista", sorted(sobrantes))


def probar_rutas_de_expediente_preguntan_de_quien_es():
    print("\n[2] Toda ruta con {id_expediente} pregunta de quien es el informe")
    sin_guardia = []
    for ruta in rutas_de_la_api():
        if "{id_expediente}" not in ruta.path:
            continue
        fuente = inspect.getsource(ruta.endpoint)
        if not any(g in fuente for g in ("_mio_o_404(", "puede_verlo(", "_soy_dueno_o_no(")):
            sin_guardia.append("{} {}".format(",".join(sorted(ruta.methods)), ruta.path))
    comprobar(not sin_guardia,
              "todas pasan por _mio_o_404 / puede_verlo / _soy_dueno_o_no", sin_guardia)

    # Y de verdad, con un informe AJENO: las tres que estaban abiertas.
    from fastapi.testclient import TestClient

    import main
    from api.dependencias import (
        obtener_repositorio_expedientes, obtener_repositorio_opcional, obtener_usuario_actual,
    )

    class RepoExpedientesFalso:
        tocado = False

        async def puede_verlo(self, id_expediente, usuario_id):
            return False                      # es de otra persona

        async def obtener(self, id_expediente):
            return SimpleNamespace(cliente="Cliente Ajeno", propietario_id=uuid.uuid4())

        async def reemplazar_plantaciones(self, *a, **k):
            RepoExpedientesFalso.tocado = True
            return 0

        async def resumen_por_especie(self, *a, **k):
            RepoExpedientesFalso.tocado = True
            return [{"especie": "secreta"}]

        async def plantaciones(self, *a, **k):
            RepoExpedientesFalso.tocado = True
            return []

    repo = RepoExpedientesFalso()
    main.app.dependency_overrides[obtener_usuario_actual] = lambda: SimpleNamespace(
        id=uuid.uuid4(), nombre="Yo", rut="111111111", activo=True)
    main.app.dependency_overrides[obtener_repositorio_expedientes] = lambda: repo
    main.app.dependency_overrides[obtener_repositorio_opcional] = lambda: repo
    try:
        cliente = TestClient(main.app)
        ajeno = str(uuid.uuid4())
        r1 = cliente.put("/expedientes/{}/catastro".format(ajeno), json={"plantaciones": []})
        r2 = cliente.get("/expedientes/{}/resumen".format(ajeno))
        r3 = cliente.get("/expedientes/{}/curvas-faltantes".format(ajeno))
    finally:
        main.app.dependency_overrides.clear()
    comprobar(r1.status_code == 404, "reemplazar el catastro de un informe ajeno: 404", r1.status_code)
    comprobar(r2.status_code == 404, "leer su resumen por especie: 404", r2.status_code)
    comprobar(r3.status_code == 404, "leer sus curvas faltantes (y su cliente): 404", r3.status_code)
    comprobar(not RepoExpedientesFalso.tocado,
              "y el repositorio ni se toco: se corta antes de leer o escribir")


def probar_redes_permitidas():
    print("\n[3] Solo entran este computador y Tailscale")
    from api.seguridad import REDES_POR_DEFECTO, RedesPermitidas, ip_permitida, leer_redes
    from config.settings import Ajustes

    redes = leer_redes(REDES_POR_DEFECTO)
    casos = {
        "127.0.0.1": True, "::1": True, "100.101.102.103": True, "100.64.0.1": True,
        "fd7a:115c:a1e0::1": True, "::ffff:100.70.1.2": True, "testclient": True,
        "192.168.1.20": False, "10.0.0.5": False, "172.16.4.4": False,
        "100.128.0.1": False, "8.8.8.8": False, "fe80::1": False, "": False,
        None: False, "no-es-una-ip": False,
    }
    malos = {ip: esperado for ip, esperado in casos.items() if ip_permitida(ip, redes) != esperado}
    comprobar(not malos, "localhost y Tailscale si; el wifi de la casa e internet no", malos)
    comprobar(Ajustes.model_fields["redes_permitidas"].default == REDES_POR_DEFECTO,
              "el valor por defecto de la configuracion es el mismo")

    # El middleware de verdad, con un cliente del wifi de la casa.
    enviados = []

    async def app_interna(scope, receive, send):
        enviados.append("llego a la app")

    async def send(mensaje):
        enviados.append(mensaje)

    async def receive():
        return {"type": "http.request"}

    puerta = RedesPermitidas(app_interna, redes)
    asyncio.run(puerta({"type": "http", "path": "/app/", "client": ("192.168.1.20", 5000)},
                       receive, send))
    comprobar("llego a la app" not in enviados and enviados
              and enviados[0].get("status") == 403,
              "desde 192.168.1.20 contesta 403 y la app ni se entera")
    cuerpo = enviados[1]["body"].decode("utf-8") if len(enviados) > 1 else ""
    comprobar("Tailscale" in cuerpo, "y la pagina dice por donde entrar")
    enviados.clear()
    asyncio.run(puerta({"type": "http", "path": "/salud", "client": ("100.88.1.2", 5000)},
                       receive, send))
    comprobar(enviados == ["llego a la app"], "desde Tailscale pasa")

    import main
    from api.seguridad import CabecerasDeSeguridad
    montados = [m.cls for m in main.app.user_middleware]
    comprobar(RedesPermitidas in montados and CabecerasDeSeguridad in montados,
              "las dos capas estan montadas en la aplicacion")
    comprobar(montados and montados[0] is RedesPermitidas,
              "y la de redes es la de mas afuera: lo primero que ve cada peticion")


def probar_cabeceras_y_politica():
    print("\n[4] Cabeceras de seguridad y politica de contenido de la pagina")
    from fastapi.testclient import TestClient

    import main
    from api.seguridad import huellas_de_guiones

    cliente = TestClient(main.app)
    api = cliente.get("/acceso/yo")       # 401, pero igual con cabeceras
    for nombre, valor in (("x-content-type-options", "nosniff"),
                          ("x-frame-options", "DENY"),
                          ("referrer-policy", "same-origin"),
                          ("cache-control", "no-store")):
        comprobar(api.headers.get(nombre) == valor,
                  "la API responde {}: {}".format(nombre, valor), api.headers.get(nombre))

    pagina = cliente.get("/app/")
    comprobar(pagina.status_code == 200, "la pagina se sirve")
    politica = pagina.headers.get("content-security-policy", "")
    comprobar(politica, "con Content-Security-Policy")
    script_src = next((p for p in politica.split(";") if p.strip().startswith("script-src")), "")
    comprobar("unsafe-inline" not in script_src and "unsafe-eval" not in script_src,
              "sin 'unsafe-inline' ni 'unsafe-eval' para guiones", script_src)
    comprobar("frame-ancestors 'none'" in politica and "object-src 'none'" in politica,
              "y no se puede meter en un marco ni cargar plugins")

    # La huella tiene que ser la del guion que el navegador VE: se calcula
    # aparte, a mano, sobre el texto normalizado.
    html = (RAIZ / "web" / "index.html").read_text(encoding="utf-8")
    normal = html.replace("\r\n", "\n").replace("\r", "\n")
    guiones = re.findall(r"<script>([\s\S]*?)</script>", normal)
    comprobar(len(guiones) == 1, "la pagina tiene un solo guion en linea", len(guiones))
    esperada = "'sha256-{}'".format(base64.b64encode(
        hashlib.sha256(guiones[0].encode("utf-8")).digest()).decode())
    comprobar(esperada in script_src, "la huella del guion coincide con el calculo propio")
    con_crlf = html.replace("\n", "\r\n")
    comprobar(huellas_de_guiones(con_crlf) == huellas_de_guiones(html),
              "y no cambia si el archivo trae saltos de linea de Windows")
    comprobar(not re.search(r"\son[a-z]+\s*=\s*[\"']", _solo_marcado_y_plantillas(html)),
              "la pagina no usa atributos onclick=/onload=... (la politica los bloquearia)")
    comprobar("javascript:" not in html, "ni enlaces javascript:")


def _solo_marcado_y_plantillas(html):
    """El HTML y los textos de plantilla del guion, sin comentarios."""
    sin_comentarios = re.sub(r"<!--[\s\S]*?-->", "", html)
    return re.sub(r"//[^\n]*", "", sin_comentarios)


def probar_freno_de_intentos():
    print("\n[5] Freno a quien prueba contrasenas")
    from services import intentos as modulo
    from services.intentos import Intentos

    reloj = [1000.0]
    i = Intentos(reloj=lambda: reloj[0])
    claves = ["rut:111111111", "ip:100.64.0.9"]
    for _ in range(modulo.FALLOS_LIBRES):
        comprobar(i.espera(claves) == 0, "los primeros fallos no frenan") if _ == 0 else None
        i.fallo(claves)
    comprobar(i.espera(claves) == modulo.ESPERA_INICIAL,
              "al quinto fallo hay que esperar {:.0f} s".format(modulo.ESPERA_INICIAL),
              i.espera(claves))
    i.fallo(claves)
    comprobar(i.espera(claves) == 2 * modulo.ESPERA_INICIAL, "y la espera se duplica",
              i.espera(claves))
    for _ in range(20):
        i.fallo(claves)
    comprobar(i.espera(claves) == modulo.ESPERA_MAXIMA, "hasta un tope de 15 minutos",
              i.espera(claves))
    comprobar(i.espera(["rut:otro", "ip:100.64.0.9"]) > 0,
              "cambiar de RUT desde el mismo equipo no sirve")
    comprobar(i.espera(["rut:111111111", "ip:otro"]) > 0,
              "ni cambiar de equipo con el mismo RUT")
    reloj[0] += modulo.OLVIDO + 1
    comprobar(i.espera(claves) == 0, "una hora despues se olvida")
    i.fallo(claves)
    i.acierto(claves)
    comprobar(i.espera(claves) == 0 and not i._fallos, "entrar bien borra los fallos")


def probar_entrar_y_cambiar_clave():
    print("\n[6] Entrar, el freno de verdad, y cambiar la clave cierra las otras sesiones")
    from fastapi.testclient import TestClient

    import main
    from api.dependencias import obtener_repositorio_usuarios
    from api.routers import acceso
    from services.intentos import Intentos

    repo = RepoUsuariosFalso()
    main.app.dependency_overrides[obtener_repositorio_usuarios] = lambda: repo
    acceso.intentos = Intentos()
    try:
        a = TestClient(main.app)
        b = TestClient(main.app)
        r = a.post("/acceso/entrar", json={"rut": "11.111.111-1", "contrasena": "clave-de-prueba-1"})
        comprobar(r.status_code == 200, "entra con la clave correcta", r.status_code)
        comprobar(a.get("/acceso/yo").status_code == 200, "y la sesion vale")
        b.post("/acceso/entrar", json={"rut": "11.111.111-1", "contrasena": "clave-de-prueba-1"})
        comprobar(b.get("/acceso/yo").status_code == 200, "una segunda sesion, en otro equipo")

        r = a.post("/acceso/clave", json={"actual": "clave-de-prueba-1", "nueva": "clave-nueva-22"})
        comprobar(r.status_code == 200, "cambia la clave", r.status_code)
        comprobar(a.get("/acceso/yo").status_code == 200,
                  "quien la cambio sigue dentro")
        r = b.get("/acceso/yo")
        comprobar(r.status_code == 401, "la OTRA sesion, abierta con la clave vieja, queda fuera",
                  r.status_code)

        # El freno, por la ruta de verdad.
        acceso.intentos = Intentos()
        c = TestClient(main.app)
        codigos = [c.post("/acceso/entrar", json={"rut": "11.111.111-1", "contrasena": "mala"}).status_code
                   for _ in range(5)]
        comprobar(codigos == [401] * 5, "cinco claves malas: 401 cada una", codigos)
        r = c.post("/acceso/entrar", json={"rut": "11.111.111-1", "contrasena": "clave-nueva-22"})
        comprobar(r.status_code == 429, "la sexta, aunque sea la buena, espera: 429", r.status_code)
        comprobar(r.headers.get("retry-after"), "con Retry-After")
        comprobar("Espere" in r.json().get("detail", ""), "y un mensaje que dice cuanto",
                  r.json().get("detail"))
        acceso.intentos = Intentos()
        r = c.post("/acceso/entrar", json={"rut": "99.999.999-9", "contrasena": "x"})
        comprobar(r.status_code == 401, "un RUT que no existe da el mismo 401", r.status_code)
        r = c.post("/acceso/entrar", json={"rut": "1", "contrasena": "x" * 5000})
        comprobar(r.status_code == 422, "una contrasena gigante se rechaza sin procesarla",
                  r.status_code)
    finally:
        main.app.dependency_overrides.clear()
        acceso.intentos = Intentos()

    # Mismo trabajo para un RUT inexistente que para uno real.
    from services.usuarios import cifrar_contrasena, comparar_con_senuelo, contrasena_correcta
    guardada = cifrar_contrasena("otra-clave-123")
    comparar_con_senuelo("calentar")
    t = time.perf_counter(); comparar_con_senuelo("x"); senuelo = time.perf_counter() - t
    t = time.perf_counter(); contrasena_correcta("x", guardada); real = time.perf_counter() - t
    comprobar(senuelo > real * 0.5,
              "un RUT inexistente tarda lo mismo que uno real ({:.0f} ms contra {:.0f} ms)".format(
                  senuelo * 1000, real * 1000))


def probar_formulas_en_el_excel():
    print("\n[7] Un texto con '=' adelante nunca llega al Excel como formula")
    from openpyxl import load_workbook

    from config.settings import obtener_ajustes
    from schemas.extraccion import ExtraccionLegal
    from schemas.informe import Antecedentes, Calificacion, InformeFitogreen
    from services.excel_engine import MotorExcel
    from services.excel_engine_rabobank import MotorRabobank
    from services.mapping import cargar_mapa
    from services.retasacion_pendientes import leer_borrador

    a = Antecedentes(cliente='=HYPERLINK("http://x.y","ver")', rut="76.543.210-3")
    comprobar(a.cliente == 'HYPERLINK("http://x.y","ver")', "el informe le quita el signo", a.cliente)
    comprobar(isinstance(Calificacion("Bueno"), Calificacion), "las enumeraciones no se tocan")
    legal = ExtraccionLegal.model_validate(
        {"gravamenes": ["=WEBSERVICE(\"http://x.y\")", "hipoteca a favor del banco"]})
    comprobar(legal.gravamenes[0].startswith("WEBSERVICE"),
              "lo que lee la IA de una escritura tambien", legal.gravamenes)
    retasacion, _ = leer_borrador({"cliente": {"nombre": "=1+1"},
                                   "observaciones_portada": "-5% por heladas"})
    comprobar(retasacion.cliente.nombre == "1+1", "la retasacion tambien", retasacion.cliente.nombre)
    comprobar(retasacion.observaciones_portada == "-5% por heladas",
              "un texto que empieza con '-' queda igual (openpyxl no lo toma por formula)")

    from services.visita import _asignar
    informe = InformeFitogreen(antecedentes=Antecedentes(cliente="Prueba", rut="76.543.210-3"))
    _asignar(informe, "antecedentes.ejecutivo", "=1+1")
    comprobar(informe.antecedentes.ejecutivo == "1+1",
              "y la visita, que asigna sin pasar por los validadores", informe.antecedentes.ejecutivo)

    ajustes = obtener_ajustes()
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "santander.xlsx"
        motor = MotorExcel(ajustes.ruta_plantilla, cargar_mapa(ajustes.ruta_mapping))
        motor.generar(InformeFitogreen(antecedentes=Antecedentes(cliente="=1+1", rut="76.543.210-3")),
                      destino)
        celda = load_workbook(destino)["Cliente"]["B10"]
        comprobar(celda.data_type != "f" and celda.value == "1+1",
                  "Santander: Cliente!B10 es texto, no formula", (celda.data_type, celda.value))

        plantilla, mapping = ajustes.rutas_de_plantilla("rabobank")
        destino = Path(tmp) / "rabobank.xlsx"
        MotorRabobank(plantilla, cargar_mapa(mapping)).generar(retasacion, destino)
        celda = load_workbook(destino)["PORTADA"]["C14"]
        comprobar(celda.data_type != "f" and celda.value == "1+1",
                  "Rabobank: PORTADA!C14 es texto, no formula", (celda.data_type, celda.value))


def probar_escapar_de_la_pagina():
    print("\n[8] escapar() de la pagina sirve tambien dentro de un atributo")
    html = (RAIZ / "web" / "index.html").read_text(encoding="utf-8")
    m = re.search(r"function escapar\(texto\) \{[\s\S]*?\n\}\n", html)
    comprobar(m is not None, "existe escapar()")
    cuerpo = m.group(0) if m else ""
    comprobar("&quot;" in cuerpo and "&#39;" in cuerpo, "escapa comillas dobles y simples")
    nodo = shutil.which("node")
    if not nodo:
        print("  (sin Node.js: no se ejecuta, solo se lee)")
        return
    guion = cuerpo + """
const casos = ['"><img src=x onerror=alert(1)>', "' onmouseover='x", "a & b < c", null, 12];
console.log(JSON.stringify(casos.map(escapar)));"""
    salida = subprocess.run([nodo, "-e", guion], capture_output=True, text=True, timeout=30)
    resultado = json.loads(salida.stdout or "[]")
    comprobar(resultado == ["&quot;&gt;&lt;img src=x onerror=alert(1)&gt;",
                            "&#39; onmouseover=&#39;x", "a &amp; b &lt; c", "", "12"],
              "ejecutada: nada sale del atributo ni abre una etiqueta", resultado)
    insertado = '<input value="{}">'.format(resultado[0] if resultado else "")
    comprobar(insertado.count('"') == 2, "un value=\"...\" con eso adentro sigue siendo un solo atributo")


def probar_subidas_y_descargas():
    print("\n[9] Toda subida tiene tope, y los enlaces de descarga no se adivinan")
    from fastapi import UploadFile

    sin_tope = []
    for ruta in rutas_de_la_api():
        firma = inspect.signature(ruta.endpoint)
        sube = any("UploadFile" in str(p.annotation) for p in firma.parameters.values())
        if sube and "leer_acotado(" not in inspect.getsource(ruta.endpoint):
            sin_tope.append(ruta.path)
    comprobar(not sin_tope, "cada ruta que recibe archivos los lee con leer_acotado", sin_tope)

    from api.routers import asistente, catastro, informes, retasaciones
    for modulo, funcion in ((asistente, "generar"), (informes, "generar"),
                            (retasaciones, "generar"), (catastro, "exportar")):
        fuente = inspect.getsource(getattr(modulo, funcion))
        comprobar("secrets.token_hex(" in fuente,
                  "{}.{}: el nombre del archivo lleva una parte al azar".format(
                      modulo.__name__.rsplit(".", 1)[-1], funcion))

    from config.settings import Ajustes
    comprobar(Ajustes.model_fields["api_cors_origenes"].default == "",
              "ningun otro sitio puede llamar a la API con la sesion del usuario (CORS vacio)")


def main():
    print("=" * 72)
    print("PRUEBA DE SEGURIDAD")
    print("=" * 72)
    probar_rutas_protegidas()
    probar_rutas_de_expediente_preguntan_de_quien_es()
    probar_redes_permitidas()
    probar_cabeceras_y_politica()
    probar_freno_de_intentos()
    probar_entrar_y_cambiar_clave()
    probar_formulas_en_el_excel()
    probar_escapar_de_la_pagina()
    probar_subidas_y_descargas()

    print("\n" + "=" * 72)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
