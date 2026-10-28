"""Prueba del generador del formato final de Santander.

    python scripts/test_motor_santander.py

Cubre las tres piezas nuevas:

- ``services/disposicion.py``: repetir las partes del molde y reescribir cada
  formula para su copia (el bloque de una plantacion lee SU fila de
  Supuestos; un total suma todas las copias; el resumen por especie suma los
  bloques de sus plantaciones).
- ``services/supuestos.py``: la hoja Supuestos propuesta desde la biblioteca
  (el agrónomo, 2026-10-14: "opcion a").
- ``services/excel_engine_santander.py``: el informe completo, con un caso
  inventado de tres especies, siete plantaciones y dos unidades prediales.

Y con el Excel real (best-effort): cero celdas en error, y el EBITDA de los
ocho anios calculado aparte en Python igual al del Excel -dos caminos
independientes al mismo numero, el contraste que usa todo el proyecto-.
"""

import io
import os
import re
import subprocess
import sys
import tempfile
import warnings
import zipfile
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

from config.settings import ajustes  # noqa: E402
from schemas.curvas import EstructuraCosto, ParametrosComerciales  # noqa: E402
from schemas.informe import (  # noqa: E402
    Antecedentes, DerechoAgua, FilaAnexo, Fotografia, InformeFitogreen, Lote, Plantacion,
    Sensibilizacion, SistemaComercializacion, SistemaProduccion, SupuestoPlantacion, UnidadPredial,
)
from services.catastro_import import cargar_reglas  # noqa: E402
from services.curvas import curva_desde_parametros  # noqa: E402
from services.disposicion import Disposicion, Parte  # noqa: E402
from services.excel_engine_santander import (  # noqa: E402
    MotorSantander, temporada_real_por_defecto, temporadas_siguientes)
from services.renderizado import (  # noqa: E402
    GUION_ARRANQUE_IMPRESORA_SEGURA, GUION_LIBERAR_COM_EXCEL, GUION_RESTAURAR_IMPRESORA)
from services.supuestos import proponer_supuestos  # noqa: E402

PLANTILLA = ajustes.raiz / "templates" / "Formato Fitogreen Santander 2026.xlsx"
MAPPING = ajustes.raiz / "config" / "mapping_santander_2026.json"
FLUJO = "Flujo proyectado "

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)
    return condicion


# ------------------------------------------------------------ el caso

def _foto(carpeta: Path, nombre: str, color) -> str:
    ruta = carpeta / nombre
    Image.new("RGB", (800, 500), color).save(ruta, "JPEG")
    return str(ruta)


def informe_de_prueba(carpeta: Path) -> InformeFitogreen:
    """Inventado: tres especies -con plantaciones de cada una repartidas en el
    catastro, para que el motor tenga que agruparlas-, dos unidades
    prediales, una plantacion en formacion y una sensibilizacion."""
    inf = InformeFitogreen(antecedentes=Antecedentes(
        cliente="Agrícola Ejemplo SpA", rut="76.543.210-3", ejecutivo="Ejecutivo Prueba",
        fecha_visita=date(2026, 9, 20), sucursal="Talca",
        fecha_emision_informe=date(2026, 10, 1), agronomo_informante="Agrónomo Prueba",
        n_solicitud="1234567"))

    def sup(pct, p1, p2, mi, fijo, var, rends, **kw):
        return SupuestoPlantacion(
            pct_exportacion=pct, precio_exportacion_usd_kg=p1,
            precio_exportacion_siguientes_usd_kg=p2, precio_mercado_interno_usd_kg=mi,
            costo_fijo_usd_ha=fijo, costo_variable_usd_kg=var, rendimientos_kg_ha=rends, **kw)

    def plant(especie, variedad, anio, ha, predio, supuesto, **kw):
        return Plantacion(especie=especie, variedad=variedad, anio_plantacion=anio,
                          superficie_ha=ha, predio=predio, riego=kw.pop("riego", "Goteo"),
                          dominio=kw.pop("dominio", "Propietario"), royalty=kw.pop("royalty", False),
                          supuestos=supuesto, **kw)

    inf.estructura_productiva.plantaciones = [
        plant("Cerezo", "Lapins", "2015", 10.5, "Unidad Norte",
              sup(0.8, 2.0, 2.0, 0.1, 7000, 0.7, [15000] * 8)),
        plant("Arándano", "Legacy", "2012", 5.2, "Unidad Norte",
              sup(0.75, 3.0, 3.0, 1.5, 9000, 1.2, [18000] * 8)),
        plant("Cerezo", "Regina", "2020", 3.0, "Unidad Sur",
              sup(0.8, 2.2, 2.0, 0.1, 7000, 0.7, [9000, 12000] + [15000] * 6),
              dominio="Arrendatario"),
        plant("Arándano", "Sekoya Pop", "2026", 2.0, "Unidad Sur",
              sup(0.8, 4.0, 4.0, 1.5, 9000, 1.2, [0, 3500, 10000, 16000] + [20000] * 4,
                  costo_formacion_1_usd_ha=2500, costo_formacion_2_usd_ha=3500), royalty=True),
        plant("Avellano", "Giffoni", "2014", 12.0, "Unidad Sur",
              sup(1.0, 4.2, 3.6, 0, 2400, 0.27, [2800] * 8)),
        plant("Cerezo", "Santina", "2024", 1.4, "Unidad Norte",
              sup(0.8, 2.5, 2.5, 0.1, 7000, 0.7, [5000, 12000] + [15000] * 6),
              riego="Microaspersion"),
        plant("Avellano", "Barcelona", "2016", 6.0, "Unidad Sur",
              sup(1.0, 3.6, 3.0, 0, 2400, 0.27, [2800] * 8)),
    ]
    dot = inf.dotacion_recursos
    dot.lotes = [
        Lote(unidad_predial="Unidad Norte", nombre="Lote A", rol="1234-5", region="VII",
             comuna="Talca", superficie_legal_ha=20.0, superficie_real_ha=20.0),
        Lote(unidad_predial="Unidad Norte", nombre="Lote B", rol="1234-6", region="VII",
             comuna="Talca", superficie_legal_ha=5.0, superficie_real_ha=5.5),
        Lote(unidad_predial="Unidad Sur", nombre="Hijuela 1", rol="2234-1", region="VII",
             comuna="Talca", superficie_legal_ha=30.0, superficie_real_ha=30.0),
    ]
    dot.unidades_prediales = [
        UnidadPredial(nombre="Unidad Norte", iii_iv_riego_ha=25.0, topografia="Plana",
                      capacidad_tranques_m3=5000),
        UnidadPredial(nombre="Unidad Sur", i_ii_riego_ha=10.0, iii_iv_riego_ha=20.0),
    ]
    dot.derechos_agua = [
        DerechoAgua(predio="Unidad Norte", tipo_fuente="Subterránea", nombre_fuente="Pozo",
                    cantidad=25, caudal_real_ls=22),
        DerechoAgua(predio="Unidad Sur", tipo_fuente="Canal", nombre_fuente="Canal",
                    cantidad=12, caudal_nominal_ls=18, caudal_real_ls=18),
        DerechoAgua(predio="Unidad Sur", tipo_fuente="Subterránea", nombre_fuente="Pozo 2",
                    cantidad=15),
    ]
    dot.notas_suelos = "Nota de prueba sobre los suelos."
    ep = inf.estructura_productiva
    ep.sistemas_produccion = [SistemaProduccion(especie="Cerezo", estado_fitosanitario="Bueno")]
    ep.sistemas_comercializacion = [SistemaComercializacion(
        especie="Arándano", mercado="Exportación", certificaciones="GlobalG.A.P.")]
    inf.anexo.resultado_real.temporada = "2025/26"
    inf.anexo.resultado_real.filas = [
        FilaAnexo(especie="Cerezo", produccion_exportacion_kg=150000,
                  produccion_mercado_interno_kg=20000, ventas_exportacion_pesos=280e6,
                  ventas_mercado_interno_pesos=2e6, costos_directos_pesos=120e6),
        FilaAnexo(especie="Arándano", produccion_exportacion_kg=70000,
                  produccion_mercado_interno_kg=15000, ventas_exportacion_pesos=190e6,
                  ventas_mercado_interno_pesos=20e6, costos_directos_pesos=80e6),
        FilaAnexo(especie="Avellano", produccion_exportacion_kg=50000,
                  produccion_mercado_interno_kg=0, ventas_exportacion_pesos=180e6,
                  ventas_mercado_interno_pesos=0, costos_directos_pesos=40e6),
    ]
    inf.anexo.resultado_real.administracion_pesos = 90e6
    inf.anexo.resultado_real.gastos_generales_pesos = 15e6
    inf.anexo.ventas_f29_pesos = [20e6, 0, 5e6, 1e6, 100e6, 300e6, 30e6, 60e6, 150e6, 80e6, 10e6, 20e6]
    inf.flujo_proyectado.tipo_cambio_pesos_usd = {"real": 940.0, "proyeccion": 950.0, "anterior": 900.0}
    inf.flujo_proyectado.uf = 39500.0
    inf.supuestos.administracion_usd_anio = 100000
    inf.supuestos.sensibilizacion = {"Cerezo": Sensibilizacion(precio_exportacion=-0.05)}
    inf.supuestos.notas = {"Arándano": "Nota de prueba."}
    for campo in ("gestion_administrativa_productiva", "canales_comerciales",
                  "calidad_activos_productivos", "conclusion_final"):
        bloque = getattr(inf.opinion, campo)
        bloque.texto, bloque.calificacion = "Texto de prueba.", "Bueno"
    inf.fotografias = [
        Fotografia(ruta=_foto(carpeta, "empl.jpg", (60, 120, 60)), categoria="emplazamiento"),
        Fotografia(ruta=_foto(carpeta, "f1.jpg", (120, 60, 60)), titulo="Cerezos", categoria="plantacion"),
        Fotografia(ruta=_foto(carpeta, "f2.jpg", (60, 60, 120)), titulo="Arándanos", categoria="plantacion"),
        Fotografia(ruta=_foto(carpeta, "f3.jpg", (120, 120, 60)), titulo="Avellanos", categoria="plantacion"),
        Fotografia(ruta=_foto(carpeta, "f4.jpg", (90, 90, 90)), titulo="Riego", categoria="riego"),
    ]
    return inf


def ebitda_en_python(inf: InformeFitogreen):
    """El EBITDA de los ocho anios, calculado sin pasar por el Excel."""
    salida = []
    for k in range(8):
        ingreso = costo = 0.0
        for p in inf.estructura_productiva.plantaciones:
            s = p.supuestos
            d = inf.supuestos.sensibilizacion.get(p.especie) or Sensibilizacion()
            rend = s.rendimientos_kg_ha[k] * (1 + d.rendimiento)
            precio = s.precio_exportacion_usd_kg if k == 0 else s.precio_exportacion_siguientes_usd_kg
            prod = p.superficie_ha * rend
            ingreso += prod * s.pct_exportacion * precio * (1 + d.precio_exportacion)
            ingreso += prod * (1 - s.pct_exportacion) * s.precio_mercado_interno_usd_kg * (
                1 + d.precio_mercado_interno)
            if k == 0 and s.costo_formacion_1_usd_ha is not None:
                c = s.costo_formacion_1_usd_ha
            elif k == 1 and s.costo_formacion_2_usd_ha is not None:
                c = s.costo_formacion_2_usd_ha
            else:
                c = s.costo_fijo_usd_ha + s.costo_variable_usd_kg * rend
            costo += p.superficie_ha * c * (1 + d.costos)
        salida.append(ingreso - costo - inf.supuestos.administracion_usd_anio)
    return salida


# ------------------------------------------------------------ pruebas

def probar_disposicion():
    print("\n[1] La disposicion: repetir partes y reescribir formulas")
    from openpyxl import load_workbook as cargar
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        molde = cargar(PLANTILLA)
    partes = [
        Parte("lista", "Plantaciones", 7, 7, "plantacion"),
        Parte("resumen", "Plantaciones", 15, 15, "especie"),
        Parte("bloques", FLUJO, 8, 17, "plantacion"),
        Parte("sup", FLUJO, 25, 25, "especie", agrega="plantacion"),
        Parte("secciones", "Supuestos", 21, 26, "especie"),
        Parte("filas", "Supuestos", 24, 24, "plantacion", dentro_de="secciones"),
    ]
    especie_de = [0, 0, 1]
    instancias = {
        "lista": lambda c: [{"plantacion": i, "especie": especie_de[i]} for i in range(3)],
        "bloques": lambda c: [{"plantacion": i, "especie": especie_de[i]} for i in range(3)],
        "resumen": lambda c: [{"especie": i} for i in range(2)],
        "sup": lambda c: [{"especie": i} for i in range(2)],
        "secciones": lambda c: [{"especie": i} for i in range(2)],
        "filas": lambda c: [{"plantacion": i, "especie": c["especie"]}
                            for i in range(3) if especie_de[i] == c["especie"]],
    }
    hojas = {h: molde[h].max_row for h in ("Plantaciones", FLUJO, "Supuestos")}
    disp = Disposicion(partes, [], instancias, hojas)
    comprobar(disp.total_filas(FLUJO) == molde[FLUJO].max_row + 20 + 1,
              "tres bloques de 10 filas y dos especies en el resumen")
    bloque_3 = disp.fila(FLUJO, 8, {"plantacion": 2, "especie": 1})
    comprobar(bloque_3 == 28, "el tercer bloque empieza 20 filas mas abajo", bloque_3)
    # Supuestos: secciones de 6 filas, la primera con dos filas de datos.
    f = disp.remapear("=Supuestos!$F$24*(1+Supuestos!$B$5)", FLUJO, 10,
                      {"plantacion": 2, "especie": 1})
    comprobar(f == "=Supuestos!$F$31*(1+Supuestos!$B$5)",
              "el bloque de la 3a plantacion lee SU fila de Supuestos", f)
    f = disp.remapear("=Plantaciones!A7", FLUJO, 8, {"plantacion": 1, "especie": 0})
    comprobar(f == "=Plantaciones!A8", "el encabezado del bloque lee su fila del catastro", f)
    f = disp.remapear("=D52-D53-D54", FLUJO, 58, {})
    comprobar(f == "=D73-D74-D75", "lo que no se repite se corre lo que crecio la hoja", f)
    total = disp.remapear("=SUM(D25:D25)", FLUJO, 24, {})
    comprobar(total == "=SUM(D45:D46)", "el total se estira a todas las especies", total)
    cadena = disp.remapear("=F9", FLUJO, 25, {"especie": 0})
    comprobar(cadena == "=F9+F19", "la especie suma los bloques de SUS plantaciones", cadena)
    rows = disp.remapear("=INDEX(X,ROWS($A$15:A15))", "Plantaciones", 15, {"especie": 1})
    comprobar(rows == "=INDEX(X,ROWS($A$17:A18))", "ROWS($A$15:A15) se ancla a la primera copia",
              rows)
    lista = disp.remapear("=SUMIFS($E$7:$E$7,$A$7:$A$7,A15)", "Plantaciones", 15, {"especie": 1})
    comprobar(lista == "=SUMIFS($E$7:$E$9,$A$7:$A$9,A18)",
              "un rango sobre la lista de plantaciones la cubre ENTERA (no solo su especie)", lista)
    seccion = disp.remapear("=SUM(D24:D24)", "Supuestos", 25, {"especie": 1})
    comprobar(seccion == "=SUM(D31:D31)", "el total de una seccion suma solo SUS filas", seccion)


def probar_supuestos():
    print("\n[2] Supuestos propuestos desde la biblioteca")
    curva = curva_desde_parametros(
        nombre="Cerezo prueba", especie="Cerezo", nivel_plena_produccion_kg_ha=15000,
        anios_a_plena_produccion=7, anios_improductivos=2, nivel_productivo_inicial_kg_ha=3000,
        comerciales=ParametrosComerciales(pct_exportacion=0.8, rnp_exportacion_usd_kg=2.1,
                                          rnp_mercado_interno_usd_kg=0.2),
        costos=EstructuraCosto(costo_base_usd_ha=7000, costo_por_kg_usd=0.7))
    joven = Plantacion(especie="Cerezo", anio_plantacion="2024", superficie_ha=1)
    adulta = Plantacion(especie="Cerezo", anio_plantacion="2010", superficie_ha=1)
    sin_curva = Plantacion(especie="Kiwi", anio_plantacion="2015", superficie_ha=1)
    temporadas = temporadas_siguientes("2025/26")
    comprobar(temporadas[0] == "2026/27" and temporadas[-1] == "2033/34",
              "las 8 temporadas siguen a la real")
    comprobar(temporada_real_por_defecto(date(2026, 10, 1)) == "2025/26"
              and temporada_real_por_defecto(date(2026, 3, 1)) == "2024/25",
              "la temporada real por defecto es la ultima cerrada (julio a junio)")
    avisos = proponer_supuestos([joven, adulta, sin_curva], [curva], temporadas)
    s = adulta.supuestos
    comprobar(s is not None and s.rendimientos_kg_ha == [15000.0] * 8,
              "una plantacion adulta rinde su meseta los 8 anios")
    r = joven.supuestos.rendimientos_kg_ha
    comprobar(r[0] < r[1] < r[4] and r[-1] == 15000.0,
              "una joven sube con la edad hasta la meseta", r)
    comprobar(s.pct_exportacion == 0.8 and s.precio_exportacion_usd_kg == 2.1
              and s.precio_exportacion_siguientes_usd_kg == 2.1
              and s.precio_mercado_interno_usd_kg == 0.2,
              "% exportado y precios salen de la curva")
    comprobar(s.costo_fijo_usd_ha == 7000 and s.costo_variable_usd_kg == 0.7,
              "costo fijo y variable salen de la curva")
    comprobar(sin_curva.supuestos is None and any("Kiwi" in a for a in avisos),
              "sin curva la fila queda vacia y se avisa")


def probar_generacion(carpeta: Path):
    print("\n[3] El informe completo con el motor nuevo")
    motor = MotorSantander(PLANTILLA, MAPPING,
                           reglas=cargar_reglas(ajustes.raiz / "config" / "reglas_agronomicas.json"))
    inf = informe_de_prueba(carpeta)
    destino = carpeta / "santander_prueba.xlsx"
    resultado = motor.generar(inf, destino)
    comprobar(destino.exists() and not resultado.avisos,
              "genera sin avisos", resultado.avisos[:3])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(destino)
    flujo = wb[FLUJO]
    bloques = [c.row for c in flujo["A"]
               if isinstance(c.value, str) and c.value.startswith("=Plantaciones!A")]
    comprobar(len(bloques) == 7, "un bloque de flujo por plantacion", len(bloques))
    especies = [flujo.cell(f, 1).value for f in bloques]
    plant = wb["Plantaciones"]
    nombres = [plant.cell(7 + i, 1).value for i in range(7)]
    comprobar(nombres == ["Cerezo"] * 3 + ["Arándano"] * 2 + ["Avellano"] * 2,
              "el catastro queda agrupado por especie, en el orden en que aparecen", nombres)
    comprobar(especies[0] == "=Plantaciones!A7" and especies[-1] == "=Plantaciones!A13",
              "cada bloque lee su plantacion")
    sup = wb["Supuestos"]
    titulos = [c.value for c in sup["A"] if c.value in ("CEREZO", "ARÁNDANO", "AVELLANO")]
    comprobar(titulos == ["CEREZO", "ARÁNDANO", "AVELLANO"], "una tabla de Supuestos por especie",
              titulos)
    comprobar([sup.cell(4, c).value for c in (2, 3, 4)] == ["Cerezo", "Arándano", "Avellano"]
              and sup.cell(5, 2).value == -0.05,
              "una columna de sensibilizacion por especie, con la de la visita")
    rec = wb["Recursos"]
    riego = [c.row for c in rec["A"] if c.value == "Unidad Predial:"]
    comprobar(len(riego) == 2, "un bloque de riego por unidad predial con plantaciones", riego)
    combinados = [str(r) for r in rec.merged_cells.ranges if r.min_col == 9 and r.min_row > riego[0]]
    comprobar(len(combinados) == 2, "la demanda del predio queda combinada a lo alto de sus especies",
              combinados)
    i_norte = rec.cell(riego[0] + 2, 9).value
    comprobar(i_norte == "=SUM(H{0}:H{1})".format(riego[0] + 2, riego[0] + 3),
              "la demanda del predio suma TODAS sus especies", i_norte)
    a1 = wb["Anexo 1 Resultado"]
    comprobar([a1.cell(5 + i, 1).value for i in range(3)] == ["Cerezo", "Arándano", "Avellano"],
              "el Anexo 1 en el mismo orden que el resumen del flujo")
    comprobar(flujo["E7"].value == "2025/26", "la temporada real en su unica celda")
    saltos = [b.id for b in flujo.row_breaks.brk]
    resumen = next(c.row for c in flujo["A"] if c.value == "RESUMEN DE INDICADORES GENERALES")
    comprobar(saltos and saltos[0] < resumen and saltos[0] > bloques[-1],
              "el salto de pagina queda antes del resumen", saltos)
    anexo4 = wb["Anexo 4 Imágenes"]
    comprobar(anexo4.max_row <= 6 and anexo4["A4"].value == "Cerezos",
              "el Anexo 4 deja solo las filas de fotos que se usan", anexo4.max_row)
    with zipfile.ZipFile(destino) as z:
        nombres_partes = z.namelist()
        hojas = {n: z.read(n).decode("utf-8") for n in nombres_partes
                 if re.match(r"xl/worksheets/sheet\d+\.xml$", n)}
        repetidos = []
        for n in nombres_partes:
            if re.match(r"xl/drawings/_rels/.*\.rels$", n):
                ids = re.findall(r'Id="([^"]+)"', z.read(n).decode("utf-8"))
                if len(ids) != len(set(ids)):
                    repetidos.append(n)
    comprobar(not repetidos, "ningun id de relacion repetido en los dibujos (imagenes rotas)",
              repetidos)
    comprobar("xl/metadata.xml" in nombres_partes
              and sum(x.count(' cm="1"') for x in hojas.values()) >= 3,
              "UNIQUE/FILTER siguen siendo formulas dinamicas")
    comprobar(sum(x.count("<dataValidation ") for x in hojas.values()) > 15,
              "las listas desplegables (incluidas las que leen de Opciones) se estiran")
    return destino, inf


def probar_con_excel(destino: Path, inf: InformeFitogreen):
    print("\n[4] Con Excel real: cero errores y el EBITDA por dos caminos")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        flujo = load_workbook(destino)[FLUJO]
    fila = next(c.row for c in flujo["A"] if c.value == "EBITDA USD")
    lecturas = "\n".join("    Write-Output ($fl.Range('{}{}').Value2)".format(c, fila)
                         for c in "FGHIJKLM")
    cuerpo = r"""
    $libro = $excel.Workbooks.Open('RUTA', 0, $true)
    $excel.CalculateFullRebuild()
    $errores = 0
    foreach ($ws in $libro.Worksheets) {
      $celdas = $null
      try { $celdas = $ws.UsedRange.SpecialCells(-4123, 16) } catch { $celdas = $null }
      if ($celdas -ne $null) { $errores += $celdas.Count }
    }
    Write-Output ('ERRORES ' + $errores)
    $fl = $libro.Worksheets.Item('Flujo proyectado ')
LECTURAS
    Write-Output 'FIN'
""".replace("RUTA", str(destino.resolve()).replace("'", "''")).replace("LECTURAS", lecturas)
    guion = ("$ErrorActionPreference = 'Stop'\n$excel = $null\n$libro = $null\ntry {\n"
             + GUION_ARRANQUE_IMPRESORA_SEGURA
             + "    $excel = New-Object -ComObject Excel.Application\n"
             + "    $excel.Visible = $false\n    $excel.DisplayAlerts = $false\n" + cuerpo
             + "\n} finally {\n    if ($libro -ne $null) { $libro.Close($false) }\n"
             + "    if ($excel -ne $null) { $excel.Quit() }\n"
             + GUION_LIBERAR_COM_EXCEL + GUION_RESTAURAR_IMPRESORA + "}\n")
    try:
        with tempfile.TemporaryDirectory() as tmp:
            archivo = Path(tmp) / "revisar.ps1"
            archivo.write_text(guion, encoding="utf-8-sig")
            salida = subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                 "-File", str(archivo)], capture_output=True, text=True, timeout=300).stdout
    except Exception as e:  # noqa: BLE001
        print("  (sin Excel en esta maquina: {})".format(e))
        return
    if "FIN" not in salida:
        print("  (Excel no respondio; se omite)")
        return
    lineas = salida.split()
    errores = int(lineas[lineas.index("ERRORES") + 1])
    comprobar(errores == 0, "cero celdas en error en todo el libro", errores)
    valores = [float(x.replace(",", ".")) for x in lineas[lineas.index("ERRORES") + 2:][:8]]
    python = ebitda_en_python(inf)
    comprobar(len(valores) == 8 and all(abs(a - b) < 0.01 for a, b in zip(valores, python)),
              "el EBITDA de los 8 anios: Excel igual a Python", list(zip(valores, python))[:2])


def main():
    probar_disposicion()
    probar_supuestos()
    with tempfile.TemporaryDirectory() as tmp:
        destino, inf = probar_generacion(Path(tmp))
        probar_con_excel(destino, inf)
    print()
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  -", f)
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
