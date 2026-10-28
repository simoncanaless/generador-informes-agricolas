"""Prueba del calculo de requerimiento hidrico.

El contraste fuerte esta en [1]: la planilla propia de Fitogreen trae un
ejemplo ya resuelto (hoja 'Hoja4') con dos especies, y el codigo tiene que
reproducir esas cifras al decimal. Si eso calza, la formula es la de ellos y
no una interpretacion mia.

Uso: python scripts/test_riego.py
"""

import glob
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config.settings import ajustes as config  # noqa: E402
from schemas.informe import Plantacion  # noqa: E402
from services.riego import (  # noqa: E402
    balance_de_predio, cargar_parametros, demanda_por_hectarea,
)

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)


def casi(a, b, tol=0.0001):
    return a is not None and abs(float(a) - float(b)) < tol


def parametros():
    return cargar_parametros(config.raiz / "config" / "parametros_riego.json")


def probar_formula():
    print("\n[1] La formula contra el ejemplo resuelto de la propia planilla")
    # Hoja4 del archivo de riego: ETo 7,02 para las dos filas.
    #   Palto      Kc 0,80  microaspersion 0,80 -> 0,8125    L/s/ha
    #   Mandarinas Kc 0,75  goteo          0,90 -> 0,677083  L/s/ha
    comprobar(casi(demanda_por_hectarea(7.02, 0.80, 0.80), 0.8125),
              "palto: 7,02 x 0,80 x 10 / 0,80 / 86,4 = 0,8125 L/s/ha",
              str(demanda_por_hectarea(7.02, 0.80, 0.80)))
    comprobar(casi(demanda_por_hectarea(7.02, 0.75, 0.90), 0.677083),
              "mandarinas con goteo = 0,677083 L/s/ha",
              str(demanda_por_hectarea(7.02, 0.75, 0.90)))
    # Y los totales del mismo ejemplo: 60 ha y 35 ha.
    comprobar(casi(demanda_por_hectarea(7.02, 0.80, 0.80) * 60, 48.75),
              "60 ha de palto piden 48,75 L/s, igual que la planilla")
    comprobar(casi(demanda_por_hectarea(7.02, 0.75, 0.90) * 35, 23.697917),
              "35 ha de mandarinas piden 23,70 L/s")

    # Un sistema mas eficiente pide menos agua para lo mismo.
    goteo = demanda_por_hectarea(6.33, 0.8, 0.9)
    surco = demanda_por_hectarea(6.33, 0.8, 0.6)
    comprobar(surco > goteo,
              "el surco (0,6) exige mas caudal que el goteo (0,9) a igualdad "
              "de cultivo")
    comprobar(casi(surco / goteo, 0.9 / 0.6),
              "y la razon es exactamente la de sus eficiencias")

    # La cobertura escala lineal: un huerto de 2 anios pide el 40%.
    comprobar(casi(demanda_por_hectarea(6.33, 0.8, 0.9, cobertura=0.4),
                   goteo * 0.4),
              "la cobertura escala la demanda de forma proporcional")

    try:
        demanda_por_hectarea(6.33, 0.8, 0)
        comprobar(False, "una eficiencia de 0 deberia ser rechazada")
    except ValueError:
        comprobar(True, "rechaza una eficiencia de 0 en vez de dividir por cero")


def probar_biblioteca():
    print("\n[2] La biblioteca extraida de la planilla")
    p = parametros()

    comprobar(len(p.eto_por_comuna) >= 120,
              "trae la ETo de 123 comunas", str(len(p.eto_por_comuna)))
    comprobar(len(p.kc_por_especie) >= 30,
              "y el Kc de 33 especies", str(len(p.kc_por_especie)))
    comprobar(casi(p.eto("Melipilla"), 6.33),
              "ETo de Melipilla = 6,33, la misma que usa la hoja MANANTIALES",
              str(p.eto("Melipilla")))
    comprobar(casi(p.eto("SAN FELIPE"), 7.02) and casi(p.eto("san felipe"), 7.02),
              "la busqueda no depende de mayusculas")
    comprobar(casi(p.eto("Vicuna"), p.eto("Vicuña")),
              "ni de las tildes: 'Vicuna' encuentra 'Vicuña'")
    comprobar(p.eto("Ninguna Parte") is None,
              "una comuna que no esta devuelve None, no un valor por defecto")

    comprobar(casi(p.kc("Palto"), 0.8), "Kc del palto")
    comprobar(casi(p.kc("Palto (rebajado)"), 0.8),
              "el calificativo no cambia el Kc: un palto rebajado bebe igual")
    comprobar(casi(p.kc("Palto - clonal"), 0.8), "ni el 'clonal'")
    comprobar(casi(p.kc("Avellano"), 0.9),
              "el avellano SI tiene Kc aunque no tenga curva de produccion")

    # Un catastro dice 'Nogales' o 'Limonero' y la tabla, 'nogal' y 'limon'.
    # La curva si calzaba -el flujo los proyectaba- y el Kc no: en el mismo
    # informe esas hectareas pedian 0 L/s (verificacion del 2026-09-11).
    nogal, nogales = p.kc("Nogal"), p.kc("Nogales")
    comprobar(nogal is not None and nogales is not None and casi(nogales, nogal),
              "el plural encuentra su Kc: 'Nogales' -> 'nogal'", str(nogales))
    limon, limonero = p.kc("Limon"), p.kc("Limonero")
    comprobar(limon is not None and limonero is not None and casi(limonero, limon),
              "el arbol encuentra el Kc de su fruto: 'Limonero' -> 'limon'",
              str(limonero))
    cerezos = p.kc("Cerezos")
    comprobar(cerezos is not None and casi(cerezos, p.kc("Cerezo")),
              "y 'Cerezos', como lo escribe el catalogo de curvas", str(cerezos))
    comprobar(p.kc("Vid") is None,
              "'Vid' sigue sin Kc: la tabla separa mesa y vinifera y no se adivina",
              str(p.kc("Vid")))

    comprobar(casi(p.eficiencia("Goteo"), 0.9), "goteo 0,9")
    comprobar(casi(p.eficiencia("Surco"), 0.6), "surco 0,6")
    comprobar(p.region_de("Melipilla") == "RM",
              "ubica la comuna en su region", str(p.region_de("Melipilla")))

    # La tabla de cobertura es POR ANIO, no por estado de madurez. Corrige lo
    # que el proyecto suponia (1,0 adulta / 0,8 creciente / 0,6 formacion).
    comprobar(casi(p.cobertura(1), 0.1), "anio 1 -> 10% de cobertura")
    comprobar(casi(p.cobertura(2), 0.4), "anio 2 -> 40%")
    comprobar(casi(p.cobertura(3), 0.7), "anio 3 -> 70%")
    comprobar(casi(p.cobertura(4), 1.0), "anio 4 -> adulta")
    comprobar(casi(p.cobertura(25), 1.0), "y de ahi en adelante sigue adulta")
    comprobar(casi(p.cobertura(None), 1.0),
              "sin edad se asume adulta: es lo conservador para el banco")


def probar_balance():
    print("\n[3] Balance de un predio: demanda contra caudal disponible")
    p = parametros()
    cartera = [
        Plantacion(especie="Palto", superficie_ha=60.0, riego="Microaspersion",
                   anio_plantacion="2010"),
        Plantacion(especie="Mandarino", superficie_ha=35.0, riego="Goteo",
                   anio_plantacion="2012"),
    ]
    balance = balance_de_predio(
        cartera, p, comuna="San Felipe", caudal_disponible_l_s=100.0,
        anio_informe=2026)

    print("      demanda total {:.2f} L/s | disponible {:.0f} | deficit {:.2f}"
          .format(balance.demanda_total_l_s, balance.caudal_disponible_l_s,
                  balance.deficit_l_s))
    comprobar(casi(balance.demanda_total_l_s, 72.447917),
              "reproduce el total del ejemplo de la planilla: 72,45 L/s",
              str(balance.demanda_total_l_s))
    comprobar(balance.suficiente is True,
              "con 100 L/s disponibles, el agua alcanza")
    comprobar(balance.deficit_l_s < 0, "y el deficit sale negativo (sobra)")

    escaso = balance_de_predio(
        cartera, p, comuna="San Felipe", caudal_disponible_l_s=50.0,
        anio_informe=2026)
    comprobar(escaso.suficiente is False, "con 50 L/s ya no alcanza")
    comprobar(casi(escaso.deficit_pct, 30.98),
              "y cuantifica el deficit: 30,98% de la demanda sin cubrir",
              str(escaso.deficit_pct))

    sin_caudal = balance_de_predio(cartera, p, comuna="San Felipe",
                                   anio_informe=2026)
    comprobar(sin_caudal.suficiente is None,
              "sin declarar el caudal disponible no se opina sobre suficiencia")

    # Una plantacion joven pide menos: cobertura del anio 2.
    joven = balance_de_predio(
        [Plantacion(especie="Palto", superficie_ha=60.0,
                    riego="Microaspersion", anio_plantacion="2024")],
        p, comuna="San Felipe", anio_informe=2026)
    comprobar(casi(joven.plantaciones[0].cobertura, 0.4),
              "un huerto plantado hace 2 anios va con cobertura 0,4")
    comprobar(casi(joven.demanda_total_l_s, 48.75 * 0.4),
              "y por eso pide el 40% del agua de uno adulto",
              str(joven.demanda_total_l_s))


def probar_faltantes():
    print("\n[4] Lo que no se puede calcular se dice, no se inventa")
    p = parametros()

    sin_curva_conocida = balance_de_predio(
        [Plantacion(especie="Dragonfruit", superficie_ha=10.0, riego="Goteo")],
        p, comuna="Melipilla")
    detalle = sin_curva_conocida.plantaciones[0]
    comprobar(not detalle.calculable, "una especie sin Kc no se calcula")
    comprobar(detalle.demanda_total_l_s is None,
              "y no se le pone una demanda por defecto")
    comprobar(any("Dragonfruit" in a for a in sin_curva_conocida.avisos),
              "se avisa nombrando la especie", str(sin_curva_conocida.avisos))

    sin_sistema = balance_de_predio(
        [Plantacion(especie="Palto", superficie_ha=10.0)], p, comuna="Melipilla")
    comprobar(not sin_sistema.plantaciones[0].calculable,
              "sin sistema de riego tampoco se calcula")
    comprobar(any("sistema de riego" in a for a in sin_sistema.avisos),
              "y se dice cual es el dato que falta")

    sin_comuna = balance_de_predio(
        [Plantacion(especie="Palto", superficie_ha=10.0, riego="Goteo")], p)
    comprobar(any("ETo" in a for a in sin_comuna.avisos),
              "sin comuna no hay ETo, y se avisa una sola vez",
              str(sin_comuna.avisos))

    # Se puede pasar la ETo a mano cuando la comuna no esta en la tabla.
    a_mano = balance_de_predio(
        [Plantacion(especie="Palto", superficie_ha=10.0, riego="Goteo")],
        p, eto_mm_dia=6.0)
    comprobar(a_mano.plantaciones[0].calculable,
              "se puede declarar la ETo a mano si la comuna no esta")

    # Los parametros usados quedan a la vista para poder defenderlos.
    detalle = a_mano.plantaciones[0]
    comprobar(all(v is not None for v in
                  (detalle.eto_mm_dia, detalle.kc, detalle.eficiencia)),
              "cada demanda guarda los tres parametros con que se calculo")


def probar_cartera_real():
    print("\n[5] Cartera real de Cliente A")
    archivos = glob.glob(str(config.raiz / "RER Agric*.xlsx"))
    if not archivos:
        print("  (no esta el RER Cliente A en la raiz; se omite)")
        return

    from services.catastro_import import cargar_reglas, leer_catastro
    reglas = cargar_reglas(config.raiz / "config" / "reglas_agronomicas.json")
    catastro = leer_catastro(archivos[0], reglas)
    p = parametros()

    balance = balance_de_predio(
        catastro.plantaciones, p, comuna="Melipilla", anio_informe=2026)

    calculadas = [d for d in balance.plantaciones if d.calculable]
    print("      {}/{} plantaciones calculadas | {:.1f} ha | demanda {:.1f} L/s"
          .format(len(calculadas), len(balance.plantaciones),
                  balance.superficie_calculada_ha, balance.demanda_total_l_s))

    comprobar(len(calculadas) == len(balance.plantaciones),
              "todas las plantaciones del catastro real tienen sus parametros",
              "sin calcular: {}".format(
                  [d.especie for d in balance.plantaciones if not d.calculable]))
    comprobar(casi(balance.superficie_calculada_ha, catastro.superficie_total_ha,
                   tol=0.5),
              "y cubren la superficie completa del cliente",
              "{} vs {}".format(balance.superficie_calculada_ha,
                                catastro.superficie_total_ha))
    comprobar(balance.demanda_total_l_s > 0, "la demanda total es positiva")
    # El avellano es 83% de la superficie y no tiene curva de produccion, pero
    # si tiene Kc: el riego se calcula igual.
    avellano = [d for d in calculadas if "Avellano" in d.especie]
    comprobar(avellano and all(d.demanda_total_l_s > 0 for d in avellano),
              "el avellano se riega aunque no tenga curva de produccion",
              "{} plantaciones".format(len(avellano)))



def probar_volcado_a_excel():
    print("\n[6] Volcado a la plantilla Bsan")
    from schemas.informe import (
        Antecedentes, DerechoAgua, DotacionRecursos, InformeFitogreen,
    )
    from services.excel_engine import MotorExcel, cargar_mapa
    from services.riego import filas_de_requerimiento
    import openpyxl

    plantilla = config.raiz / "templates" / "Formato Fitogreen Bsan.xlsx"
    if not plantilla.exists():
        print("  (no esta la plantilla Bsan; se omite)")
        return

    p = parametros()
    cartera = [
        Plantacion(especie="Palto", variedad="Hass", superficie_ha=16.6,
                   riego="Microaspersion", anio_plantacion="2010"),
        Plantacion(especie="Cerezo", superficie_ha=8.3, riego="Goteo",
                   anio_plantacion="2024"),
    ]
    balance = balance_de_predio(cartera, p, comuna="Melipilla",
                                anio_informe=2026)
    filas = filas_de_requerimiento(balance, cartera)

    comprobar(len(filas) == 2, "una fila de tabla por plantacion")
    comprobar(casi(filas[0].eto_maximo_mm_enero, 6.33),
              "la fila lleva la ETo de la comuna, no un valor tecleado")
    comprobar(casi(filas[1].cobertura, 0.4),
              "y la cobertura del cerezo de 2 anios es 0,4",
              str(filas[1].cobertura))

    informe = InformeFitogreen(
        antecedentes=Antecedentes(cliente="Prueba Riego", rut="76.543.210-3"),
        dotacion_recursos=DotacionRecursos(
            requerimientos_riego=filas,
            derechos_agua=[DerechoAgua(predio="Predio", nombre_fuente="Canal",
                                       caudal_real_ls=40.0)]),
    )
    motor = MotorExcel(plantilla, cargar_mapa(
        config.raiz / "config" / "mapping_santander.json"))
    destino = config.ruta_salida_pruebas / "test_riego_bsan.xlsx"
    motor.generar(informe, destino)

    ws = openpyxl.load_workbook(destino)["Dotación de recursos"]

    comprobar(str(ws["N32"].value or "").startswith("Cobertura"),
              "la columna de cobertura queda rotulada: en el informe impreso "
              "no se ven las formulas", repr(ws["N32"].value))
    comprobar(casi(ws["N34"].value, 0.4),
              "y el 0,4 del cerezo joven se ve como dato, no escondido")

    # Lo derivado va como formula viva, igual que en el flujo.
    for celda, concepto in (("H33", "requerimiento por hectarea"),
                            ("I33", "requerimiento total")):
        valor = str(ws[celda].value or "")
        comprobar(valor.startswith("="),
                  "{} ({}) es formula, no un numero muerto".format(concepto, celda),
                  valor)
    comprobar("86.4" in str(ws["H33"].value),
              "la formula lleva la constante de conversion a la vista",
              str(ws["H33"].value))
    comprobar("N33" in str(ws["H33"].value),
              "y referencia la celda de cobertura, no un literal")
    comprobar(str(ws["K33"].value) == "=$H$27",
              "la disponibilidad se toma del total de derechos de agua",
              str(ws["K33"].value))
    comprobar(ws["J33"].value == "=SUM(I33:I36)",
              "la suma que ya traia la plantilla queda intacta",
              str(ws["J33"].value))


def main():
    print("=" * 72)
    print("PRUEBA DEL REQUERIMIENTO HIDRICO")
    print("=" * 72)

    probar_formula()
    probar_biblioteca()
    probar_balance()
    probar_faltantes()
    probar_cartera_real()
    probar_volcado_a_excel()

    print("\n" + "=" * 72)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
