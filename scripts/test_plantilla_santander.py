"""Prueba de templates/Formato Fitogreen Santander 2026.xlsx: el formato
vacio que scripts/derivar_plantilla_santander.py saca de un informe
Santander terminado.

No necesita el informe de origen (trae datos de un cliente y no esta en el
repositorio): comprueba que la plantilla sea un MOLDE -una vez cada parte que
se repite, sin datos escritos, sin años a mano, sin nada que dependa del
computador donde se armo- y que las temporadas salgan de su unica celda.

Dos niveles, como el resto del proyecto:
- Sin Excel: lee el paquete. Corre siempre.
- Con Excel real: escribe una temporada y recalcula. Best-effort.

Uso: python scripts/test_plantilla_santander.py
"""

import os
import re
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from config.settings import ajustes  # noqa: E402
from services import dibujos_xlsx as dx  # noqa: E402
from services.renderizado import (  # noqa: E402
    GUION_ARRANQUE_IMPRESORA_SEGURA,
    GUION_LIBERAR_COM_EXCEL,
    GUION_RESTAURAR_IMPRESORA,
)

PLANTILLA = ajustes.raiz / "templates" / "Formato Fitogreen Santander 2026.xlsx"
FLUJO = "Flujo proyectado "
HOJAS = ["Portada", "Recursos", "Inversión", "Plantaciones", FLUJO, "Supuestos",
         "Comentarios generales Opinion", "Anexo 1 Resultado", "Anexo 2 Ventas F29",
         "Anexo 3 Precios", "Anexo 4 Imágenes", "Opciones", "ETo",
         "Listas Validación", "Conciliación Ventas"]
OCULTAS = {"Opciones", "ETo", "Listas Validación", "Conciliación Ventas"}

fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)
    return condicion


def _libro():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(PLANTILLA)


def _texto(valor):
    return getattr(valor, "text", valor)


def probar_hojas(wb):
    print("\n[1] Las quince hojas del formato, en su orden")
    comprobar(wb.sheetnames == HOJAS, "mismas hojas y mismo orden", wb.sheetnames)
    ocultas = {ws.title for ws in wb.worksheets if ws.sheet_state != "visible"}
    comprobar(ocultas == OCULTAS, "las cuatro hojas de apoyo siguen ocultas", ocultas)


def probar_una_vez(wb):
    print("\n[2] Una sola vez cada parte que se repite")
    flujo = wb[FLUJO]
    bloques = [c.row for c in flujo["A"]
               if isinstance(c.value, str) and c.value.startswith("=Plantaciones!A")]
    comprobar(bloques == [8], "un bloque de plantacion en el flujo, en la fila 8", bloques)
    sup = wb["Supuestos"]
    filas = [c.row for c in sup["A"]
             if isinstance(c.value, str) and c.value.startswith("=Plantaciones!A")]
    comprobar(filas == [24], "una fila de plantacion en los supuestos", filas)
    comprobar(sup["B4"].value is None and sup["C5"].value is None,
              "una columna de sensibilizacion (la de la primera especie)")
    resumen = [c.row for c in flujo["A"] if c.row > 20 and c.row < 50
               and isinstance(c.value, str) and re.fullmatch(r"=A\d+", c.value)]
    # Nueve cuadros (siete indicadores, ingresos y costos): el primero lleva
    # el nombre de la especie y los otros ocho lo repiten.
    comprobar(resumen == [27, 29, 31, 33, 35, 37, 42, 44],
              "una fila de especie en cada cuadro del resumen y del EERR", resumen)
    portada = wb["Portada"]
    comprobar(_texto(portada["A22"].value) == "Total (ha)",
              "una fila de especie en la portada")
    rec = wb["Recursos"]
    comprobar(rec["L9"].value == "=SUM(L8:L8)", "un lote en los datos legales")
    comprobar(rec["A14"].value == "Total (ha)", "una unidad predial en la clasificacion")
    comprobar(rec["H25"].value == "=SUM(H24:H24)", "una fuente de agua")
    unidades = [c.row for c in rec["A"] if c.value == "Unidad Predial:"]
    comprobar(unidades == [33], "un bloque de requerimiento de riego", unidades)
    plan = wb["Plantaciones"]
    comprobar(plan["E10"].value == "=SUM(E7:E9)",
              "una plantacion en la lista (mas las dos filas en blanco del formato)")
    comprobar(plan["A16"].value == "Total ", "una especie en el resumen de superficie")
    comprobar(wb["Anexo 1 Resultado"]["E6"].value == "=SUM(E5:E5)",
              "una especie en el Anexo 1")


def probar_sin_datos(wb):
    print("\n[3] Sin datos escritos en las celdas que se llenan")
    vacias = {
        "Portada": ["B10", "H10", "B11", "H11", "B12", "H12", "B14", "H13", "H14"],
        FLUJO: ["D4", "F4", "E7", "A25"],
        "Supuestos": ["A22", "B10", "E24", "F24", "I24", "K24", "R24"],
        "Anexo 1 Resultado": ["F3", "A5", "B5", "E5", "H5"],
        "Recursos": ["A8", "E8", "L8", "A13", "B24"],
        "Plantaciones": ["A7", "E7", "B33", "B39"],
        "Comentarios generales Opinion": ["A5", "H4", "H29"],
    }
    llenas = [(h, c) for h, celdas in vacias.items() for c in celdas
              if wb[h][c].value not in (None, "")]
    comprobar(not llenas, "las celdas de datos estan vacias", llenas)

    numeros = []
    for ws in wb.worksheets:
        if ws.title in OCULTAS:
            continue
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) \
                        and c.value != 0:
                    numeros.append("{}!{}".format(ws.title, c.coordinate))
    comprobar(not numeros, "ningun numero escrito en las hojas visibles", numeros[:10])


def probar_temporadas(wb):
    print("\n[4] Las temporadas salen de UNA celda (CLAUDE.md §13)")
    flujo = wb[FLUJO]
    comprobar(all("$E$7" in str(flujo["{}7".format(c)].value) for c in "DFGHIJKLM"),
              "los rotulos de la fila 7 se cuentan desde E7")
    copias = [(f, c) for f in (23, 40, 51, 57, 64, 70) for c in "DFGHIJKLM"
              if flujo["{}{}".format(c, f)].value != "={}$7".format(c)]
    comprobar(not copias, "las otras seis filas de rotulos repiten la fila 7", copias[:5])

    anios = []
    for ws in wb.worksheets:
        if ws.title in OCULTAS:
            continue
        for fila in ws.iter_rows():
            for c in fila:
                v = _texto(c.value)
                if isinstance(v, str) and not v.startswith("=") \
                        and re.search(r"(19|20)\d\d", v) and (ws.title, c.coordinate) != ("Portada", "I7"):
                    anios.append("{}!{}={!r}".format(ws.title, c.coordinate, v[:30]))
    comprobar(not anios, "ningun año escrito a mano en las hojas visibles", anios[:10])


def probar_sumas(wb):
    print("\n[5] Sin sumas encadenadas a mano (CLAUDE.md §13)")
    cadenas = []
    for fila in wb[FLUJO].iter_rows():
        for c in fila:
            v = c.value
            if isinstance(v, str) and re.fullmatch(r"=\+?[A-Z]+\d+(\*[A-Z]+\d+)?"
                                                     r"(\+[A-Z]+\d+(\*[A-Z]+\d+)?){2,}", v):
                cadenas.append(c.coordinate)
    comprobar(not cadenas, "ninguna suma de tres o mas bloques escrita a mano", cadenas[:5])


def probar_paquete():
    print("\n[6] Nada que dependa del computador donde se armo")
    with zipfile.ZipFile(PLANTILLA) as z:
        nombres = z.namelist()
        hojas = {n: z.read(n).decode("utf-8") for n in nombres
                 if re.match(r"xl/worksheets/sheet\d+\.xml$", n)}
        libro = z.read("xl/workbook.xml").decode("utf-8")
    for prefijo, motivo in (("xl/externalLinks/", "sin vinculos a otros archivos"),
                            ("xl/comments", "sin comentarios de celda"),
                            ("xl/webextensions/", "sin complementos de Excel"),
                            ("docProps/custom.xml", "sin etiqueta de confidencialidad")):
        comprobar(not any(n.startswith(prefijo) for n in nombres), motivo)
    comprobar("externalReferences" not in libro, "el libro no declara referencias externas")
    comprobar(not any("Confidential" in x for x in hojas.values()),
              "sin la marca 'Confidential' en los encabezados")
    comprobar("xl/metadata.xml" in nombres and sum(x.count(' cm="1"') for x in hojas.values()) == 2,
              "las formulas de Excel 365 siguen siendo dinamicas (UNIQUE/FILTER)")
    comprobar(sum(x.count("<x14:dataValidation ") for x in hojas.values()) == 3,
              "las tres listas desplegables que leen de la hoja Opciones")


def probar_dibujos():
    print("\n[7] Imagenes y cuadros de texto")
    formas = dx.nombres_de_formas(PLANTILLA)
    anexo4 = [n for _, n in formas.get("Anexo 4 Imágenes", [])]
    comprobar(len(anexo4) == 15, "quince huecos de fotografia en el Anexo 4", len(anexo4))
    portada = [n for _, n in formas.get("Portada", [])]
    comprobar("Imagen 1" in portada and "Imagen 2" in portada,
              "la portada conserva el logo y el hueco de emplazamiento", portada)
    comprobar(not dx.texto_de_forma(PLANTILLA, "Recursos", "CuadroTexto 1"),
              "el cuadro de notas de suelos queda vacio")
    regla = dx.texto_de_forma(PLANTILLA, "Recursos", "CuadroTexto 2") or ""
    comprobar("derechos permanentes" in regla,
              "la regla del balance hidrico (del formato) se conserva")


def probar_con_excel():
    print("\n[8] Con Excel real: una temporada escrita mueve todos los rotulos")
    ruta = str(PLANTILLA.resolve()).replace("'", "''")
    cuerpo = r"""
    $excel.AskToUpdateLinks = $false
    $libro = $excel.Workbooks.Open('RUTA', 0, $true)
    $fl = $libro.Worksheets.Item('Flujo proyectado ')
    $fl.Range('E7').Value2 = '2025/26'
    $excel.CalculateFullRebuild()
    Write-Output ('D7=' + $fl.Range('D7').Text)
    Write-Output ('M7=' + $fl.Range('M7').Text)
    Write-Output ('E70=' + $fl.Range('E70').Text)
    Write-Output ('I26=' + $libro.Worksheets.Item('Portada').Range('I26').Text)
    Write-Output ('K23=' + $libro.Worksheets.Item('Supuestos').Range('K23').Text)
    Write-Output ('D3=' + $libro.Worksheets.Item('Anexo 1 Resultado').Range('D3').Text)
    Write-Output ('A11=' + $libro.Worksheets.Item('Anexo 2 Ventas F29').Range('A11').Text)
    Write-Output 'FIN'
""".replace("RUTA", ruta)
    guion = ("$ErrorActionPreference = 'Stop'\n$excel = $null\n$libro = $null\ntry {\n"
             + GUION_ARRANQUE_IMPRESORA_SEGURA
             + "    $excel = New-Object -ComObject Excel.Application\n"
             + "    $excel.Visible = $false\n    $excel.DisplayAlerts = $false\n" + cuerpo
             + "\n} finally {\n    if ($libro -ne $null) { $libro.Close($false) }\n"
             + "    if ($excel -ne $null) { $excel.Quit() }\n"
             + GUION_LIBERAR_COM_EXCEL + GUION_RESTAURAR_IMPRESORA + "}\n")
    try:
        with tempfile.TemporaryDirectory() as tmp:
            archivo = Path(tmp) / "temporadas.ps1"
            archivo.write_text(guion, encoding="utf-8-sig")
            salida = subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy",
                 "Bypass", "-File", str(archivo)],
                capture_output=True, text=True, timeout=240, encoding="utf-8",
                errors="replace").stdout
    except Exception as e:  # noqa: BLE001
        print("  (sin Excel en esta maquina: {})".format(e))
        return
    if "FIN" not in salida:
        print("  (Excel no respondio; se omite)")
        return
    leido = dict(l.split("=", 1) for l in salida.splitlines() if "=" in l)
    esperado = {"D7": "2024/25", "M7": "2033/34", "E70": "2025/26", "I26": "26/27",
                "K23": "2026/27", "D3": "2025/26", "A11": "Enero 2026"}
    for celda, valor in esperado.items():
        comprobar(leido.get(celda, "").strip() == valor,
                  "{} = {}".format(celda, valor), leido.get(celda))


def main():
    if not PLANTILLA.exists():
        print("No existe {}".format(PLANTILLA))
        return 1
    wb = _libro()
    probar_hojas(wb)
    probar_una_vez(wb)
    probar_sin_datos(wb)
    probar_temporadas(wb)
    probar_sumas(wb)
    probar_paquete()
    probar_dibujos()
    probar_con_excel()
    print()
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  -", f)
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
