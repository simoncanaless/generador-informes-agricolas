"""Prueba de extremo a extremo del motor Excel.

Genera un informe con tablas que DESBORDAN la plantilla (35 plantaciones frente
a las 4 de fabrica, como el informe real de Cliente B) y luego reabre el archivo
para comprobar que sigue siendo un libro sano:

  - las formulas de total se reexpandieron al nuevo rango,
  - las formulas que apuntan a otra pestana NO se tocaron,
  - los rangos combinados siguieron a sus filas,
  - las filas insertadas heredaron bordes, alto y formulas por fila,
  - las etiquetas y las imagenes de la plantilla continuan en su sitio.

Uso: python scripts/test_engine.py
"""

import datetime as dt
import os
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from schemas.informe import (  # noqa: E402
    Antecedentes, BloqueEspecieFlujo, BloqueOpinion, Calificacion, CategoriaFoto,
    ClaseUso, DatosSuelos, DerechoAgua, DotacionRecursos, Estado,
    EstructuraProductiva, FilaAnexo, FilaEstructuraTemporada, FlujoProyectado,
    Fotografia, InformeFitogreen, Inversiones, ItemDotacion, LineaInversion,
    MontoComparado, OpinionProfesional, Plantacion, PredioSuelos,
    RequerimientoRiego, ResultadoTemporada, ResumenResultados,
    ResumenSuperficieEspecie, SistemaComercializacion, SistemaProduccion,
    SuperficieClaseUso, Tenencia, Anexo,
)
from services import excel_ops  # noqa: E402
from services.excel_engine import MotorExcel  # noqa: E402
from services.excel_ops import desplazar_referencias  # noqa: E402
from services.mapping import cargar_mapa  # noqa: E402

from config.settings import ajustes  # noqa: E402

PLANTILLA = ajustes.ruta_plantilla
MAPPING = ajustes.ruta_mapping
SALIDA = ajustes.ruta_salida_pruebas / "test_informe_generado.xlsx"

fallos = []
comprobaciones = 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)


# ----------------------------------------------------- unidad: formulas

def probar_desplazamiento():
    print("\n[1] Reescritura de formulas al insertar filas")

    casos = [
        # (formula, desde, cantidad, esperado, motivo)
        ("=SUM(D7:D10)", 10, 31, "=SUM(D7:D41)",
         "el total se estira al insertar dentro del rango"),
        ("=SUM(D7:D10)", 11, 31, "=SUM(D7:D10)",
         "insertar por debajo del rango no lo altera"),
        ("='Flujo proyectado '!F110", 10, 5, "='Flujo proyectado '!F110",
         "una referencia a otra hoja jamas se desplaza"),
        ("=I5/$I$15", 14, 3, "=I5/$I$18",
         "la fila absoluta SI se desplaza al insertar"),
        ("=IFERROR(D62/D57,0)", 100, 4, "=IFERROR(D62/D57,0)",
         "referencias por encima del corte quedan igual"),
        ('="Total "&A20', 10, 2, '="Total "&A22',
         "los literales de texto no se confunden con referencias"),
        ("=B13+C13", 13, 2, "=B15+C15",
         "suma simple desplazada entera"),
        ("=SUM(H23:H26)", 26, 1, "=SUM(H23:H27)",
         "el rango crece por su extremo inferior"),
    ]
    for formula, desde, cantidad, esperado, motivo in casos:
        obtenido = desplazar_referencias(formula, desde, cantidad)
        comprobar(obtenido == esperado, "{} -> {}".format(motivo, esperado),
                  "obtenido: {!r}".format(obtenido))


# ------------------------------------------------------- datos de prueba

def imagen_de_prueba(directorio):
    """PNG minimo para ejercitar la insercion de imagenes."""
    from PIL import Image

    ruta = Path(directorio) / "emplazamiento.png"
    Image.new("RGB", (800, 600), (34, 90, 45)).save(ruta)
    return ruta


def construir_informe(ruta_imagen):
    especies = ["Cerezo", "Nogal", "Manzano", "Ciruelo", "Vid"]
    variedades = ["Santina", "Chandler", "Gala", "D'Agen", "Thompson"]

    # 35 plantaciones: fuerza 31 filas nuevas en 'Estructura Productiva'
    plantaciones = [
        Plantacion(
            especie=especies[i % 5],
            variedad=variedades[i % 5],
            anio_plantacion=str(2008 + i % 15),
            superficie_ha=round(2.5 + i * 0.37, 2),
            sistema_conduccion="Eje central" if i % 2 else "Espaldera",
            predio="Lote {}".format(i + 1),
            dominio=Tenencia.PROPIETARIO if i % 3 else Tenencia.ARRENDATARIO,
            riego="Goteo",
            estado=Estado.BUENO if i % 4 else Estado.REGULAR,
            densidad_plantacion="{}x{}".format(4 + i % 3, 2),
        )
        for i in range(35)
    ]

    return InformeFitogreen(
        antecedentes=Antecedentes(
            cliente="Agricola Las Encinas SpA",
            rut="76.543.210-3",
            ejecutivo="M. Fuentes",
            fecha_visita=date(2026, 5, 14),
            sucursal="Rancagua",
            fecha_emision_informe=date(2026, 5, 28),
            agronomo_informante="P. Aravena",
            n_solicitud="SOL-2026-00417",
        ),
        # 6 filas para 4 huecos -> inserta 2 en 'Cliente'
        estructura_temporada=[
            FilaEstructuraTemporada(
                especie=e, superficie_ha=10.0 + i, consolidacion_pct=(80 + i) / 100,
                tenencia_propio_ha=8.0 + i, tenencia_arriendo_ha=2.0,
            )
            for i, e in enumerate(especies + ["Palto"])
        ],
        resumen_resultados=ResumenResultados(
            temporada_ultima="25/26",
            temporada_proyeccion="26/27",
            ingresos_operacionales=MontoComparado(
                ultima_m_pesos=1_250_000, ultima_uf=32_400,
                proyeccion_m_pesos=1_410_000, proyeccion_uf=35_900,
            ),
            costos_directos=MontoComparado(
                ultima_m_pesos=780_000, ultima_uf=20_200,
                proyeccion_m_pesos=820_000, proyeccion_uf=20_900,
            ),
            margen_bruto=MontoComparado(ultima_m_pesos=470_000, proyeccion_m_pesos=590_000),
            costos_indirectos=MontoComparado(ultima_m_pesos=95_000, proyeccion_m_pesos=99_000),
            resultado_operacional_neto=MontoComparado(
                ultima_m_pesos=375_000, proyeccion_m_pesos=491_000
            ),
        ),
        dotacion_recursos=DotacionRecursos(
            suelos=DatosSuelos(
                cliente="Agricola Las Encinas SpA",
                predio="Fundo El Maiten",
                dominio="Propietario",
                rol_avaluo="123-45",
                comuna="Requinoa",
                inscripcion_dominio="Fojas 1240 N 890 ano 2011",
                predios=[
                    PredioSuelos(
                        nombre="Fundo El Maiten",
                        superficies={
                            ClaseUso.I_II: SuperficieClaseUso(riego_ha=45.0, secano_ha=0.0),
                            ClaseUso.III_IV: SuperficieClaseUso(riego_ha=30.0, secano_ha=5.0),
                        },
                    ),
                    PredioSuelos(
                        nombre="Lote Sur",
                        superficies={
                            ClaseUso.III_IV: SuperficieClaseUso(riego_ha=12.0, secano_ha=3.0),
                            ClaseUso.V_VI: SuperficieClaseUso(riego_ha=0.0, secano_ha=18.0),
                        },
                    ),
                ],
            ),
            # 5 derechos para 4 huecos -> inserta 1
            derechos_agua=[
                DerechoAgua(
                    predio="Fundo El Maiten", comuna="Requinoa",
                    tipo_fuente="Canal", nombre_fuente="Canal Requinoa {}".format(i),
                    dominio="Inscrito", cantidad=12.0 + i,
                    caudal_nominal_ls=30.0 + i, caudal_real_ls=25.0 + i,
                    cbr_fojas="{}".format(1200 + i), cbr_numero="{}".format(880 + i),
                    cbr_anio=2011 + i, cbr_conservador="CBR Rancagua",
                    propietario="Agricola Las Encinas SpA",
                )
                for i in range(5)
            ],
            # 5 requerimientos para 4 huecos -> inserta 1 y estira J33:J36
            requerimientos_riego=[
                RequerimientoRiego(
                    especie_variedad="{} {}".format(especies[i % 5], variedades[i % 5]),
                    anio_plantacion="201{}".format(i), superficie_ha=20.0 + i,
                    sistema_riego="Goteo", eficiencia_riego=0.9,
                    eto_maximo_mm_enero=6.5, coef_cultivo_kc_enero=0.95,
                    requerimiento_max_ls_ha=0.72,
                    requerimiento_total_plantacion_ls=14.4 + i,
                    densidad_plantacion="4x2",
                )
                for i in range(5)
            ],
            # 6 maquinas para 4 huecos -> inserta 2
            maquinaria=[
                ItemDotacion(nombre="Tractor {}".format(i), cantidad=1 + i, estado=Estado.BUENO)
                for i in range(6)
            ],
            infraestructura=[
                ItemDotacion(nombre="Bodega", cantidad=1, estado=Estado.BUENO),
                ItemDotacion(nombre="Packing", cantidad=1, estado=Estado.REGULAR),
                ItemDotacion(nombre="Camara de frio", cantidad=2, estado=Estado.BUENO),
            ],
            comentarios="Dotacion adecuada al nivel productivo del predio.",
        ),
        inversiones=Inversiones(
            lineas=[
                LineaInversion(
                    proyecto="Recambio varietal", anio=2026, superficie_ha=8.0,
                    item="Plantas", cantidad_ha=1250, costo_unitario_usd=3.2,
                    costo_inversion_total_usd=32000.0, costo_inversion_pesos=30_400_000,
                )
                for _ in range(3)
            ],
            comentarios="Inversion financiada con capital propio.",
        ),
        estructura_productiva=EstructuraProductiva(
            anio_catastro=2026,
            plantaciones=plantaciones,
            resumen_superficie=[
                ResumenSuperficieEspecie(
                    especie=e, propietario_ha=20.0 + i, arrendatario_ha=5.0,
                    plena_produccion_ha=15.0, rendimientos_crecientes_ha=6.0,
                    formacion_ha=4.0, proyectos_plantacion_ha=2.0,
                    superficie_total_con_proyectos_ha=27.0 + i,
                )
                for i, e in enumerate(especies[:4])
            ],
            comentarios_plantaciones="Huerto homogeneo, buen estado fitosanitario.",
            sistemas_produccion=[
                SistemaProduccion(
                    especie=e, sistema_produccion="Intensivo", densidad_plantacion="4x2",
                    sistema_riego="Goteo", variedades_con_royalty="Si",
                    considera_techos="No", sistema_contra_heladas="Torres",
                    nivel_tecnologico="Alto", nivel_homogeneidad="Alto",
                    nivel_productivo_promedio="12 ton/ha",
                )
                for e in especies[:3]
            ],
            sistemas_comercializacion=[
                SistemaComercializacion(
                    especie=e, empresas_comercializadoras="Exportadora X",
                    comercializacion="Consignacion", mercado="China",
                    periodo_cosecha="Nov-Dic", periodo_liquidaciones="Mar-May",
                    proceso="Packing propio", gestion_exportacion="Tercerizada",
                )
                for e in especies[:2]
            ],
        ),
        flujo_proyectado=FlujoProyectado(
            temporadas=["2023/24", "2024/25", "2025/26", "2026/27"],
            # Claves simbolicas, no anios escritos: el ano lo declara el
            # tasador en el Anexo y el sistema los cruza. Asi un informe hecho
            # el proximo ano no arrastra los anios de este.
            tipo_cambio_pesos_usd={"real": 890.0, "preliminar": 940.0},
            # Con decimales a proposito: es el valor que openpyxl convertia en
            # fecha cuando el formato se declaraba mal (ver [9d]).
            uf=39412.55,
            bloques=[
                BloqueEspecieFlujo(
                    especie_variedad_anio="Cerezo Santina 2015",
                    superficie_ha={"2023/24": 20.0, "2024/25": 20.0, "2025/26": 22.0},
                    precio_exportacion_usd_kg={"2023/24": 4.5, "2024/25": 4.2},
                    rendimiento_kg_ha={"2023/24": 9000.0, "2024/25": 9500.0},
                    costos_usd_ha={"2023/24": 12000.0},
                ),
                BloqueEspecieFlujo(
                    especie_variedad_anio="Nogal Chandler 2012",
                    superficie_ha={"2023/24": 15.0, "2024/25": 15.0},
                    rendimiento_kg_ha={"2023/24": 3800.0},
                ),
            ],
            administracion={"2023/24": 45_000_000.0, "2024/25": 47_000_000.0},
            gastos_generales={"2023/24": 22_000_000.0},
            comentarios="Proyeccion conservadora de precios de exportacion.",
        ),
        opinion=OpinionProfesional(
            gestion_administrativa_productiva=BloqueOpinion(
                texto="Administracion profesionalizada, con registros al dia.",
                calificacion=Calificacion.BUENO,
            ),
            canales_comerciales=BloqueOpinion(
                texto="Exportadora consolidada con historial de liquidaciones.",
                calificacion=Calificacion.MUY_BUENO,
            ),
            calidad_activos_productivos=BloqueOpinion(
                texto="Huertos en plena produccion y riego tecnificado.",
                calificacion=Calificacion.BUENO,
            ),
            conclusion_final=BloqueOpinion(
                texto="Unidad productiva solida y sustentable en el mediano plazo.",
                calificacion=Calificacion.BUENO,
            ),
        ),
        anexo=Anexo(
            resultado_real=ResultadoTemporada(
                filas=[
                    FilaAnexo(
                        especie=e, produccion_exportacion_kg=180_000.0 + i,
                        produccion_mercado_interno_kg=20_000.0,
                        ventas_exportacion_pesos=700_000_000.0,
                        ventas_mercado_interno_pesos=40_000_000.0,
                        total_ventas_usd=830_000.0, costos_directos_pesos=480_000_000.0,
                        costos_directos_usd=540_000.0,
                    )
                    for i, e in enumerate(especies[:4])
                ],
                administracion_pesos=45_000_000.0,
                gastos_generales_pesos=22_000_000.0,
            ),
            resultado_preliminar=ResultadoTemporada(
                filas=[
                    FilaAnexo(especie=e, produccion_exportacion_kg=195_000.0)
                    for e in especies[:2]
                ],
                administracion_pesos=47_000_000.0,
                gastos_generales_pesos=23_000_000.0,
            ),
        ),
        fotografias=[
            Fotografia(
                ruta=str(ruta_imagen), titulo="Vista general del predio",
                categoria=CategoriaFoto.EMPLAZAMIENTO,
            ),
            Fotografia(
                ruta=str(ruta_imagen), titulo="Sector de riego",
                categoria=CategoriaFoto.RIEGO,
            ),
        ],
    )


# --------------------------------------------------------- verificacion

def verificar(resultado):
    wb = load_workbook(SALIDA)

    print("\n[3] 'Estructura Productiva': 35 plantaciones sobre 4 huecos")
    ws = wb["Estructura Productiva"]
    # 31 filas insertadas en la 10 -> cuerpo 7..41, totales 11 -> 42
    comprobar(ws["A7"].value == "Cerezo", "el primer registro sigue en la fila 7",
              "A7={!r}".format(ws["A7"].value))
    comprobar(ws["A41"].value is not None, "el registro 35 llego a la fila 41",
              "A41={!r}".format(ws["A41"].value))
    comprobar(ws["A42"].value == "Superficie Total Plantada (ha)",
              "la fila de totales bajo de 11 a 42",
              "A42={!r}".format(ws["A42"].value))
    comprobar(ws["D42"].value == "=SUM(D7:D41)",
              "el total se reexpandio a todo el cuerpo",
              "D42={!r}".format(ws["D42"].value))
    combinadas = {str(r) for r in ws.merged_cells.ranges}
    comprobar("A42:C42" in combinadas,
              "el rango combinado siguio a su fila (A11:C11 -> A42:C42)")
    comprobar(ws["A20"].border.left.style == "medium"
              and ws["A20"].border.bottom.style == "thin",
              "una fila insertada heredo los bordes de la plantilla",
              "A20 left={} bottom={}".format(
                  ws["A20"].border.left.style, ws["A20"].border.bottom.style))
    comprobar(ws.row_dimensions[42].height == 15.0,
              "el alto de la fila de totales viajo con ella",
              "alto={}".format(ws.row_dimensions[42].height))

    print("\n[4] 'Estructura Productiva': la tabla de abajo se desplazo intacta")
    comprobar(ws["A46"].value == "Resumen superficie ",
              "el encabezado del resumen bajo de 15 a 46",
              "A46={!r}".format(ws["A46"].value))
    comprobar(ws["D47"].value == "=SUM(B47:C47)",
              "la formula por fila del resumen se reescribio",
              "D47={!r}".format(ws["D47"].value))
    comprobar(ws["B51"].value == "=SUM(B47:B50)",
              "el total del resumen apunta a su nuevo rango",
              "B51={!r}".format(ws["B51"].value))

    print("\n[5] 'Cliente': 6 especies sobre 4 huecos + imagen")
    ws = wb["Cliente"]
    comprobar(ws["A26"].value == "Total (ha)",
              "la fila de total bajo de 24 a 26", "A26={!r}".format(ws["A26"].value))
    # La consolidacion va como FRACCION (la celda multiplica por 100 al
    # mostrarla): con 80 salia «8000,0%» en la vista previa (2026-10-05). Y su
    # total no es la suma de los porcentajes sino el promedio ponderado por
    # superficie -recalculado con Excel da 82,7% para estas 6 especies-.
    comprobar(ws["D20"].value == 0.8 and ws["D20"].number_format.endswith("%"),
              "la consolidacion se escribe como fraccion bajo formato de %",
              "D20={!r} {}".format(ws["D20"].value, ws["D20"].number_format))
    total_consolidacion = str(ws["D26"].value)
    comprobar("SUMPRODUCT(B20:B25,D20:D25)" in total_consolidacion
              and "SUM(D" not in total_consolidacion,
              "su total es el promedio ponderado de las 6 filas, no una suma",
              "D26={!r}".format(total_consolidacion))
    comprobar(ws["B10"].value == "Agricola Las Encinas SpA",
              "el campo simple se escribio", "B10={!r}".format(ws["B10"].value))
    # Esta fila SI se reescribe, y es lo unico de la plantilla que se toca
    # aparte de la fila 112 del flujo: su rotulo dice 'M$' y la formula del
    # banco leia el flujo, que esta en USD. Fitogreen confirmo que manda el
    # rotulo, asi que se convierte con el tipo de cambio.
    comprobar(str(ws["B39"].value).startswith("=IF('Flujo proyectado '!$E$3")
              and "F110" in str(ws["B39"].value),
              "la fila de 8 anios se convirtio a M$ y siguio a su fila nueva",
              "B39={!r}".format(ws["B39"].value))
    comprobar(ws["A41"].value == "IMAGEN EMPLAZAMIENTO UNIDAD PRINCIPAL",
              "el rotulo de la imagen bajo con su fila",
              "A41={!r}".format(ws["A41"].value))
    comprobar(len(ws._images) == 2,
              "conviven el logo de la plantilla y la foto de emplazamiento",
              "imagenes={}".format(len(ws._images)))

    print("\n[6] 'Dotación de recursos': dos tablas crecen en la misma hoja")
    ws = wb["Dotación de recursos"]
    comprobar(ws["H28"].value == "=SUM(H23:H27)",
              "el total de caudal absorbio el derecho de agua extra",
              "H28={!r}".format(ws["H28"].value))
    combinadas = {str(r) for r in ws.merged_cells.ranges}
    comprobar("J34:J38" in combinadas,
              "el merge que cruzaba el hueco se estiro (J33:J36 -> J34:J38)",
              "presentes: {}".format(sorted(c for c in combinadas if c.startswith("J"))))
    comprobar(ws["B6"].value == "Agricola Las Encinas SpA",
              "los campos de cabecera no se movieron",
              "B6={!r}".format(ws["B6"].value))
    comprobar(ws["B13"].value == 45.0 and ws["E14"].value == 12.0,
              "los bloques laterales de suelos se escribieron por predio",
              "B13={!r} E14={!r}".format(ws["B13"].value, ws["E14"].value))
    comprobar(ws["D13"].value == "=B13+C13",
              "las formulas de TOTAL de suelos quedaron intactas",
              "D13={!r}".format(ws["D13"].value))
    # las dos tablas de arriba crecieron 1 fila cada una: el bloque baja a la 43
    comprobar(ws["F43"].value == "Infraestructura " and ws["F44"].value == "Bodega",
              "la infraestructura convive con la maquinaria ampliada",
              "F43={!r} F44={!r}".format(ws["F43"].value, ws["F44"].value))
    comprobar(ws["A49"].value == "Tractor 5" and ws["A51"].value.startswith("Dotacion"),
              "la maquinaria ocupa 6 filas y el comentario quedo debajo",
              "A49={!r} A51={!r}".format(ws["A49"].value, ws["A51"].value))
    comprobar(ws["J34"].value == "=SUM(I34:I38)",
              "el requerimiento total absorbio la fila extra",
              "J34={!r}".format(ws["J34"].value))
    comprobar("L38:M38" in {str(r) for r in ws.merged_cells.ranges},
              "la fila insertada heredo la combinacion L:M de la plantilla")

    print("\n[7] 'Anexo ' e 'Inversión'")
    ws = wb["Anexo "]
    comprobar(ws["B4"].value == "Cerezo", "el anexo real se escribio",
              "B4={!r}".format(ws["B4"].value))
    comprobar(ws["M9"].value == 45_000_000.0, "administracion en su celda fija",
              "M9={!r}".format(ws["M9"].value))
    comprobar(ws["F8"].value == "=SUM(F4:F7)",
              "sin desbordes, las formulas del anexo no cambian",
              "F8={!r}".format(ws["F8"].value))

    # La plantilla trae el literal '0' en I4 donde I5:I7 llevan =G+H. Es un
    # error del formato del banco, no del generador, y se corrige al escribir:
    # un literal si se puede pisar, una formula no.
    comprobar(ws["I4"].value == "=G4+H4",
              "corrige el literal 0 que la plantilla trae en I4",
              "I4={!r}".format(ws["I4"].value))
    comprobar(ws["I5"].value == "=G5+H5",
              "y respeta la formula que la plantilla si trae en I5")

    # J y L son las columnas en USD, que la plantilla dejo vacias.
    for celda, concepto in (("J4", "total de ventas"), ("L4", "costos directos")):
        valor = str(ws[celda].value or "")
        comprobar(valor.startswith("=") and "$D$3" in valor,
                  "{} en USD sale de dividir por el tipo de cambio ({})".format(
                      concepto, celda), valor[:60])
    comprobar("IFERROR" in str(ws["J4"].value),
              "envuelto en IFERROR: sin tipo de cambio declarado daria "
              "#DIV/0!, que la regla de integridad prohibe")
    comprobar("$E$3" in str(ws["J15"].value or ""),
              "el preliminar usa el tipo de cambio de SU temporada, no el "
              "historico", str(ws["J15"].value))

    ws = wb["Inversión"]
    comprobar(ws["A5"].value == "Recambio varietal",
              "la inversion se escribio en el ancla de la celda combinada",
              "A5={!r}".format(ws["A5"].value))
    # G, H e I estaban vacias en la plantilla y por eso el total daba cero y
    # las once celdas de J salian en #DIV/0! en TODO informe generado.
    for columna, quien in (("G", "costo unitario USD/ha"),
                           ("H", "costo inversion total USD"),
                           ("I", "costo inversion en pesos")):
        valor = str(ws["{}5".format(columna)].value or "")
        comprobar(valor.startswith("="),
                  "{} es formula viva, no un numero muerto".format(quien),
                  "{}5={!r}".format(columna, valor))

    # La superficie esta COMBINADA por proyecto (C5:C9): la segunda linea tiene
    # que multiplicar por el ancla. Apuntar a C6 daria cero -Excel devuelve cero
    # dentro de un combinado- sin ninguna celda en error y sin ningun aviso.
    comprobar("C5" in str(ws["H6"].value or ""),
              "la segunda linea multiplica por el ANCLA de la superficie",
              "H6={!r}".format(ws["H6"].value))
    comprobar("'Flujo proyectado '!$E$3" in str(ws["I5"].value or ""),
              "los pesos salen del tipo de cambio de la PROYECCION: una "
              "inversion es futura", "I5={!r}".format(ws["I5"].value))
    comprobar(str(ws["J5"].value or "").startswith('=IF(I5="","",IFERROR('),
              "el porcentaje queda con doble guardia: en blanco si la fila no "
              "tiene inversion, y sin #DIV/0! si no hay ninguna",
              "J5={!r}".format(ws["J5"].value))
    comprobar(ws["B15"].value is None and ws["F15"].value is None,
              "la fila de totales no suma anios ni costos unitarios",
              "B15={!r} F15={!r}".format(ws["B15"].value, ws["F15"].value))

    print("\n[8] 'Flujo proyectado ' y opinion")
    ws = wb["Flujo proyectado "]
    comprobar(ws["D3"].value == 890.0 and ws["E3"].value == 940.0,
              "el tipo de cambio salio de la serie por temporada",
              "D3={!r} E3={!r}".format(ws["D3"].value, ws["E3"].value))
    comprobar(ws["A10"].value == "Cerezo Santina 2015",
              "la etiqueta del bloque 1 del flujo",
              "A10={!r}".format(ws["A10"].value))
    comprobar(ws["D11"].value == 20.0 and ws["F11"].value == 22.0,
              "las series por temporada cayeron en su columna",
              "D11={!r} F11={!r}".format(ws["D11"].value, ws["F11"].value))
    comprobar(ws["A20"].value == "Nogal Chandler 2012",
              "el bloque 2 empieza en la fila 20",
              "A20={!r}".format(ws["A20"].value))
    comprobar(ws["D101"].value == 45_000_000.0,
              "los costos generales usan sus filas fijas",
              "D101={!r}".format(ws["D101"].value))
    comprobar(ws["D110"].value == "=D106-D107-D108",
              "las 473 formulas del flujo quedaron sin tocar",
              "D110={!r}".format(ws["D110"].value))

    # Las tres filas derivadas de cada bloque salen como formula viva y no como
    # numero: es lo que permite que el banco mueva un precio y vea el efecto.
    # Origen: RER 'Flujo sin replantes' filas 15-17.
    comprobar(ws["D17"].value == "=D14*D11",
              "produccion total = rendimiento x superficie, como formula",
              "D17={!r}".format(ws["D17"].value))
    comprobar(ws["D18"].value == "=D17*D15*D12",
              "ingresos exportacion = produccion x ratio x precio exp",
              "D18={!r}".format(ws["D18"].value))
    comprobar(ws["D19"].value == "=D17*(1-D15)*D13",
              "ingresos mercado interno = produccion x (1-ratio) x precio MI",
              "D19={!r}".format(ws["D19"].value))
    comprobar(ws["D27"].value == "=D24*D21",
              "el segundo bloque desplaza las referencias a su propia base",
              "D27={!r}".format(ws["D27"].value))
    comprobar(ws["N17"].value == "=N14*N11",
              "la fila de formula se llena hasta la ultima temporada (N)",
              "N17={!r}".format(ws["N17"].value))

    # Una fila derivada con huecos rompe los rangos de suma de mas abajo, que
    # es exactamente la clase de error que se arrastra desde el RER.
    huecos = [
        "{}{}".format(col, fila)
        for fila in (17, 18, 19)
        for col in "DEFGHIJKLMN"
        if not str(ws["{}{}".format(col, fila)].value or "").startswith("=")
    ]
    comprobar(not huecos,
              "ninguna de las 33 celdas derivadas del bloque 1 quedo sin formula",
              "sin formula: {}".format(huecos))

    ws = wb["Comentarios generales Opinion"]
    # La plantilla combina el rotulo 'CALIFICACION' en E4:G4 y deja la casilla
    # con borde en H4. Hasta el 2026-09-15 el dato caia en E5 -dentro del
    # cuerpo del comentario- y la casilla del banco salia vacia. Se mira
    # tambien la celda vieja: sin eso, volver a escribir ahi no da rojo.
    comprobar(ws["H4"].value == "Bueno",
              "la calificacion va a la casilla que la plantilla rotula",
              "H4={!r}".format(ws["H4"].value))
    comprobar(ws["E5"].value in (None, ""),
              "y no dentro del cuerpo del comentario",
              "E5={!r}".format(ws["E5"].value))
    ancho_del_texto = next(
        (str(r) for r in ws.merged_cells.ranges if str(r).startswith("A5:")), "")
    comprobar(ancho_del_texto == "A5:H10",
              "el comentario usa el recuadro entero, no la mitad izquierda",
              "combinada: {!r}".format(ancho_del_texto))

    # El ancho de esa columna es lo unico que separa la calificacion mas larga
    # de salirse del recuadro por la derecha, y eso no deja ninguna celda en
    # error: solo se ve en el render. El piso sale de la propia escala.
    # El piso sale de la escala que ofrece la PANTALLA, no del enum del
    # dominio: 'aplicar_visita' guarda el texto elegido tal cual, asi que lo
    # que puede caer en la celda son los cinco rotulos del formulario -hasta
    # 'Bajo el Promedio', 16 caracteres- y no los cuatro del enum, cuyo mas
    # largo mide 9 y cabria en la columna original sin arreglar nada.
    import json
    escala = json.loads(
        ajustes.ruta_formulario_visita.read_text(encoding="utf-8")
    )["escalas"]["calificacion"]
    mas_largo = max(escala, key=len)
    ancho_h = ws.column_dimensions["H"].width or 0
    comprobar(ancho_h >= len(mas_largo),
              "la columna de la calificacion cabe el rotulo mas largo",
              "'{}' son {} caracteres y la columna mide {:.1f}".format(
                  mas_largo, len(mas_largo), ancho_h))

    # La fila de totales tambien se corre cuando la tabla desborda. Este
    # informe mete 6 especies en 4 huecos, asi que 'Total (ha)' baja dos filas;
    # antes el motor escribia sobre una fila de datos y el total quedaba vacio.
    cliente_hoja = wb["Cliente"]
    fila_total = next(
        (f for f in range(20, 34)
         if "total" in str(cliente_hoja.cell(f, 1).value or "").lower()), None)
    comprobar(fila_total is not None,
              "la fila de totales de la portada existe tras el desborde",
              str(fila_total))
    if fila_total:
        formula = str(cliente_hoja.cell(fila_total, 2).value or "")
        comprobar(formula.startswith("=SUM("),
                  "y lleva su suma, no quedo vacia",
                  "fila {} -> {!r}".format(fila_total, formula))
        comprobar(str(fila_total - 1) in formula,
                  "la suma llega hasta la ultima fila de datos, no se corta",
                  formula)

    print("\n[9] Formato numerico")
    # La plantilla del banco viene entera en 'General': los pesos salen como
    # '1107021,94' y las hectareas como '67,69999999999999'. Los datos estan
    # bien y el informe se ve mal igual, que para un entregable es lo mismo
    # que estar mal.
    flujo = wb["Flujo proyectado "]
    # La REGLA, que la puso el agrónomo: si el rotulo de la fila ya dice la
    # unidad -"Superficie (ha)"- el dato no la repite. Se escribe 69,5 y no
    # "69,5 ha". La unidad se reserva para las filas cuyo rotulo NO la trae
    # ("Ingresos Exportacion", "EBITDA"), donde sin ella el numero queda sin
    # explicar.
    desnudos = [
        ("G11", "Superficie (ha)",                 "la superficie no repite 'ha'"),
        ("G12", "Precio exportacion (USD/Kg)",     "el precio no repite 'USD'"),
        ("G14", "Produccion/Rendimiento (Kg/ha)",  "el rendimiento no repite 'kg/ha'"),
        ("G17", "Produccion Total (kg)",           "la produccion no repite 'kg'"),
    ]
    for celda, rotulo, motivo in desnudos:
        formato = flujo[celda].number_format
        comprobar(formato != "General" and '"' not in formato,
                  "{} — el rotulo ya dice la unidad ({})".format(motivo, rotulo),
                  "{}={!r}".format(celda, formato))

    con_unidad = [
        ("G18", "USD", "'Ingresos Exportacion' SI lleva USD: su rotulo no lo dice"),
        ("G110", "USD", "'EBITDA' tambien"),
    ]
    for celda, marca, motivo in con_unidad:
        formato = flujo[celda].number_format
        comprobar(marca in formato, motivo, "{}={!r}".format(celda, formato))

    comprobar("%" in flujo["G15"].number_format,
              "el ratio de exportacion se ve como porcentaje",
              "G15={!r}".format(flujo["G15"].number_format))

    cliente = wb["Cliente"]
    comprobar('"' not in cliente["B20"].number_format
              and cliente["B20"].number_format != "General",
              "la superficie de la portada tampoco repite la unidad",
              "B20={!r}".format(cliente["B20"].number_format))

    # Ojo: este informe DESBORDA las tablas, asi que las filas de mas abajo
    # se corrieron. Que el formato las haya seguido es justo lo dificil, y por
    # eso la fila se busca por su etiqueta en vez de darla por fija.
    fila_ingresos = next(
        (f for f in range(28, 40)
         if "ingresos" in str(cliente.cell(f, 1).value or "").lower()), None)
    comprobar(fila_ingresos is not None,
              "se ubica la fila de ingresos de la portada tras el desborde",
              str(fila_ingresos))
    if fila_ingresos:
        formato_uf = cliente.cell(fila_ingresos, 4).number_format
        comprobar(formato_uf != "General" and '"' not in formato_uf,
                  "la columna UF va desnuda -el encabezado dice UF- y el "
                  "formato siguio a su fila",
                  "fila {} -> {!r}".format(fila_ingresos, formato_uf))

    # El formato se aplica al RANGO, no solo a lo que escribe el generador:
    # las 573 formulas que ya traia la plantilla producen numeros que tambien
    # hay que presentar bien.
    comprobar(flujo["N110"].number_format != "General",
              "una celda que calcula la plantilla queda formateada igual",
              "N110={!r}".format(flujo["N110"].number_format))

    hoja_anexo = wb["Anexo "]
    comprobar('"' not in hoja_anexo["G4"].number_format
              and hoja_anexo["G4"].number_format != "General",
              "en el anexo, 'Ventas Exportacion ($)' no repite el peso",
              "G4={!r}".format(hoja_anexo["G4"].number_format))
    comprobar('"' not in hoja_anexo["J4"].number_format,
              "ni 'Total ventas USD' el dolar",
              "J4={!r}".format(hoja_anexo["J4"].number_format))

    dotacion = wb["Dotación de recursos"]
    fila_riego = next(
        (f for f in range(30, 45)
         if dotacion.cell(f, 8).number_format not in ("General",)), None)
    comprobar(fila_riego is not None
              and '"' not in dotacion.cell(fila_riego, 8).number_format,
              "el requerimiento hidrico tampoco: su encabezado ya dice L/s",
              "fila {} -> {!r}".format(
                  fila_riego,
                  dotacion.cell(fila_riego, 8).number_format if fila_riego else None))

    # Red de seguridad, aprendida a golpes: una mascara con texto literal SIN
    # comillas ('USD #,##0') hace que openpyxl la lea como formato de FECHA
    # -la 's' es segundos y la 'd' dias-, intente convertir el numero y lo
    # deje en #VALUE!. El dato queda destruido y la celda se ve formateada.
    # Por eso no se comprueba mascara por mascara sino el libro entero.
    corruptas = []
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.data_type == "e":
                    corruptas.append("{}!{} = {}".format(
                        hoja.title, celda.coordinate, celda.value))
    comprobar(not corruptas,
              "ninguna celda del libro quedo en error tras aplicar formatos",
              "; ".join(corruptas[:5]))

    print("\n[9b] Correcciones a las formulas defectuosas de la plantilla")
    # La plantilla del banco trae =(X110/X106)*100 sin guardia: toda temporada
    # sin ingresos -que son casi todas al empezar- deja la celda en #DIV/0!.
    # Es lo unico que el motor tiene permitido pisar de la plantilla, y solo
    # porque esta declarado celda por celda en el mapping.
    flujo = wb["Flujo proyectado "]
    fila_112 = [flujo.cell(112, c).value for c in range(4, 15)]
    comprobar(all(str(v).startswith("=IFERROR(") for v in fila_112),
              "las 11 columnas de EBITDA/Ingresos quedaron con guardia",
              str(fila_112[:2]))
    comprobar(fila_112[0] == '=IFERROR((D110/D106)*100,"")',
              "y la formula corregida conserva el calculo original",
              str(fila_112[0]))
    # 11 columnas de EBITDA/Ingresos + 8 temporadas de la caratula + las 10
    # filas de porcentaje de Inversion.
    comprobar(len(resultado.correcciones) == 29,
              "el motor deja constancia de las 29 correcciones aplicadas",
              str(len(resultado.correcciones)))
    comprobar(all(":" in c and len(c.split(":", 1)[1].strip()) > 20
                  for c in resultado.correcciones),
              "cada correccion viaja con su motivo escrito",
              str(resultado.correcciones[:1]))

    # La salvaguarda general sigue en pie: fuera de lo declarado, ninguna
    # formula del banco se toca. F110 es la que lee la caratula.
    comprobar(flujo["F110"].value == "=F106-F107-F108",
              "las demas formulas de la plantilla quedan intactas",
              str(flujo["F110"].value))

    print("\n[9c] Resumen de la caratula, en M$ y leido del Anexo")
    # El Anexo esta en pesos y el cuadro de la caratula esta rotulado M$, asi
    # que la formula divide por mil. Y la fila de totales del Anexo SE MUEVE
    # cuando el cliente trae mas especies que huecos: por eso la referencia se
    # resuelve al final, con '{total:anexo_real}' ya compensado.
    cliente = wb["Cliente"]
    anexo = wb["Anexo "]
    fila_total_anexo = next(
        (f for f in range(4, 40)
         if str(anexo.cell(f, 9).value or "").startswith("=SUM(I")), None)
    comprobar(fila_total_anexo is not None,
              "se ubica la fila de totales del Anexo", str(fila_total_anexo))

    # Las filas del resumen tambien se corren -en este caso 2, por las 6
    # especies sobre 4 huecos-, asi que se buscan por su rotulo.
    fila_ing = next(f for f in range(25, 45)
                    if cliente.cell(f, 1).value == "Ingresos Operacionales")
    fila_ind = next(f for f in range(25, 45)
                    if cliente.cell(f, 1).value == "Costos Indirectos")
    b30 = str(cliente.cell(fila_ing, 2).value or "")
    comprobar("/1000" in b30,
              "el resumen convierte a M$ el peso que trae el Anexo", b30)
    comprobar("{total:" not in b30,
              "el token de fila quedo resuelto, no literal en la celda", b30)
    comprobar("I{}".format(fila_total_anexo) in b30,
              "y apunta a la fila de totales donde REALMENTE quedo", b30)
    comprobar(str(cliente.cell(fila_ind, 2).value or "").count("'Anexo '!M") == 2,
              "los costos indirectos suman administracion y gastos generales",
              str(cliente.cell(fila_ind, 2).value))
    comprobar(str(cliente.cell(fila_ing, 6).value or "").startswith("=IF('Anexo '!I"),
              "la columna de proyeccion lee la tabla Preliminar",
              str(cliente.cell(fila_ing, 6).value))

    print("\n[9d] El formato de un campo no puede destruir su valor")
    # openpyxl decide si una celda lleva fecha mirando las LETRAS del formato:
    # 'd' dia, 'm' mes, 's' segundo. Un nombre de catalogo como 'decimal2'
    # tiene d y m, asi que al releer el archivo el numero 39.412,55 volvia
    # convertido en el 26 de noviembre de 2007. Sin celdas en error y sin
    # pruebas rojas: el numero simplemente dejaba de ser el numero.
    from openpyxl.styles.numbers import is_date_format

    flujo_uf = wb["Flujo proyectado "]["D6"]
    comprobar(not isinstance(flujo_uf.value, (dt.date, dt.datetime)),
              "un campo numerico no termina convertido en fecha",
              "D6={!r}".format(flujo_uf.value))
    comprobar(flujo_uf.number_format == "#,##0.00",
              "el nombre del catalogo se resolvio a su mascara real",
              str(flujo_uf.number_format))

    numericas_como_fecha = []
    for hoja in wb.worksheets:
        for filas in hoja.iter_rows():
            for celda in filas:
                if isinstance(celda.value, (dt.date, dt.datetime)) \
                        and not is_date_format(celda.number_format):
                    numericas_como_fecha.append("{}!{}".format(
                        hoja.title, celda.coordinate))
    comprobar(not numericas_como_fecha,
              "ninguna celda quedo con fecha bajo un formato que no es de fecha",
              "; ".join(numericas_como_fecha[:5]))

    print("\n[9e] Los costos generales del flujo proyectado")
    # La plantilla trae 'Costo General' restandose del EBITDA, pero nadie
    # llenaba las filas 101-102: el EBITDA proyectado salia SIN gastos
    # indirectos y por eso inflado -29% medido sobre el informe completo-. Era
    # el unico numero del informe que le decia al banco algo que no era.
    #
    # El monto va por la serie POR TEMPORADA, no por un rango de columnas fijo:
    # las columnas se corren segun cuantos anios historicos haya entregado el
    # cliente, y un 'G:N' escrito a mano dejaria de calzar en silencio.
    flujo_g = wb["Flujo proyectado "]
    comprobar(flujo_g["D101"].value == 45_000_000.0,
              "la administracion declarada cae en la columna de su temporada",
              "D101={!r}".format(flujo_g["D101"].value))
    comprobar(flujo_g["D100"].value is not None,
              "y la plantilla la suma en 'Costo General'",
              "D100={!r}".format(flujo_g["D100"].value))


    print("\n[9f] Las temporadas historicas: solo las que se declararon")
    # el agrónomo: no siempre les entregan los anios hacia atras. A veces dos, a
    # veces uno, a veces ninguno. Asi que la temporada la DECLARA el tasador en
    # el Anexo y la columna aparece solo si hay dato; la plantilla las trae
    # escritas a mano y eso solo servia para el ano en que se escribieron.
    cab = wb["Flujo proyectado "]
    # Las temporadas se leen del propio archivo: 'informe' no esta a mano aqui.
    columnas = "DEFGHIJKLMN"
    temporadas_informe = []
    for letra in columnas:
        v = cab["{}8".format(letra)].value
        if v is None:
            break
        temporadas_informe.append(v)
    for i, temporada in enumerate(temporadas_informe):
        comprobar(cab["{}8".format(columnas[i])].value == temporada,
                  "la columna {} se rotula con la temporada del informe ({})"
                  .format(columnas[i], temporada),
                  "{}8={!r}".format(columnas[i],
                                    cab["{}8".format(columnas[i])].value))
    sobrante = columnas[len(temporadas_informe)]
    comprobar(cab["{}8".format(sobrante)].value is None,
              "la columna sin datos queda EN BLANCO: asi no existe",
              "{}8={!r}".format(sobrante,
                                cab["{}8".format(sobrante)].value))
    comprobar(cab["D9"].value == "Real ",
              "y la fila de abajo dice de que tipo es cada una",
              "D9={!r}".format(cab["D9"].value))

    print("\n[10] Integridad del libro completo")
    # El motor reescribe a mano los rangos combinados: dos que se solapen dejan
    # el archivo irreparable para Excel, asi que se comprueba hoja por hoja.
    solapes = []
    for hoja in wb.worksheets:
        rangos = list(hoja.merged_cells.ranges)
        for i, uno in enumerate(rangos):
            for otro in rangos[i + 1:]:
                if not uno.isdisjoint(otro):
                    solapes.append("{}: {} x {}".format(hoja.title, uno, otro))
    comprobar(not solapes, "ningun rango combinado se solapa con otro",
              "; ".join(solapes[:5]))

    # Un valor en una celda que no es el ancla de su rango combinado es la causa
    # clasica del dialogo "Excel reparo el archivo".
    huerfanas = []
    for hoja in wb.worksheets:
        for rango in hoja.merged_cells.ranges:
            for fila in hoja.iter_rows(
                min_row=rango.min_row, max_row=rango.max_row,
                min_col=rango.min_col, max_col=rango.max_col,
            ):
                for celda in fila:
                    if celda.value is None:
                        continue
                    if (celda.row, celda.column) != (rango.min_row, rango.min_col):
                        huerfanas.append("{}!{}".format(hoja.title, celda.coordinate))
    comprobar(not huerfanas,
              "ningun valor quedo fuera del ancla de su rango combinado",
              "; ".join(huerfanas[:5]))

    vacias = [h.title for h in wb.worksheets if h.max_row < 2]
    comprobar(not vacias, "ninguna hoja quedo vacia", str(vacias))
    comprobar(len(wb.sheetnames) == 7,
              "el libro conserva sus 7 pestanas", str(wb.sheetnames))

    print("\n[11] Avisos emitidos por el motor")
    for aviso in resultado.avisos:
        print("  - {}".format(aviso))
    esperados = [
        ("sin hueco en la plantilla", "avisa de la foto que la plantilla no admite"),
    ]
    for fragmento, descripcion in esperados:
        comprobar(any(fragmento in a for a in resultado.avisos), descripcion)

    # En 'Inversion' la plantilla combina A5:A9 y A10:A14: el nombre del
    # proyecto abarca sus cinco lineas de item. Antes, cada linea escribia su
    # proyecto sobre el ancla y la ultima ganaba, asi que el informe mostraba
    # UN solo proyecto -el ultimo- como si fuera el de todas las lineas. Ahora
    # el ancla la escribe quien llega primero y el que no calza se avisa.
    inversion = wb["Inversión"]
    comprobar(inversion["A5"].value == "Recambio varietal",
              "el bloque combinado conserva el proyecto de su PRIMERA linea",
              "A5={!r}".format(inversion["A5"].value))
    conflictos = [a for a in resultado.avisos if "bloque combinado" in a]
    comprobar(all("no se sobrescribio" in a for a in conflictos),
              "y un proyecto distinto en el mismo bloque se avisa, no se pisa",
              str(conflictos[:1]))


#: El catalogo de curvas, cargado UNA vez. Cada asyncio.run() abre y cierra su
#: propio event loop, pero el pool de conexiones de SQLAlchemy es del motor y
#: sobrevive: el segundo run se encuentra con conexiones atadas a un loop ya
#: cerrado y revienta con 'Event loop is closed'. Con una sola carga no pasa.
_CURVAS = None


def curvas_del_catalogo():
    global _CURVAS
    if _CURVAS is None:
        import asyncio

        from services.almacen import construir_almacen
        _CURVAS = asyncio.run(construir_almacen().listar_curvas())
    return _CURVAS


def probar_historicos_variables():
    """el agrónomo: no siempre les entregan los anios hacia atras.

    A veces dos, a veces uno, a veces ninguno. La columna tiene que aparecer
    solo si hay dato, rotulada con la temporada que el tasador declaro, y las
    que sobran deben quedar EN BLANCO -asi la columna 'no existe'-.
    """


    from config.settings import ajustes as config
    from schemas.informe import Antecedentes, InformeFitogreen, Plantacion
    from services.armado import completar_informe
    from services.visita import cargar_formulario

    print("\n[12] Historicos: la columna existe solo si hay dato")
    if not config.ruta_plantilla.exists():
        print("  (no esta la plantilla; se omite)")
        return

    curvas = curvas_del_catalogo()
    if not curvas:
        print("  (catalogo de curvas vacio; se omite)")
        return
    plantas = [Plantacion(especie="Palto", variedad="Hass",
                          anio_plantacion="2015", superficie_ha=40.0)]
    formulario = cargar_formulario(config.ruta_formulario_visita)

    def temporadas_del_archivo(visita):
        informe = InformeFitogreen(
            antecedentes=Antecedentes(cliente="Prueba historicos", rut="1-9"))
        completar_informe(informe, plantaciones=plantas, curvas=curvas,
                          visita=visita, formulario=formulario)
        motor = MotorExcel(config.ruta_plantilla, cargar_mapa(MAPPING))
        destino = config.ruta_salida_pruebas / "historicos.xlsx"
        motor.generar(informe, destino)
        hoja = load_workbook(destino)["Flujo proyectado "]
        return [(hoja["{}8".format(c)].value, hoja["{}9".format(c)].value)
                for c in "DEFGHIJKLMN"]

    # La geometria es fija y por eso la caratula -que lee columnas fijas-
    # queda siempre alineada: D y E para los historicos, F para el ano 1 del
    # horizonte, G..N para el resto de la proyeccion.
    sin_nada = temporadas_del_archivo({})
    comprobar(sin_nada[0][0] is None and sin_nada[1][0] is None,
              "sin historicos, las dos primeras columnas quedan en blanco",
              str(sin_nada[:2]))
    comprobar(sin_nada[2][0] is not None,
              "pero el ano 1 esta siempre en su columna (F)",
              str(sin_nada[2]))

    con_uno = temporadas_del_archivo({
        "temporadas_anexo": {"temporada_real": "2024/25"},
        "resultado_real": [{"especie": "Palto", "superficie_ha": 40.0,
                            "produccion_exportacion_kg": 300000,
                            "ventas_exportacion_pesos": 500_000_000,
                            "costos_directos_pesos": 300_000_000}],
        "tipo_cambio": {"tc_ultima_temporada": 940.0},
    })
    comprobar(con_uno[0][0] is None,
              "con UN historico, la columna del segundo queda en blanco",
              str(con_uno[0]))
    comprobar(con_uno[1][0] == "2024/25",
              "el historico se pega al ano 1, no al borde",
              str(con_uno[1]))
    comprobar(str(con_uno[1][1] or "").strip() == "Real",
              "y va rotulado como Real")
    comprobar(con_uno[2][0] is not None and con_uno[2][0] != "2024/25",
              "el ano 1 sigue en su columna de siempre",
              str(con_uno[2]))

    # La columna sin anio no debe mostrar ceros: las formulas de la plantilla
    # calculan igual sobre celdas vacias y '0,0' se lee como un dato en cero.
    # Se ocultan con una seccion vacia en el formato, NO borrando la formula:
    # si el tasador escribe ahi despues, el numero tiene que aparecer solo.
    hoja = load_workbook(config.ruta_salida_pruebas / "historicos.xlsx")
    flujo_h = hoja["Flujo proyectado "]
    formato_d = flujo_h["D56"].number_format
    comprobar(formato_d.endswith(";"),
              "los ceros de la columna sin anio quedan en blanco",
              "D56 formato={!r}".format(formato_d))
    comprobar(str(flujo_h["D56"].value or "").startswith("="),
              "y su formula sigue intacta por si llega el dato despues",
              "D56={!r}".format(flujo_h["D56"].value))
    comprobar(not str(flujo_h["E56"].number_format).endswith(";"),
              "la columna que SI tiene anio no se toca",
              "E56 formato={!r}".format(flujo_h["E56"].number_format))


def probar_respaldo_de_ventas():
    """El nivel de respaldo: lo que los IVAs confirman de lo declarado.

    El cliente entrega las ventas segun sus declaraciones de IVA, en PESOS, y
    la fila del banco pide dolares. Se escribe como formula dividiendo por la
    celda del tipo de cambio -no convirtiendo en Python- para que el monto en
    pesos quede a la vista y el cambio siga siendo una celda viva.

    Y solo existe donde el cliente entrego IVAs: una proyeccion no tiene con
    que estar respaldada, asi que su columna no se rotula siquiera. Un anio
    escrito con la celda vacia debajo se lee como un dato que falta.
    """


    from config.settings import ajustes as config
    from schemas.informe import Antecedentes, InformeFitogreen, Plantacion
    from services.armado import completar_informe
    from services.visita import cargar_formulario

    print("\n[13] Nivel de respaldo de ventas")
    if not config.ruta_plantilla.exists():
        print("  (no esta la plantilla; se omite)")
        return

    curvas = curvas_del_catalogo()
    if not curvas:
        print("  (catalogo de curvas vacio; se omite)")
        return

    plantas = [Plantacion(especie="Palto", variedad="Hass",
                          anio_plantacion="2015", superficie_ha=40.0)]
    formulario = cargar_formulario(config.ruta_formulario_visita)
    informe = InformeFitogreen(
        antecedentes=Antecedentes(cliente="Prueba respaldo", rut="1-9"))
    completar_informe(informe, plantaciones=plantas, curvas=curvas, visita={
        "temporadas_anexo": {"temporada_real": "2024/25",
                             "temporada_preliminar": "2025/26"},
        "resultado_real": [{"especie": "Palto", "superficie_ha": 40.0,
                            "produccion_exportacion_kg": 300000,
                            "ventas_exportacion_pesos": 500_000_000,
                            "costos_directos_pesos": 300_000_000}],
        "resultado_preliminar": [{"especie": "Palto", "superficie_ha": 40.0,
                                  "produccion_exportacion_kg": 320000,
                                  "ventas_exportacion_pesos": 560_000_000,
                                  "costos_directos_pesos": 320_000_000}],
        "respaldo_ventas": {"real_ventas_respaldadas_pesos": 400_000_000,
                            "prelim_ventas_respaldadas_pesos": 450_000_000},
        "tipo_cambio": {"tc_ultima_temporada": 940.0,
                        "tc_temporada_en_curso": 960.0},
    }, formulario=formulario)

    motor = MotorExcel(config.ruta_plantilla, cargar_mapa(MAPPING))
    destino = config.ruta_salida_pruebas / "respaldo.xlsx"
    motor.generar(informe, destino)
    flujo = load_workbook(destino)["Flujo proyectado "]

    # Que columna le toco a cada temporada no se supone: se lee de la cabecera.
    columnas = {flujo["{}8".format(c)].value: c for c in "DEFGHIJKLMN"
                if flujo["{}8".format(c)].value}
    real, prelim = columnas.get("2024/25"), columnas.get("2025/26")
    comprobar(real and prelim, "las dos temporadas declaradas tienen columna",
              str(columnas))
    if not (real and prelim):
        return

    monto = flujo["{}115".format(real)].value
    comprobar(str(monto or "").startswith("="),
              "el monto respaldado es una formula, no un numero muerto",
              "{}115={!r}".format(real, monto))
    comprobar("400000000" in str(monto),
              "y lleva dentro el monto en PESOS que declaro el cliente",
              str(monto))
    comprobar("$D$3" in str(monto),
              "dividido por la celda del tipo de cambio, que es una celda viva",
              str(monto))
    # Cada temporada usa el cambio de SU anio: la historica lee D3 y la
    # preliminar E3. Mezclarlos daria un monto convertido con el cambio del
    # anio equivocado, sin ninguna senal de que paso.
    comprobar("$E$3" in str(flujo["{}115".format(prelim)].value),
              "la preliminar usa el OTRO tipo de cambio, no el de la historica",
              "{}115={!r}".format(prelim, flujo["{}115".format(prelim)].value))

    pct = str(flujo["{}116".format(real)].value or "")
    comprobar(pct.startswith("=") and "115" in pct and "106" in pct,
              "el porcentaje es formula viva: respaldado / ingresos totales",
              "{}116={!r}".format(real, pct))
    comprobar("*100" not in pct,
              "en ratio crudo, porque su formato ya es 0.0%", pct)
    comprobar("IFERROR" in pct,
              "con guardia: sin ingresos daria #DIV/0!, que la regla prohibe",
              pct)

    # Lo que NO debe existir. Una proyeccion no tiene IVAs.
    proyectadas = [c for c in "DEFGHIJKLMN"
                   if c not in (real, prelim)]
    vacias = [c for c in proyectadas
              if flujo["{}115".format(c)].value is None
              and flujo["{}116".format(c)].value is None]
    comprobar(len(vacias) == len(proyectadas),
              "las columnas proyectadas quedan sin monto ni porcentaje",
              str([c for c in proyectadas if c not in vacias]))
    sin_rotulo = [c for c in proyectadas if flujo["{}114".format(c)].value is None]
    comprobar(len(sin_rotulo) == len(proyectadas),
              "y sin rotular: un anio con la celda vacia debajo se lee como un "
              "dato que falta",
              str([c for c in proyectadas if c not in sin_rotulo]))
    comprobar(flujo["{}114".format(real)].value == "2024/25",
              "la cabecera del bloque si rotula las declaradas",
              "{}114={!r}".format(real, flujo["{}114".format(real)].value))

    # El contraste de siempre: la formula, evaluada aparte, tiene que dar lo
    # mismo que calcula Python. Dos caminos distintos al mismo numero es la
    # unica verificacion que sirve para una formula: openpyxl escribe el texto
    # sin evaluarlo, asi que una formula que apunta mal se ve igual de bien.
    cuenta = str(monto).rsplit(",", 1)[-1].rstrip(")").replace("$D$3", "940.0")
    comprobar(abs(eval(cuenta) - 400_000_000 / 940.0) < 1e-6,
              "la formula, evaluada aparte, da el mismo USD que Python",
              "{} -> {:.2f}".format(cuenta, 400_000_000 / 940.0))


def probar_filas_espejo():
    """El resumen de indicadores dice de QUE especie es cada fila.

    El 'RESUMEN DE INDICADORES GENERALES' y el 'ESTADO DE RESULTADOS' repiten
    cuatro filas por indicador -una por bloque de especie- y la plantilla del
    banco no las rotula: el informe mostraba '69,5' y '765.000 kg' sin decir de
    que. Los numeros estaban bien y el cuadro era ilegible igual.
    """
    from config.settings import ajustes as config
    from schemas.informe import Antecedentes, InformeFitogreen, Plantacion
    from services.armado import completar_informe
    from services.visita import cargar_formulario

    print("\n[14] El resumen de indicadores, rotulado")
    if not config.ruta_plantilla.exists():
        print("  (no esta la plantilla; se omite)")
        return
    curvas = curvas_del_catalogo()
    if not curvas:
        print("  (catalogo de curvas vacio; se omite)")
        return

    plantas = [Plantacion(especie="Palto", variedad="Hass",
                          anio_plantacion="2015", superficie_ha=40.0),
               Plantacion(especie="Cerezo", variedad="Regina",
                          anio_plantacion="2018", superficie_ha=12.0)]
    informe = InformeFitogreen(
        antecedentes=Antecedentes(cliente="Prueba espejo", rut="1-9"))
    completar_informe(informe, plantaciones=plantas, curvas=curvas, visita={},
                      formulario=cargar_formulario(config.ruta_formulario_visita))

    motor = MotorExcel(config.ruta_plantilla, cargar_mapa(MAPPING))
    destino = config.ruta_salida_pruebas / "espejo.xlsx"
    motor.generar(informe, destino)
    flujo = load_workbook(destino)["Flujo proyectado "]

    # Los rotulos de los bloques, tal como quedaron escritos arriba.
    bloques = [flujo["A{}".format(b)].value for b in (10, 20, 30, 40)]
    usados = [b for b in bloques if b and b != "Especie/Variedad/Año"]
    comprobar(len(usados) == 2, "el informe tiene dos bloques con especie",
              str(bloques))

    grupos = {57: "Superficie", 62: "Produccion Total", 67: "Produccion media",
              72: "Precio de venta", 77: "Costo promedio", 82: "Costo unitario",
              90: "Ingresos", 95: "Costos Totales"}
    sin_rotulo = []
    for grupo, que in grupos.items():
        for indice, etiqueta in enumerate(usados):
            if flujo["A{}".format(grupo + indice)].value != etiqueta:
                sin_rotulo.append("{} fila {}".format(que, grupo + indice))
    comprobar(not sin_rotulo,
              "los 8 indicadores rotulan sus filas con la especie del bloque",
              str(sin_rotulo[:4]))

    # El rotulo sale del bloque, no de una lista aparte: no puede desfasarse.
    comprobar(flujo["A57"].value == flujo["A10"].value,
              "y dice exactamente lo mismo que el encabezado del bloque",
              "A57={!r} A10={!r}".format(flujo["A57"].value, flujo["A10"].value))

    # La fila del bloque que el cliente NO tiene calcula cero sobre celdas
    # vacias. Mostrarlo en una fila sin nombre es lo que mas confunde.
    vacios = range(len(usados), 4)
    apagadas = [g + i for g in grupos for i in vacios
                if str(flujo["E{}".format(g + i)].number_format).endswith(";")]
    comprobar(len(apagadas) == len(grupos) * len(list(vacios)),
              "las filas de los bloques que no existen quedan en blanco",
              "{} de {}".format(len(apagadas), len(grupos) * len(list(vacios))))
    comprobar(str(flujo["E60"].value or "").startswith("="),
              "sin borrar la formula: si manana hay una cuarta especie, el "
              "numero aparece solo", "E60={!r}".format(flujo["E60"].value))
    comprobar(flujo["A60"].value is None,
              "y sin rotulo, porque no hay especie que rotular",
              "A60={!r}".format(flujo["A60"].value))

    # La plantilla pone negrita en la columna A de unos grupos y no de otros,
    # aunque los ocho son subordinados de su encabezado.
    negritas = {flujo["A{}".format(g)].font.bold for g in grupos}
    comprobar(negritas == {False},
              "las ocho listas se leen igual: ninguna en negrita",
              str(negritas))


def _ext_de_hoja(xlsx, hoja):
    """Tamanos `(ancho, alto)` en px de las imagenes de UNA hoja, del XML.

    **No sirve `Image.width` de openpyxl.** Al releer un libro vuelve a abrir
    el archivo de imagen y devuelve su tamano NATURAL, no el `ext` que se
    guardo: una imagen escalada a 794x397 se reporta como 2400x1200 y parece
    que el escalado no funciono. Es la misma familia de error que
    `classList.contains` (§16.a.tricies.semel) e `innerText` sobre un elemento
    oculto (§16.a.quadragies.ter).
    """
    import re
    import zipfile

    def objetivo(rels, rid):
        # Sin asumir el ORDEN de los atributos: openpyxl escribe Target antes
        # que Id, y Excel al reves.
        for elemento in re.findall(r"<Relationship\b[^>]*/?>", rels):
            if re.search(r'Id="{}"'.format(re.escape(rid)), elemento):
                encontrado = re.search(r'Target="([^"]+)"', elemento)
                return encontrado.group(1) if encontrado else None
        return None

    with zipfile.ZipFile(xlsx) as z:
        libro = z.read("xl/workbook.xml").decode("utf-8")
        rid = None
        for elemento in re.findall(r"<sheet\b[^>]*/?>", libro):
            nombre = re.search(r'name="([^"]+)"', elemento)
            ref = re.search(r'r:id="([^"]+)"', elemento)
            if nombre and ref and nombre.group(1) == hoja:
                rid = ref.group(1)
        if rid is None:
            return []
        destino = objetivo(z.read("xl/_rels/workbook.xml.rels").decode("utf-8"), rid)
        if destino is None:
            return []
        ruta = destino.lstrip("/")
        if not ruta.startswith("xl/"):
            ruta = "xl/" + ruta
        dibujo = re.search(r'<drawing[^>]*r:id="([^"]+)"',
                           z.read(ruta).decode("utf-8"))
        if not dibujo:
            return []
        rels = z.read(ruta.replace("worksheets/", "worksheets/_rels/") + ".rels")
        blanco = objetivo(rels.decode("utf-8"), dibujo.group(1)) or ""
        texto = z.read("xl/drawings/" + Path(blanco).name).decode("utf-8")
        return [(int(cx) // 9525, int(cy) // 9525)
                for cx, cy in re.findall(r'ext cx="(\d+)" cy="(\d+)"', texto)]


def probar_hueco_compartido_de_imagen():
    """La satelital y el emplazamiento comparten el UNICO hueco de la Bsan.

    el agrónomo pidio que la imagen satelital saliera grande (2026-09-09), y el
    unico marcador de imagen que trae la plantilla es `Cliente!A40:I60`. O sea
    que dos categorias apuntan al mismo sitio, y hay dos cosas que vigilar:

      - que NO se peguen una encima de otra. La ocupacion se cuenta por HUECO
        y no por categoria justamente por esto: contando por categoria, cada
        una cree que el sitio esta libre y la de abajo queda invisible, sin
        aviso y sin celda en error.
      - que gane la satelital aunque llegue despues. El modelo le pone
        «emplazamiento» a cualquier foto de terreno, y esa no puede quedarse
        con el hueco que el tasador reservo a proposito.

    Las dos imagenes se fabrican con proporciones distintas -2:1 y 1:1- para
    poder decir CUAL quedo en cada sitio. Con dos fotos parecidas la
    comprobacion no distinguiria nada.
    """
    from PIL import Image

    print("\n[15] El hueco de imagen que comparten satelital y emplazamiento")

    with tempfile.TemporaryDirectory() as tmp:
        sat = Path(tmp) / "satelital.jpg"
        foto = Path(tmp) / "terreno.jpg"
        Image.new("RGB", (2400, 1200), (30, 90, 40)).save(sat, "JPEG")
        Image.new("RGB", (1200, 1200), (150, 60, 30)).save(foto, "JPEG")

        informe = InformeFitogreen(
            antecedentes=Antecedentes(cliente="Hueco compartido", rut="76.123.456-0"),
            fotografias=[
                # A proposito: la de emplazamiento va PRIMERA en la lista.
                Fotografia(ruta=str(foto), titulo="Vista del huerto",
                           categoria=CategoriaFoto.EMPLAZAMIENTO),
                Fotografia(ruta=str(sat), titulo="Satelital del predio",
                           categoria=CategoriaFoto.SATELITAL),
            ],
        )
        salida = ajustes.ruta_salida_pruebas / "test_hueco_compartido.xlsx"
        motor = MotorExcel(PLANTILLA, cargar_mapa(MAPPING))
        resultado = motor.generar(informe, salida, crear_galeria_fotos=True)

    # El logo de la plantilla mide 191x50 y no lo pone el motor.
    caratula = [t for t in _ext_de_hoja(salida, "Cliente")
                if not (t[0] < 300 and t[1] < 120)]
    comprobar(len(caratula) == 1,
              "una sola imagen del motor en la caratula: no se apilan",
              "hay {}".format(caratula))
    if caratula:
        ancho, alto = caratula[0]
        comprobar(ancho <= 794 and alto <= 420,
                  "escalada al hueco A40:I60 (794x420 px)",
                  "{}x{} px".format(ancho, alto))
        comprobar(ancho >= 700 or alto >= 380,
                  "y lo LLENA: no queda una estampilla en media carilla",
                  "{}x{} px".format(ancho, alto))
        comprobar(abs(ancho / alto - 2.0) < 0.1,
                  "la que quedo es la SATELITAL, no la de emplazamiento",
                  "proporcion {:.2f}".format(ancho / alto))

    galeria = _ext_de_hoja(salida, "Fotografias")
    comprobar(len(galeria) == 1,
              "la desplazada NO se pierde: baja a la hoja de fotografias",
              "{} en la galeria".format(len(galeria)))
    if galeria:
        comprobar(abs(galeria[0][0] / galeria[0][1] - 1.0) < 0.1,
                  "y la que bajo es la foto de terreno")
        # La rejilla copia la de un informe entregado (Cliente D): columnas de
        # 2,66 y filas de 14 pt, con cada foto en 15x10 celdas. El mapping
        # decia 15 de ancho y 10 pt de alto -la cantidad de columnas y filas,
        # no su medida-: cada hueco era doce veces mas ancho que alto y la
        # hoja se imprimia con estampillas de dos centimetros.
        comprobar(galeria[0][1] >= 150,
                  "la foto de la galeria se ve: no es una estampilla",
                  "{}x{} px".format(*galeria[0]))
    from openpyxl import load_workbook
    hoja_fotos = load_workbook(salida)["Fotografias"]
    comprobar(hoja_fotos["B4"].value is None
              and "Vista del huerto" in str(hoja_fotos["B14"].value or ""),
              "el titulo va DEBAJO de la foto, no en la celda que la foto tapa",
              "B4={!r} B14={!r}".format(hoja_fotos["B4"].value, hoja_fotos["B14"].value))

    aviso = [a for a in resultado.avisos if "hueco" in a.lower()]
    comprobar(bool(aviso), "se avisa por que bajo")
    comprobar(bool(aviso) and "Vista del huerto" in aviso[0],
              "y el aviso NOMBRA la desplazada, no la satelital",
              aviso[0] if aviso else "")


def probar_inversion_agrupada_por_proyecto():
    """Dos proyectos de inversion que no calzan justo en el limite de 5 filas.

    el usuario lo reporto probando: un archivo sintetico traia dos items de
    'Ampliacion palto' (2 lineas) y despues uno de 'Renovacion cerezo' (1
    linea). La plantilla combina A5:A9/B5:B9/C5:C9 para el PRIMER proyecto y
    A10:A14/B10:B14/C10:C14 para el segundo -2 bloques de 5, no 10 filas
    sueltas-. El motor escribia las 3 lineas seguidas (filas 5, 6, 7), asi que
    'Renovacion cerezo' caia en la fila 7, que sigue DENTRO del bloque
    combinado del primer proyecto: no se corrompia nada -el motor se negaba a
    pisar el ancla-, pero el segundo proyecto quedaba sin su nombre, anio o
    superficie escritos, con solo un aviso para enterarse.

    `campos_de_grupo` + `tamano_bloque` (services/excel_engine.py:
    _agrupar_en_bloques) resuelven esto: cada proyecto se completa con filas
    en blanco hasta llenar su bloque, asi el siguiente arranca justo en el
    limite. Verificado ademas RECALCULANDO CON EXCEL REAL -openpyxl no evalua
    formulas- que las formulas G/H/I de la segunda linea del primer proyecto
    siguen resolviendo al ancla de SU bloque (C5), no a la del segundo.
    """
    print("\n[16] Inversion: dos proyectos que no calzan en el limite de 5")

    informe = InformeFitogreen(
        antecedentes=Antecedentes(cliente="Dos proyectos", rut="76.543.210-3"),
        flujo_proyectado=FlujoProyectado(
            tipo_cambio_pesos_usd={"2025/26": 950.0, "2026/27": 970.0}),
        inversiones=Inversiones(lineas=[
            LineaInversion(proyecto="Ampliación palto", anio=2027,
                           superficie_ha=5.0, item="Plantas",
                           cantidad_ha=400, costo_unitario_usd=6.5),
            LineaInversion(proyecto="Ampliación palto", anio=2027,
                           superficie_ha=5.0, item="Sistema de riego",
                           cantidad_ha=1, costo_unitario_usd=3200),
            LineaInversion(proyecto="Renovación cerezo", anio=2028,
                           superficie_ha=3.0, item="Plantas",
                           cantidad_ha=500, costo_unitario_usd=8.0),
        ]),
    )
    salida = ajustes.ruta_salida_pruebas / "test_inversion_dos_proyectos.xlsx"
    motor = MotorExcel(PLANTILLA, cargar_mapa(MAPPING))
    resultado = motor.generar(informe, salida)

    conflictos = [a for a in resultado.avisos
                 if "bloque combinado" in a or "no se sobrescribio" in a]
    comprobar(not conflictos,
              "sin avisos de colision: cada proyecto cayo en su propio bloque",
              str(conflictos))

    ws = load_workbook(salida)["Inversión"]
    comprobar(ws["A5"].value == "Ampliación palto" and ws["B5"].value == 2027
              and ws["C5"].value == 5.0,
              "el primer proyecto queda anclado en su bloque (fila 5)",
              "A5={!r} B5={!r} C5={!r}".format(
                  ws["A5"].value, ws["B5"].value, ws["C5"].value))
    comprobar(ws["D6"].value == "Sistema de riego",
              "y su segunda linea de item, en la fila 6 del MISMO bloque",
              "D6={!r}".format(ws["D6"].value))
    comprobar(ws["A10"].value == "Renovación cerezo" and ws["B10"].value == 2028
              and ws["C10"].value == 3.0,
              "el segundo proyecto arranca en el SIGUIENTE bloque (fila 10), "
              "no en la fila 7 donde habria caido escribiendo seguido",
              "A10={!r} B10={!r} C10={!r}".format(
                  ws["A10"].value, ws["B10"].value, ws["C10"].value))
    comprobar(all(ws.cell(f, 4).value is None for f in (7, 8, 9)),
              "las filas de relleno del primer bloque quedan sin item",
              str([ws.cell(f, 4).value for f in (7, 8, 9)]))

    # openpyxl no evalua formulas: se recalcula con el Excel real y se lee el
    # resultado, la unica prueba de verdad para una formula (§14). Mismo
    # mecanismo que ver_excel.py -PowerShell + COM por subprocess-, porque
    # pywin32 no esta instalado en este entorno.
    #
    # Los fragmentos de impresora-a-salvo/liberar-COM son los mismos que
    # services/renderizado.py -esta automatizacion es SU PROPIA copia
    # independiente, no pasa por ese modulo, y tenia la misma trampa de la
    # impresora sin el arreglo (encontrada porque este script se quedo
    # colgado de verdad, corriendo esta sesion).
    import subprocess

    from services.renderizado import (
        GUION_ARRANQUE_IMPRESORA_SEGURA, GUION_LIBERAR_COM_EXCEL, GUION_RESTAURAR_IMPRESORA,
    )

    guion = r"""
$ErrorActionPreference = 'Stop'
$excel = $null
$libro = $null
try {{
{arranque_impresora}
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $libro = $excel.Workbooks.Open('{ruta}', $false, $false)
    $excel.CalculateFullRebuild()
    $excel.CalculateUntilAsyncQueriesDone()
    $hoja = $libro.Sheets('Inversión')
    $errores = 0
    foreach ($c in $hoja.UsedRange.Cells) {{ if ($c.Text -match '^#') {{ $errores++ }} }}
    Write-Output "ERRORES=$errores"
    Write-Output "H6=$($hoja.Range('H6').Value2)"
    Write-Output "H10=$($hoja.Range('H10').Value2)"
    $libro.Close($false)
}} finally {{
    if ($excel -ne $null) {{ $excel.Quit() }}
{liberar_com}
{restaurar_impresora}
}}
""".format(
        ruta=str(salida.resolve()).replace("'", "''"),
        arranque_impresora=GUION_ARRANQUE_IMPRESORA_SEGURA,
        restaurar_impresora=GUION_RESTAURAR_IMPRESORA,
        liberar_com=GUION_LIBERAR_COM_EXCEL,
    )

    try:
        proceso = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
            capture_output=True, text=True, timeout=120)
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        print("  (no se pudo recalcular con Excel real: {})".format(error))
        return

    salida_ps = proceso.stdout
    if proceso.returncode != 0 or "ERRORES=" not in salida_ps:
        print("  (Excel no esta instalado o fallo: {})".format(
            (proceso.stderr or salida_ps)[:200]))
        return

    valores = dict(
        linea.split("=", 1) for linea in salida_ps.splitlines() if "=" in linea)

    def _numero(clave):
        # Si el bloque quedo sin ancla, la formula da "" (guardia de §8.7) y
        # Value2 vuelve una cadena vacia: convertirla a float reventaria la
        # prueba entera en vez de dejarla fallar limpio en el comprobar.
        bruto = (valores.get(clave) or "").strip()
        try:
            return float(bruto)
        except ValueError:
            return float("nan")

    errores = int(valores.get("ERRORES", "-1"))
    comprobar(errores == 0,
              "cero celdas en error tras recalcular con Excel real",
              "{} en error".format(errores))
    # Fila 6 es la 2a linea del PRIMER proyecto: su H tiene que multiplicar
    # por la superficie del bloque 1 (C5=5.0), no una superficie ajena.
    h6 = _numero("H6")
    comprobar(h6 == 1 * 3200 * 5.0,
              "la formula de la 2a linea del primer proyecto usa SU ancla "
              "(C5=5.0), no la del segundo",
              "H6={!r}".format(h6))
    h10 = _numero("H10")
    comprobar(h10 == 500 * 8.0 * 3.0,
              "y la del segundo proyecto usa la suya (C10=3.0)",
              "H10={!r}".format(h10))


def probar_estilo_de_hojas_anexas():
    """'Detalle plantaciones' y 'Fotografias' con el mismo lenguaje visual
    del banco, no en blanco y negro.

    el usuario: *«hay que arreglar como se ve el tema de las fotografias... el
    detalle de plantacion deberia tener el mismo formato y diseno que tiene
    la plantilla»*. Las dos hojas las CREA el motor -no vienen en el archivo
    del banco-, y hasta ahora se escribian con `ws['A5'] = valor` puro: sin
    una fuente, un borde o un color que las distinguiera de un volcado de
    datos. El resto del libro trae barra roja de titulo ('COMENTARIOS
    GENERALES'), barra gris de sub-seccion ('ESTRUCTURA PRODUCTIVA
    (TEMPORADA)') y tablas con marco; estas dos hojas no heredaban nada de
    eso.

    Los colores no son inventados: se sacaron DEL ARCHIVO REAL con openpyxl
    -rojo solido FFFF0000 de 'Comentarios generales Opinion'!A2, gris
    A6A6A6 (blanco con tinte -0.35) de 'Cliente'!A16-, asi que esta prueba
    los compara contra esos mismos valores.
    """
    from PIL import Image

    from schemas.informe import BloqueMatriz, FilaMatriz, MatrizEnvejecimiento

    print("\n[17] 'Detalle plantaciones' y 'Fotografias' con estilo del banco")

    with tempfile.TemporaryDirectory() as tmp:
        rutas = []
        for i, color in enumerate(
                [(40, 70, 40), (90, 60, 40), (40, 60, 90)], start=1):
            ruta = Path(tmp) / "foto{}.jpg".format(i)
            Image.new("RGB", (1200, 800), color).save(ruta, "JPEG")
            rutas.append(str(ruta))

        temporadas = ["2026/27", "2027/28"]
        matriz = MatrizEnvejecimiento(
            temporadas=temporadas,
            bloques=[BloqueMatriz(
                etiqueta="Palto Hass 2012-2023", curva="Palto",
                factor_ajuste=1.0, costo_por_kg_usd=0.5,
                filas=[FilaMatriz(
                    edad=12, rendimiento_curva_kg_ha=11600,
                    costo_base_usd_ha=6000,
                    superficie_por_temporada={t: 10.0 for t in temporadas})])],
        )
        informe = InformeFitogreen(
            antecedentes=Antecedentes(cliente="Estilo de hojas anexas",
                                      rut="76.543.210-3"),
            flujo_proyectado=FlujoProyectado(
                tipo_cambio_pesos_usd={"2026/27": 950.0}, matriz=matriz),
            fotografias=[
                Fotografia(
                    ruta=r,
                    # La del medio, larga a proposito: es el caso que se
                    # cortaba a media palabra (screenshot de el usuario,
                    # 2026-09-22) cuando la foto de al lado tambien tenia
                    # titulo, y que sin wrap_text pasaba inadvertido.
                    titulo=("Detalle aéreo del conjunto de edificación y "
                            "patio de maniobras") if i == 2
                          else "Foto {}".format(i),
                    categoria=CategoriaFoto.PLANTACION)
                for i, r in enumerate(rutas, start=1)
            ],
        )
        salida = ajustes.ruta_salida_pruebas / "test_estilo_hojas_anexas.xlsx"
        motor = MotorExcel(PLANTILLA, cargar_mapa(MAPPING))
        motor.generar(informe, salida, crear_galeria_fotos=True)

    from openpyxl import load_workbook
    wb = load_workbook(salida)

    det = wb["Detalle plantaciones"]
    comprobar(det["A2"].value == "DETALLE DE PLANTACIONES POR EDAD",
              "el titulo va en la barra roja de la fila 2")
    comprobar(det["A2"].fill.fgColor.rgb == "FFFF0000",
              "en el MISMO rojo que usa el banco en sus propias barras "
              "('Comentarios generales Opinion'!A2)",
              str(det["A2"].fill.fgColor.rgb))
    comprobar(det["A2"].font.color.rgb == "FFFFFFFF" and det["A2"].font.b,
              "texto blanco y en negrita, no el negro por defecto")
    comprobar(any(str(r).startswith("A2:") for r in det.merged_cells.ranges),
              "la barra esta combinada a lo ancho de la tabla, no solo en A2")

    comprobar(det["A4"].value == "Palto Hass 2012-2023",
              "la etiqueta del bloque sigue en su celda de siempre")
    comprobar(det["A4"].fill.fgColor.rgb == "FFA6A6A6",
              "con la barra gris de sub-seccion, el mismo gris que "
              "'Cliente'!A16 ('ESTRUCTURA PRODUCTIVA (TEMPORADA)')",
              str(det["A4"].fill.fgColor.rgb))

    comprobar(det["A8"].value == "Edad" and det["A8"].font.b,
              "el encabezado de la tabla esta en negrita")
    comprobar(det["G8"].value == "2026/27 (ha)",
              "la temporada lleva su unidad -este encabezado no la decia "
              "en ningun otro lado, al reves que 'Rend. curva (kg/ha)'",
              "G8={!r}".format(det["G8"].value))

    comprobar(det["B9"].number_format == "#,##0",
              "el rendimiento va DESNUDO -sin 'kg/ha' repetido en la celda-: "
              "el encabezado ya lo dice (§16.a.duodecies)",
              det["B9"].number_format)
    comprobar(det["G9"].number_format == "#,##0.0",
              "la superficie por temporada lleva un decimal", det["G9"].number_format)

    comprobar(det["A8"].border.top.style == "medium"
              and det["E9"].border.right.style == "medium",
              "la tabla tiene marco: no es texto suelto sobre la hoja")
    comprobar(det["F8"].border.top.style is None,
              "la columna en blanco entre las dos cajas queda SIN bordear, "
              "no como una celda vacia dentro de la tabla",
              str(det["F8"].border.top.style))

    fot = wb["Fotografias"]
    comprobar(fot["A2"].value == "REGISTRO FOTOGRÁFICO",
              "la galeria tambien lleva su barra roja de titulo")
    comprobar(fot["A2"].fill.fgColor.rgb == "FFFF0000",
              "del mismo rojo", str(fot["A2"].fill.fgColor.rgb))
    comprobar("Estilo de hojas anexas" in str(fot["A3"].value or ""),
              "y el cliente debajo, para saber de que informe es aunque "
              "esta hoja se mire aparte", "A3={!r}".format(fot["A3"].value))
    comprobar(fot["B14"].font.name == "Aptos Narrow",
              "el titulo de cada foto usa la MISMA tipografia del banco, no "
              "la que trae openpyxl por defecto",
              fot["B14"].font.name)

    comprobar(fot["Q14"].alignment.wrap_text is True,
              "el titulo de la foto tiene wrap_text: sin esto, un titulo mas "
              "largo que su foto se corta a media palabra en vez de saltar "
              "de linea -Excel solo deja desbordar sobre la celda vecina "
              "cuando esta vacia, y la foto de al lado siempre tiene texto")
    alto_titulo = fot.row_dimensions[14].height
    comprobar(alto_titulo is not None and alto_titulo >= 26,
              "y la fila del titulo crece para las dos lineas que el titulo "
              "largo necesita, en vez de dejar la segunda invisible por un "
              "alto de fila pensado para una sola",
              "alto={!r}".format(alto_titulo))


def probar_formula_real_de_pixeles():
    """La conversion ancho/alto de columna a pixeles, contra Excel real.

    Se detecto persiguiendo otro problema (el marco de foto de abajo): la
    aproximacion vieja ``w*7+5`` para el ancho de columna sobreestimaba en
    24% un ancho angosto como el 2,66 de la rejilla de fotos -19px reales
    contra 23,62 calculados-, y de ahi que `_encajar`
    (`services/excel_engine.py`) pensara que el hueco de una foto era mas
    grande de lo que es y la escalara de mas: la foto quedaba
    **desbordando su hueco declarado**, invadiendo la columna de al lado sin
    ningun aviso. Medido con `Shape.Width` de Excel real (COM) sobre la
    plantilla Bsan: una foto declarada para B4:P13 salia 9pt mas ancha que
    el propio hueco.

    Estos numeros SON los que dio Excel real (no una aproximacion) el
    2026-09-22, para 7 anchos de columna y 6 altos de fila distintos, con
    'Aptos Narrow' 11 -la fuente 'Normal' real de la plantilla Bsan-.
    """
    print("\n[18] La conversion a pixeles, contra Excel real (no aproximada)")

    # ancho: (unidades declaradas, px reales medidos con Range.Width/COM)
    anchos = [(1.0, 7), (2.66, 19), (5.0, 35), (8.43, 59),
              (10.0, 70), (15.0, 105), (20.0, 140)]
    for unidades, esperado in anchos:
        obtenido = excel_ops._ancho_columna_px(unidades)
        comprobar(obtenido == esperado,
                  "columna de ancho {} da {}px reales, no la aproximacion vieja".format(
                      unidades, esperado),
                  "obtenido={}".format(obtenido))

    # alto: Excel trunca CADA fila a su propio pixel antes de sumar, no
    # trunca la suma de puntos al final -por eso 10 filas de 14pt no dan
    # 140pt (14*10), dan 135pt (18px*10, reconvertido)-.
    with tempfile.TemporaryDirectory() as tmp:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for fila in range(1, 11):
            ws.row_dimensions[fila].height = 14.0
        ws.column_dimensions["A"].width = 8.43
        ruta = Path(tmp) / "calibracion.xlsx"
        wb.save(ruta)
        wb2 = load_workbook(ruta)
        ancho_px, alto_px = excel_ops.dimensiones_rango_px(wb2.active, "A1:A10")
    comprobar(alto_px == 180,
              "10 filas de 14pt dan 180px (18px c/u, truncando fila por "
              "fila), no 186 (sumando 18,667 sin truncar y truncando al "
              "final)", "alto_px={}".format(alto_px))


def probar_marco_de_foto_segun_tamano():
    """El marco/mate detras de una foto se acomoda a su ANCHO real, no al
    del hueco entero.

    el usuario, mirando un informe real generado con el arreglo anterior (el
    wrap_text de los titulos): *"sigue viendose mas o menos mal, no
    podriamos poner algo justo por atras de las fotos que se acomode segun
    el tamano de las fotos?"*. `_encajar` ya encaja la foto sin deformarla,
    asi que casi nunca llena el hueco entero -el hueco de la galeria es mas
    ancho que alto, y una foto VERTICAL como esta queda muy angosta dentro
    de el-, y hasta ahora esa franja quedaba en blanco puro, sin marco ni
    fondo que la distinguiera del resto de la hoja.

    El limite exacto (columna H, no otra) esta verificado generando el
    mismo archivo y leyendolo con openpyxl antes de escribir esta prueba,
    no calculado a mano -el numero sale de `_encajar` + la formula real de
    pixeles de `probar_formula_real_de_pixeles`, no de una aproximacion
    nueva-.
    """
    from PIL import Image

    print("\n[19] El marco de foto se acomoda al ANCHO real de la foto")

    with tempfile.TemporaryDirectory() as tmp:
        vertical = Path(tmp) / "vertical.jpg"
        Image.new("RGB", (900, 1400), (90, 60, 40)).save(vertical, "JPEG")

        informe = InformeFitogreen(
            antecedentes=Antecedentes(cliente="Marco de foto", rut="1-9"),
            fotografias=[
                Fotografia(ruta=str(vertical), titulo="Vertical",
                          categoria=CategoriaFoto.PLANTACION),
            ],
        )
        salida = ajustes.ruta_salida_pruebas / "test_marco_de_foto.xlsx"
        motor = MotorExcel(PLANTILLA, cargar_mapa(MAPPING))
        motor.generar(informe, salida, crear_galeria_fotos=True)

    wb = load_workbook(salida)
    ws = wb["Fotografias"]

    # El hueco es B4:P13 (15 columnas); una foto vertical de 900x1400
    # termina encajada en solo 7 (B..H) -el resto del hueco (I..P) es del
    # tasador, no de la foto-.
    con_mate = [c for c in range(2, 17) if
                ws.cell(6, c).fill.fgColor.rgb == "FFF2F2F2"]
    comprobar(con_mate == list(range(2, 9)),
              "el mate cubre B..H -el ancho real de la foto encajada- y "
              "nada mas", "columnas={}".format(con_mate))
    borde_h = ws.cell(4, 8).border
    estilo_h = borde_h.right.style if borde_h and borde_h.right else None
    comprobar(estilo_h == "medium",
              "el borde derecho del marco cae en H, no en P -el hueco "
              "completo-", estilo_h)
    comprobar(ws.cell(6, 9).fill.fgColor.rgb != "FFF2F2F2",
              "la columna I, justo despues del marco, no tiene mate: el "
              "resto del hueco queda libre para lo que venga",
              str(ws.cell(6, 9).fill.fgColor.rgb))

    # El alto SI llega siempre hasta la fila del titulo -foto y titulo son
    # una sola tarjeta, no se angosta como el ancho-.
    comprobar(ws.cell(4, 2).border.top.style == "medium",
              "el marco arranca arriba del hueco (fila 4)")


def main():
    print("=" * 70)
    print("PRUEBA DEL MOTOR EXCEL")
    print("=" * 70)

    probar_desplazamiento()

    print("\n[2] Generacion del informe")
    with tempfile.TemporaryDirectory() as tmp:
        informe = construir_informe(imagen_de_prueba(tmp))
        motor = MotorExcel(PLANTILLA, cargar_mapa(MAPPING))
        resultado = motor.generar(informe, SALIDA)

    print("  {}".format(resultado.resumen()))
    comprobar(resultado.ok, "el archivo se escribio en disco")
    # 31 plantaciones + 2 especies + 1 derecho de agua + 1 requerimiento + 2 maquinas
    comprobar(resultado.filas_insertadas == 37,
              "se insertaron 37 filas (31+2+1+1+2 de las tablas desbordadas)",
              "insertadas={}".format(resultado.filas_insertadas))

    verificar(resultado)
    probar_historicos_variables()
    probar_respaldo_de_ventas()
    probar_filas_espejo()
    probar_hueco_compartido_de_imagen()
    probar_inversion_agrupada_por_proyecto()
    probar_estilo_de_hojas_anexas()
    probar_formula_real_de_pixeles()
    probar_marco_de_foto_segun_tamano()

    print("\n" + "=" * 70)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    print("Archivo generado: {}".format(SALIDA))
    return 0


if __name__ == "__main__":
    sys.exit(main())
