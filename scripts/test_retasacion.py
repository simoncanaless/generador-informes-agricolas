"""Prueba del motor de retasaciones Rabobank (services/excel_engine_rabobank.py).

el usuario: pidio un modulo nuevo para actualizar una tasacion anterior al
formato de RaboFinance -motor de garantia (hipoteca+prenda), distinto del
de Santander (capacidad de pago). El mapping y las capacidades de fila
salen de medir celda por celda el archivo real que mando (Tasacion_Base_
Rabo.xlsx, de un cliente real); la plantilla en blanco
(templates/Formato Fitogreen Rabo.xlsx) se deriva de ese archivo quitando
sus datos.

Uso: python scripts/test_retasacion.py
"""

import io
import os
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import fitz  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from config.settings import ajustes  # noqa: E402
from schemas.retasacion import (  # noqa: E402
    Cliente, Construccion, DerechoAgua, DescripcionSuelos, FilaSuelo,
    Instalacion, InscripcionCBR, ParametrosRetasacion, PlantacionConRoyalty,
    PlantacionSinRoyalty, Predio, RetasacionRabobank, Tasador, TestigoCanal,
    TestigoPropiedad,
)
from services.excel_engine_rabobank import MotorRabobank
from services.importar_retasacion import leer_retasacion_excel
from services.importar_retasacion_pdf import leer_retasacion_pdf
from services.mapping import cargar_mapa
from services.retasacion_pendientes import leer_borrador
from services.renderizado import (  # noqa: E402
    GUION_ARRANQUE_IMPRESORA_SEGURA, GUION_LIBERAR_COM_EXCEL, GUION_RESTAURAR_IMPRESORA,
)

PLANTILLA = ajustes.raiz / "templates" / "Formato Fitogreen Rabo.xlsx"
MAPPING = ajustes.raiz / "config" / "mapping_rabobank.json"

#: Cuenta las celdas en error de una hoja ($hoja) en $errores RECORRIENDO sus
#: valores. No con SpecialCells: en una hoja protegida -las dos de
#: referencias del formato real lo estan- SpecialCells lanza una excepcion,
#: y el catch la contaba como "cero errores". Asi paso de largo, por meses,
#: la estadistica en #¡DIV/0! con una sola referencia (2026-10-03). Excel
#: devuelve un error como un entero negativo (#¡DIV/0! = -2146826281).
CONTAR_ERRORES_PS = r"""
        foreach ($v in @($hoja.UsedRange.Value2)) {{
            if ($v -is [int] -and $v -lt -2146820000) {{ $errores += 1 }}
        }}"""

#: Rastros del cliente real del que salio la plantilla -sus nombres y los RUT
#: de la empresa, la propietaria y la tasadora-, como huella SHA-256 de cada
#: palabra o numero normalizado (sin tildes, en mayuscula, sin puntos) y NO en
#: texto: un RUT o un nombre de un cliente no tiene que quedar escrito en el
#: repositorio, ni siquiera en la prueba que verifica que no este en la
#: plantilla (2026-10-03: estaban en claro y hubo que reescribir el historial).
HUELLAS_CLIENTE_REAL = {
    "c75a3f2136fce2e062422beb5b78453522c514af8097cce5095738210424941e",
    "68e09011e4bbd562fa4aac0e6a3e66607f9da0a2aec7733ed3523ba934bf67fc",
    "bb93265b493fdeee76e71f2fbb565941db24a82f33849cd0e14d795cbf26f67f",
    "757ebcb4a0aeffb9b2498e360f2573a9781a5b9d1e97ce500f0ab2cf8a42f439",
    "93fb57dde9f072d7db2041a8a488dd27614f48502472eda60ed76ec173581479",
    "e83ab5f43a03577d8f8875874d771044ca2d8f3c3c4285075274b370ee5fdaa9",
    "a00380b54df2ded331771c4eb57c6b2097300a413db2c96982c099fb6c8fc696",
}


def huellas_de(texto):
    """La huella de cada palabra y cada numero de un texto, normalizados como
    las de HUELLAS_CLIENTE_REAL ("11.111.111-1" -> "11111111" y "1")."""
    import hashlib
    import re
    import unicodedata
    for token in re.findall(r"[^\W\d_]+|\d[\d.]*\d|\d", str(texto)):
        normal = "".join(c for c in unicodedata.normalize("NFD", token)
                         if unicodedata.category(c) != "Mn").upper().replace(".", "")
        yield hashlib.sha256(normal.encode()).hexdigest()


fallos, comprobaciones = [], 0


def comprobar(condicion, descripcion, detalle=""):
    global comprobaciones
    comprobaciones += 1
    if condicion:
        print("  OK   {}".format(descripcion))
    else:
        print("  FALLA {} {}".format(descripcion, detalle))
        fallos.append(descripcion)
    return condicion


def informe_de_prueba(especie_valida=True):
    """Una retasacion sintetica completa -no datos de ningun cliente real."""
    return RetasacionRabobank(
        parametros=ParametrosRetasacion(
            fecha_retasacion=date(2026, 9, 23), uf_base=40844.79,
            valor_usd=911.58, region="VI", comuna="Rancagua"),
        cliente=Cliente(nombre="Agrícola Prueba SpA", rut="76.543.210-3",
                        propietario="Juan Pérez Soto", rut_propietario="11.111.111-1"),
        predio=Predio(
            nombre="Parcela N°2 Prueba", rol_avaluo="1001-1", comuna="Rancagua",
            superficie_legal_ha=14.15,
            inscripcion=InscripcionCBR(fojas="1234", numero="5678", anio=2003,
                                       conservador="Rancagua"),
            afecto_expropiacion=False),
        suelos_descripcion=DescripcionSuelos(ciudad_cercana="Rancagua"),
        suelos=[FilaSuelo(clase_uso="ZEUES", grado_habilitacion="Completa",
                          superficie_m2=141500, valor_unitario_uf_m2=0.24)],
        aguas=[DerechoAgua(
            fuente="Canal de Prueba", rol_avaluo="1001-1",
            derechos_cantidad=16.63, hipotecado_rabofinance=True,
            inscripcion=InscripcionCBR(fojas="23", numero="45", anio=2005,
                                       conservador="Rancagua"),
            caudal_nominal_ls=24.945, valor_unitario_clp=3500000)],
        construcciones_hipotecables=[
            Construccion(edificacion="Casa patronal", estado_conservacion="Bueno",
                        ubicacion_rol="1001-1", superficie_m2=180,
                        costo_reposicion_uf_m2=25, vida_util_total_anios=60,
                        vida_remanente_anios=40)],
        instalaciones=[
            Instalacion(item="Riego tecnificado por goteo", ubicacion_rol="1001-1",
                       cantidad=16, unidad="ha", costo_reposicion_unitario_uf=60,
                       vida_util_total_anios=20, vida_remanente_anios=14)],
        plantaciones_sin_royalty=[
            PlantacionSinRoyalty(
                especie="Vid Vinifera" if especie_valida else "Especie Inventada",
                variedad="Pinot Gris", anio_plantacion=2018,
                densidad_plantas_ha=3704, estructura="No", rol_avaluo="1001-1",
                superficie_ha=2)],
        plantaciones_con_royalty=[
            PlantacionConRoyalty(especie="Arándano", variedad="Legacy",
                                 anio_plantacion=2021, marco_plantacion="3x1",
                                 valor_unitario_clp_ha=15000000,
                                 rol_avaluo="1001-1", superficie_ha=1.5)],
        tasador=Tasador(nombre="Prueba Tasador", rut="22.222.222-2"),
        # Cuatro de cada una, como un informe real: con una sola, la
        # desviacion estandar de la estadistica es #¡DIV/0! -con razon-.
        testigos_propiedades=[TestigoPropiedad(
            ubicacion="Sector Prueba {}".format(i), region="VI", comuna="Rancagua",
            superficie_m2=50000 + 5000 * i, valor_total_uf=12000 + 900 * i,
            tipo_referencia="Oferta") for i in range(4)],
        testigos_derechos_agua=[TestigoCanal(
            fuente="Canal Prueba {}".format(i), tipo_fuente="Canal", tipo="L/s",
            caudal_nominal_ls=10 + 3 * i, valor_total_m_pesos=40000 + 2500 * i,
            tipo_referencia="Oferta") for i in range(4)],
    )


def probar_plantilla_sin_rastro_de_cliente_real():
    print("\n[1] La plantilla en blanco no tiene ningun dato del cliente real")
    wb = load_workbook(PLANTILLA)
    encontrados = []
    for nombre in wb.sheetnames:
        for fila in wb[nombre].iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and (
                        HUELLAS_CLIENTE_REAL.intersection(huellas_de(celda.value))):
                    encontrados.append("{}!{}".format(nombre, celda.coordinate))
    comprobar(not encontrados,
              "sin nombres ni RUT del cliente de origen en ninguna celda de la plantilla",
              encontrados)


def probar_motor_genera_sin_errores():
    print("\n[2] El motor genera un archivo, sin pisar formulas del banco")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "retasacion.xlsx"
        resultado = motor.generar(informe_de_prueba(), destino)
        comprobar(resultado.ok, "el archivo se generó")
        comprobar(resultado.celdas_escritas > 100,
                  "escribió un número razonable de celdas", resultado.celdas_escritas)
        # Hasta 2026-10-03 la plantilla traia en AGUA!G25 la formula del
        # cliente de origen (=D25*1.5, la equivalencia de SU canal) y el motor
        # se negaba a escribir el caudal. Ahora la celda es del informe.
        agua = load_workbook(destino)[load_workbook(PLANTILLA).sheetnames[3]]
        comprobar(agua["G25"].value == 24.945,
                  "el caudal se escribe en AGUA!G25 -la plantilla ya no trae la "
                  "equivalencia del canal del cliente de origen-", agua["G25"].value)
        comprobar(not [a for a in resultado.avisos if "no se sobrescribe" in a],
                  "ninguna celda que el motor escribe choca con una fórmula del banco",
                  resultado.avisos)


def probar_especie_invalida_avisa_y_no_inventa_valor():
    print("\n[3] Una especie que no está en PLANT avisa, no inventa un valor")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "retasacion.xlsx"
        resultado = motor.generar(informe_de_prueba(especie_valida=False), destino)
        avisos_especie = [a for a in resultado.avisos if "no está en la biblioteca" in a]
        comprobar(len(avisos_especie) == 1, "avisó exactamente una vez", resultado.avisos)


def _recalcular_con_excel(destino: Path):
    """Abre con Excel real, fuerza el recalculo, y devuelve (errores, valores)
    de un puñado de celdas clave. None si no hay Excel disponible.

    Usa los mismos fragmentos de PowerShell que services/renderizado.py
    -impresora por defecto a salvo, COM liberado de verdad- porque esta
    automatizacion es SU PROPIA copia independiente, no pasa por ese modulo.
    Costó encontrarlo: este mismo script se quedó colgado de verdad con el
    diálogo "Configurar impresora" corriendo esta sesión, aunque
    renderizado.py ya tenía el arreglo -era ESTE script el que no lo
    tenía."""
    guion = r"""
$ErrorActionPreference = 'Stop'
$excel = $null
$libro = $null
try {{
{arranque_impresora}
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $libro = $excel.Workbooks.Open('{ruta}', $false, $false)
    $excel.CalculateFullRebuild()
    $excel.CalculateUntilAsyncQueriesDone()
    $errores = 0
    foreach ($hoja in $libro.Sheets) {{{contar_errores}
    }}
    Write-Output "ERRORES=$errores"
    Write-Output "RESUMEN_C19=$($libro.Sheets.Item('RESUMEN VALOR').Range('C19').Value2)"
    Write-Output "PORTADA_F31=$($libro.Sheets.Item('PORTADA').Range('F31').Value2)"
    Write-Output "PLANT_AA37=$($libro.Sheets.Item('PLANT').Range('AA37').Value2)"
    $libro.Close($false)
}} finally {{
    if ($excel -ne $null) {{ $excel.Quit() }}
{liberar_com}
{restaurar_impresora}
}}
""".format(
        ruta=str(destino.resolve()).replace("'", "''"),
        arranque_impresora=GUION_ARRANQUE_IMPRESORA_SEGURA,
        restaurar_impresora=GUION_RESTAURAR_IMPRESORA,
        liberar_com=GUION_LIBERAR_COM_EXCEL,
        contar_errores=CONTAR_ERRORES_PS.format(),
    )
    try:
        proceso = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
            capture_output=True, text=True, timeout=170)
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        return None, "no se pudo recalcular con Excel real: {}".format(error)
    if proceso.returncode != 0:
        return None, "Excel no está instalado o falló: {}".format(
            (proceso.stderr or proceso.stdout)[:200])
    valores = {}
    for linea in proceso.stdout.splitlines():
        if "=" in linea:
            clave, _, valor = linea.partition("=")
            valores[clave.strip()] = valor.strip()
    return valores, None


def probar_recalculo_con_excel_real():
    print("\n[4] Recalculo con Excel real: 0 errores, y los dos caminos de "
          "valor comercial coinciden")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    destino = ajustes.ruta_salida_pruebas / "test_retasacion_rabobank.xlsx"
    resultado = motor.generar(informe_de_prueba(), destino)
    comprobar(resultado.ok, "el archivo se generó")

    valores, motivo = _recalcular_con_excel(destino)
    if valores is None:
        print("  ({})".format(motivo))
        return

    errores = int(valores.get("ERRORES", "-1"))
    comprobar(errores == 0, "cero celdas en error en todo el libro",
              "{} en error".format(errores))

    resumen = float(valores.get("RESUMEN_C19", "nan"))
    portada = float(valores.get("PORTADA_F31", "nan"))
    comprobar(abs(resumen - portada) < 1,
              "RESUMEN VALOR!C19 y PORTADA!F31 -dos fórmulas independientes "
              "de la propia plantilla- coinciden exacto ({} vs {})".format(
                  resumen, portada))

    plant = float(valores.get("PLANT_AA37", "0"))
    comprobar(plant > 0,
              "el motor de costo de reposición (PLANT) calculó un valor real, "
              "no cero, para la plantación con especie válida", plant)


def probar_lectura_recupera_lo_que_escribio_el_motor():
    print("\n[5] La lectura del Paso 1 (importar_retasacion) recupera lo que "
          "el motor acaba de escribir -el mismo mapping, en reversa-")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    original = informe_de_prueba()
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "retasacion.xlsx"
        resultado = motor.generar(original, destino)
        comprobar(resultado.ok, "el archivo se generó")

        datos, avisos = leer_retasacion_excel(destino, mapa)

        comprobar(not any("ningún dato" in a for a in avisos),
                  "no avisa 'sin datos' -reconoció el propio formato que acaba "
                  "de escribir", avisos)
        comprobar(datos.get("cliente", {}).get("nombre") == original.cliente.nombre,
                  "cliente.nombre vuelve igual", datos.get("cliente"))
        comprobar(datos.get("cliente", {}).get("rut") == original.cliente.rut,
                  "cliente.rut vuelve igual")
        comprobar(datos.get("predio", {}).get("nombre") == original.predio.nombre,
                  "predio.nombre vuelve igual")
        comprobar(
            datos.get("predio", {}).get("inscripcion", {}).get("fojas")
            == original.predio.inscripcion.fojas,
            "predio.inscripcion.fojas vuelve igual -aplanado desde LEGAL!C6")
        comprobar(
            datos.get("predio", {}).get("afecto_expropiacion") is False,
            "predio.afecto_expropiacion vuelve como booleano False, no como "
            "el texto 'No' que trae la celda -normalización de §CAMPOS_BOOLEANOS",
            datos.get("predio", {}).get("afecto_expropiacion"))

        suelos = datos.get("suelos", [])
        comprobar(len(suelos) == 1, "una fila de suelos vuelve -no las 7 vacías",
                  suelos)
        if suelos:
            comprobar(suelos[0].get("clase_uso") == original.suelos[0].clase_uso,
                      "suelos[0].clase_uso vuelve igual")

        aguas = datos.get("aguas", [])
        comprobar(len(aguas) == 1, "un derecho de agua vuelve -no los 4 vacíos",
                  aguas)
        if aguas:
            comprobar(
                aguas[0].get("hipotecado_rabofinance") is True,
                "aguas[0].hipotecado_rabofinance vuelve como booleano True "
                "-el mismo campo se lee de DOS tablas (4.1 y 4.2) y no se "
                "pisa entre sí", aguas[0])
            comprobar(
                aguas[0].get("inscripcion", {}).get("fojas")
                == original.aguas[0].inscripcion.fojas,
                "aguas[0].inscripcion.fojas vuelve igual -viene solo de la "
                "sub-tabla 4.1, la 4.2 no la trae")

        plant_sin = datos.get("plantaciones_sin_royalty", [])
        comprobar(len(plant_sin) == 1, "una plantación sin royalty vuelve")
        if plant_sin:
            comprobar(
                plant_sin[0].get("especie") == original.plantaciones_sin_royalty[0].especie,
                "plantaciones_sin_royalty[0].especie vuelve igual")


def probar_lectura_de_archivo_ajeno_no_inventa_nada():
    print("\n[6] Un .xlsx de otro formato -sin las hojas del mapping- no "
          "inventa datos, solo avisa")
    mapa = cargar_mapa(MAPPING)
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "ajeno.xlsx"
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Hoja1"
        wb.active["A1"] = "esto no es una retasación Rabobank"
        wb.save(destino)

        datos, avisos = leer_retasacion_excel(destino, mapa)
        comprobar(not datos, "vuelve sin datos -ningún campo inventado", datos)
        comprobar(any("ningún dato" in a for a in avisos),
                  "avisa que no reconoció el formato", avisos)


def _generar_pdf_sintetico(ruta):
    """Un PDF de retasación de mentira -ningún dato de ningún cliente
    real-, para probar services/importar_retasacion_pdf.py sin comprometer
    los dos PDF reales que el usuario mandó de referencia (esos NO se suben al
    repositorio: son de clientes reales del banco).

    Reproduce la geometría mínima que el lector necesita: portada con
    etiqueta a la izquierda y valor a la derecha, la tabla de Inmuebles con
    el nombre del predio en su propia columna angosta -la misma trampa que
    costó resolver con los PDF reales, donde el nombre se colaba en la
    columna de Fs/Nº/Año si el texto era largo-, los deslindes envueltos
    en cuatro columnas (con `insert_textbox`, no `insert_text`: un texto
    largo en una sola línea desborda a la columna de al lado, que es
    justo el bug que tenía esta prueba la primera vez que se escribió), y
    un «Diferencias Tasación Anterior» por sección."""
    def fila(pagina, y, pares, tam=9):
        for x, texto in pares:
            pagina.insert_text((x, y), texto, fontsize=tam)

    doc = fitz.open()

    p1 = doc.new_page()
    fila(p1, 60, [(50, "EJECUTIVO :"), (150, "PRUEBA EJECUTIVO")])
    fila(p1, 75, [(50, "FECHA TASACION:"), (150, "01-01-2025")])
    fila(p1, 90, [(50, "UF BASE:"), (150, "UF 39,000.00")])
    fila(p1, 105, [(50, "USD:"), (150, "900.00")])
    fila(p1, 120, [(50, "N° SOLICITUD :"), (150, "9999 / 00000000")])
    fila(p1, 135, [(50, "Región"), (150, "Prueba")])
    fila(p1, 160, [(50, "Cliente:"), (100, "CLIENTE SINTETICO SPA")])
    fila(p1, 175, [(50, "Rut:"), (100, "76543210-3")])
    fila(p1, 190, [(50, "Propietario:"), (100, "CLIENTE SINTETICO SPA")])
    fila(p1, 205, [(50, "Rut:"), (100, "76543210-3")])
    fila(p1, 230, [(50, "Predio N°"), (110, "Nombre Propiedad"), (250, "Rol Avalúo"),
                  (340, "Comuna"), (460, "Superficies (ha)")])
    fila(p1, 245, [(50, "1"), (110, "Fundo de Prueba"), (250, "111-11"),
                  (340, "Comuna Prueba"), (460, "25.50")])

    p2 = doc.new_page()
    p2.insert_text((50, 30), "2.-INMUEBLES", fontsize=11)
    fila(p2, 60, [(70, "Inscripción dominio"), (250, "Superficie"),
                 (330, "Tasación"), (380, "Contribuciones")])
    fila(p2, 72, [(50, "PREDIO 1"), (200, "C.B.R."), (240, "Rol SII"),
                 (300, "Gravámenes")])
    fila(p2, 84, [(60, "Fs"), (90, "Nº"), (120, "Año")])
    # "111-11" arranca 1pt ANTES que su propio encabezado ("Rol SII"@x240)
    # -a proposito, igual que el valor real medido en Fundo Ejemplo
    # (139-95@x304.75 con "Rol"@x305.59-: sin el margen de tolerancia de
    # _inscripcion_y_rol_predio(), este valor se cuela a la columna de al
    # lado y el rol de avaluo se pierde en silencio.
    fila(p2, 98, [(50, "Fundo"), (60, "1500"), (90, "200"), (120, "2015"),
                 (200, "Prueba"), (239, "111-11"), (250, "25.50"),
                 (300, "NO"), (330, "SI"), (380, "SÍ")])
    p2.insert_text((50, 108), "de Prueba", fontsize=9)
    p2.insert_text((50, 140), "Inmueble", fontsize=9)
    p2.insert_text((50, 160), "Deslindes legales de la propiedad según sus títulos", fontsize=9)
    fila(p2, 175, [(80, "Norte"), (200, "Sur"), (300, "Oriente"), (400, "Poniente")])
    # Ancho de cada caja BIEN antes del punto medio hacia la columna
    # siguiente (el límite que usa _deslindes()): 80->128 (medio
    # Norte/Sur=140), 200->238 (medio Sur/Oriente=250), 300->338 (medio
    # Oriente/Poniente=350).
    p2.insert_textbox(fitz.Rect(80, 180, 128, 235), "Deslinde norte de prueba.", fontsize=8)
    p2.insert_textbox(fitz.Rect(200, 180, 238, 235), "Deslinde sur de prueba.", fontsize=8)
    p2.insert_textbox(fitz.Rect(300, 180, 338, 235), "Deslinde oriente de prueba.", fontsize=8)
    p2.insert_textbox(fitz.Rect(400, 180, 490, 235), "Deslinde poniente de prueba.", fontsize=8)
    p2.insert_text((70, 230), "Inscripción dominio", fontsize=9)  # cierra la ventana del predio 1

    def pagina_diferencias(titulo, texto):
        pagina = doc.new_page()
        y = 30
        for token in titulo:
            pagina.insert_text((50, y), token, fontsize=11)
            y += 12
        pagina.insert_text((50, 80), "Diferencias Tasación Anterior", fontsize=10)
        pagina.insert_text((50, 95), texto, fontsize=9)
        return pagina

    # --- 3.- SUELOS: titulo + tabla "3.2 Valores" + diferencias --------
    # La tabla va ANTES de "Diferencias Tasación Anterior" -como en los
    # PDF reales-: _bloque_de_texto() lee todo lo que hay DEBAJO de esa
    # etiqueta, así que si la tabla quedara después se leería como si
    # fuera parte del párrafo de diferencias.
    p_suelos = doc.new_page()
    p_suelos.insert_text((50, 30), "3.- SUELOS", fontsize=11)
    p_suelos.insert_text((50, 120), "Valores", fontsize=10)
    fila(p_suelos, 140, [(50, "Predio"), (150, "Rol SII"), (230, "Superficie"),
                        (300, "Valor"), (340, "$ /ha"), (400, "Total $"), (460, "Valor UF")])
    fila(p_suelos, 155, [(50, "Clase III R Plantado"), (150, "222-22"), (230, "5.00"),
                        (295, "$"), (300, "10,000,000"), (395, "$"), (400, "50,000,000"),
                        (455, "UF"), (460, "1,336.90")])
    fila(p_suelos, 170, [(50, "Clase IV"), (150, "222-22"), (230, "3.00"),
                        (295, "$"), (300, "8,000,000"), (395, "$"), (400, "24,000,000"),
                        (455, "UF"), (460, "641.71")])
    p_suelos.insert_text((50, 200), "Diferencias Tasación Anterior", fontsize=10)
    p_suelos.insert_text((50, 215), "Se mantiene criterio de prueba en suelos.", fontsize=9)

    # --- 4.- Derechos de Aguas: titulo + tabla + observaciones + diferencias
    p_aguas = doc.new_page()
    p_aguas.insert_text((50, 30), "4.- Derechos de Aguas", fontsize=11)
    fila(p_aguas, 140, [(50, "Derecho"), (170, "Origen"), (230, "Rol"),
                       (270, "Lts/seg"), (320, "Valor"), (340, "$"), (348, "lt/s"),
                       (400, "Total"), (420, "$"), (460, "Valor"), (480, "UF")])
    # Fila 1: cantidad entera, simple -y los dos "$" con un jitter de Y
    # A PROPOSITO (el de Total, más a la derecha, con Y más chica que el de
    # Valor): reproduce el bug real encontrado en Hijuela, donde ordenar
    # por (y, x) en vez de por x invertía cuál "$" era cuál.
    fila(p_aguas, 160, [(50, "3"), (58, "acciones"), (95, "Canal"), (130, "Prueba")])
    p_aguas.insert_text((175, 160), "Canal", fontsize=9)
    p_aguas.insert_text((175, 172), "Prueba", fontsize=9)  # origen envuelve 2 líneas
    p_aguas.insert_text((230, 163), "111-11", fontsize=9)
    p_aguas.insert_text((270, 163), "25.5", fontsize=9)
    p_aguas.insert_text((305, 163.4), "$", fontsize=9)
    p_aguas.insert_text((315, 163.4), "2,000,000", fontsize=9)
    p_aguas.insert_text((400, 163.0), "$", fontsize=9)  # Y mas chica, X mas grande
    p_aguas.insert_text((410, 163.0), "51,000,000", fontsize=9)
    p_aguas.insert_text((460, 163), "UF", fontsize=9)
    p_aguas.insert_text((470, 163), "1,300.00", fontsize=9)
    # Fila 2: cantidad con COMA DECIMAL ('4,175' = 4.175, no 4175 -verificado
    # contra los PDF reales, ver _numero_coma_decimal), descripción que
    # envuelve en 2 líneas ANTES de la línea del ancla (como en los PDF
    # reales: el wrap no siempre arranca en la misma línea que el Lts/seg),
    # y un total en "-" (sin dato, como el PDF real también trae).
    fila(p_aguas, 185, [(50, "4,175"), (75, "Regadores"), (130, "de"), (145, "Prueba")])
    fila(p_aguas, 197, [(50, "y"), (58, "Testigos")])
    p_aguas.insert_text((175, 190), "Canal", fontsize=9)
    p_aguas.insert_text((175, 202), "Prueba", fontsize=9)
    p_aguas.insert_text((230, 200), "111-11", fontsize=9)
    p_aguas.insert_text((270, 200), "62.6", fontsize=9)
    p_aguas.insert_text((305, 200), "$", fontsize=9)
    p_aguas.insert_text((315, 200), "3,000,000", fontsize=9)
    p_aguas.insert_text((400, 200), "$", fontsize=9)
    p_aguas.insert_text((410, 200), "-", fontsize=9)
    p_aguas.insert_text((460, 200), "UF", fontsize=9)
    p_aguas.insert_text((470, 200), ".00", fontsize=9)
    p_aguas.insert_text((50, 220), "Total", fontsize=9)
    # Gap generoso entre etiqueta y contenido: con un gap chico (probado y
    # roto a propósito) el borde inferior de "Observaciones:" queda POR
    # DEBAJO del borde superior de su propio texto -se superponen un pelo-
    # y `_bloque_de_texto` (que exige contenido estrictamente debajo de la
    # etiqueta) no encuentra nada.
    p_aguas.insert_text((50, 235), "Observaciones:", fontsize=9)
    p_aguas.insert_text((50, 255), "Se recomienda revisar aforo de prueba.", fontsize=9)
    p_aguas.insert_text((50, 280), "Diferencias Tasación Anterior", fontsize=10)
    p_aguas.insert_text((50, 295), "Sin diferencias en aguas.", fontsize=9)

    # --- 5.1- Construcciones: titulo + tabla + total sin rotulo + diferencias
    # Ancho explicito: las tablas reales de construcciones/instalaciones
    # llegan a x=846 (son paginas mas anchas que la A4 por defecto de
    # fitz.new_page, 595pt) -sin esto, insert_text mas alla del borde de la
    # hoja simplemente no queda en el texto extraible.
    p_constr = doc.new_page(width=900, height=450)
    p_constr.insert_text((50, 30), "5.1- Construcciones", fontsize=11)
    fila(p_constr, 150, [(50, "Tipo"), (190, "Material"), (330, "Tipo"),
                        (390, "Superficie"), (580, "Asegurable")])
    fila(p_constr, 162, [(50, "Construcción"), (240, "Año"), (280, "Descripción"),
                        (330, "de"), (430, "Valor"), (490, "Total"), (540, "Total")])
    fila(p_constr, 174, [(280, "ubicación"), (330, "Garantía"), (430, "por"),
                        (490, "$"), (540, "UF")])
    # Fila 1: Hipoteca, simple -una sola linea.
    fila(p_constr, 195, [(50, "Bodega Prueba"), (190, "Albañilería"), (240, "2020"),
                        (280, "100-10"), (330, "Hipoteca"), (390, "50.0"),
                        (540, "UF"), (560, "106.50"), (600, "Sí")])
    # Fila 2: Prenda, y el rol partido en dos líneas CON otra columna
    # metida en el medio por Y -mismo caso real que costó en instalaciones
    # ('45 KVA' quedaba entre las dos mitades de '139-'/'95')-: acá "2015"
    # (el año) cae, por Y, entre "100-" y "20".
    fila(p_constr, 215, [(50, "Bodega Prenda"), (190, "Metal"), (280, "100-")])
    fila(p_constr, 222, [(240, "2015")])
    fila(p_constr, 229, [(280, "20"), (330, "Prenda"), (390, "80.0"),
                        (540, "UF"), (560, "212.00"), (600, "No")])
    # Fila 3: SIN "Tipo de Garantía" -no se sabe si es hipotecable o
    # prendable, se descarta entera (no se inventa a cuál lista pertenece).
    fila(p_constr, 249, [(50, "Sin Garantia"), (190, "Madera"), (240, "2000"),
                        (280, "100-30"), (390, "30.0")])
    # Fila de SUMA sin la palabra "Total" -pasa en el PDF real de Ejemplo-:
    # solo el numero (50+80+30=160), en la misma columna de Superficie.
    # Tiene que descartarse por no traer nombre de bien, no por buscar un
    # rótulo que en este caso no existe.
    p_constr.insert_text((390, 269), "160.0", fontsize=9)
    p_constr.insert_text((50, 290), "Observaciones:", fontsize=9)
    p_constr.insert_text((50, 310), "Sin observaciones de prueba.", fontsize=9)
    p_constr.insert_text((50, 335), "Diferencias Tasación Anterior", fontsize=10)
    p_constr.insert_text((50, 350), "Sin diferencias en construcciones.", fontsize=9)

    # --- 5.2.- Instalaciones: titulo + tabla + diferencias -----------------
    p_instal = doc.new_page(width=900, height=350)
    p_instal.insert_text((50, 30), "5.2.- Instalaciones", fontsize=11)
    fila(p_instal, 150, [(50, "Instalación"), (250, "Unidad"), (300, "Descripción"),
                        (450, "Tipo"), (500, "Valor"), (600, "Total"), (650, "Total"),
                        (700, "Asegurable")])
    fila(p_instal, 162, [(300, "y"), (310, "ubicación"), (450, "de"), (500, "por"),
                        (600, "$"), (650, "UF")])
    fila(p_instal, 174, [(450, "Garantía")])
    # Fila 1: el "$" viene PEGADO al número (como en los dos PDF reales) y
    # el Total UF SIN el rótulo "UF" -solo el número, como en Ejemplo-.
    fila(p_instal, 195, [(50, "Riego tecnificado"), (250, "40.0"), (310, "222-22"),
                        (450, "Prenda"), (500, "$50,000"), (600, "$2,000,000"),
                        (680, "53.30"), (730, "No")])
    p_instal.insert_text((50, 220), "Diferencias Tasación Anterior", fontsize=10)
    p_instal.insert_text((50, 235), "Se agrega una instalación de prueba.", fontsize=9)

    # --- 6.- Plantaciones: titulo + tabla sin-royalty + diferencias ----
    p_plant = doc.new_page()
    p_plant.insert_text((50, 30), "6.- Plantaciones", fontsize=11)
    p_plant.insert_text((50, 100), "Royalties", fontsize=10)
    fila(p_plant, 120, [(50, "Tipo de Cultivo"), (150, "Variedad"), (210, "Año"),
                       (250, "Marco"), (300, "Rol"), (340, "ha"), (380, "$ x ha"),
                       (430, "Total $"), (480, "Total UF"), (530, "Estado")])
    # una especie CORTA (una sola linea) y otra LARGA que se parte en dos
    # palabras -la misma trampa real que costo encontrar: un margen
    # izquierdo adivinado en vez de anclado al encabezado clasificaba mal
    # 'MESA' de 'VID DE MESA' como si fuera la variedad.
    fila(p_plant, 135, [(50, "PALTO"), (150, "Hass"), (210, "2018"), (250, "5X3"),
                       (300, "222-22"), (335, "2.50"), (378, "$"), (383, "4,000,000"),
                       (425, "$"), (430, "10,000,000"), (475, "UF"), (480, "267.38"),
                       (525, "Bueno")])
    fila(p_plant, 150, [(50, "VID DE MESA"), (210, "2019"), (250, "3X2"),
                       (300, "222-22"), (335, "1.80"), (378, "$"), (383, "3,500,000"),
                       (425, "$"), (430, "6,300,000"), (475, "UF"), (480, "168.45"),
                       (525, "Bueno")])
    p_plant.insert_text((165, 158), "Thompson", fontsize=9)  # variedad, linea propia
    p_plant.insert_text((50, 180), "Total frutales sin royalti", fontsize=9)
    p_plant.insert_text((50, 220), "Diferencias Tasación Anterior", fontsize=10)
    p_plant.insert_text((50, 235), "Sin diferencias en plantaciones.", fontsize=9)

    # --- 8.- REFERENCIAS DE MERCADO: testigos de propiedades, en pesos y
    # SIN tasa UF declarada -el caso mas dificil, el de Ejemplo: sin tasa,
    # el valor queda en 'valor_total_clp_sin_convertir' en vez de un UF
    # inventado-. Columnas espaciadas generosamente (>=60pt): un espacio
    # mas chico ya hizo que "Valor"+"Total" se fundieran en un solo token
    # ("ValorTotal") en un intento anterior de este mismo archivo.
    p_testprop = doc.new_page(width=850, height=280)
    p_testprop.insert_text((50, 30), "8.- REFERENCIAS DE MERCADO", fontsize=11)
    p_testprop.insert_text((50, 50), "OFERTAS PREDIOS AGRÍCOLAS", fontsize=9)
    fila(p_testprop, 80, [(60, "Propiedad"), (170, "Superficie"), (250, "Valor"),
                         (290, "Total"), (380, "Región"), (450, "Comuna"),
                         (530, "Observaciones")])
    fila(p_testprop, 92, [(170, "Hás"), (290, "$")])
    # Fila 1: numero corto (11 caracteres) cerca del encabezado "Total".
    # Fila 2: numero LARGO (13 caracteres) que arranca mas a la izquierda
    # -columna alineada a la derecha, mismo caso real medido en Malloa
    # (Ejemplo)-: si el margen de busqueda no le da cabida, se pierde y
    # cae a la columna vecina.
    # Los nombres de prueba de comuna NO contienen la palabra "Comuna"
    # literal a propósito -si la contuvieran, la búsqueda del bloque de
    # encabezado (que busca esa misma palabra) confundiría el propio DATO
    # de la fila 1 con el encabezado, y esa fila se perdía por completo:
    # bug real encontrado armando esta misma prueba.
    fila(p_testprop, 110, [(60, "Prueba Uno"), (170, "22.0"), (285, "550,000,000"),
                          (380, "VI"), (450, "Rancagua")])
    p_testprop.insert_text((530, 110), "https://www.ejemplo.cl/predio-uno", fontsize=8)
    fila(p_testprop, 140, [(60, "Prueba Dos"), (170, "23.0"), (276, "1,789,607,350"),
                          (380, "VI"), (450, "Peumo")])
    p_testprop.insert_text((530, 140), "https://www.ejemplo.cl/predio-dos", fontsize=8)
    p_testprop.insert_text((60, 200), "Total Muestra", fontsize=9)
    p_testprop.insert_text((170, 200), "2 unidades", fontsize=9)

    # --- 8.1- REFERENCIAS DE COMPRAVENTA DERECHOS AGUAS SUPERFICIALES -----
    # Acá el "Valor Total" SÍ viene en UF directo (como Hijuela): no hace
    # falta tasa de conversión. valor_unitario_uf_ls y valor_total_m_pesos
    # se derivan por aritmética simple de los DOS números que ya trae la
    # tabla -Total UF y el valor unitario en pesos-, no de una tasa externa.
    p_testagua = doc.new_page(width=900, height=250)
    p_testagua.insert_text((50, 30), "8.1- REFERENCIAS DE COMPRAVENTA DERECHOS AGUAS",
                          fontsize=10)
    fila(p_testagua, 80, [(150, "Caudal"), (230, "Valor"), (270, "Total"),
                         (360, "Valor"), (400, "unit."), (490, "Región"), (560, "Comuna")])
    fila(p_testagua, 92, [(150, "L/s"), (270, "UF")])
    fila(p_testagua, 110, [(50, "1. Río Prueba"), (150, "5.0"),
                          (265, "500"), (390, "2,000,000"), (490, "VI"),
                          (560, "Coltauco")])
    p_testagua.insert_text((650, 110), "https://www.ejemplo.cl/agua-uno", fontsize=8)
    p_testagua.insert_text((70, 170), "Total Muestra", fontsize=9)
    p_testagua.insert_text((170, 170), "1 unidad", fontsize=9)

    p3 = doc.new_page()
    p3.insert_text((50, 30), "RESUMEN DE VALORES", fontsize=12)
    fila(p3, 60, [(50, "1) PREDIO DENOMINADO:"), (250, "Fundo de Prueba")])
    fila(p3, 200, [(50, "FECHA"), (120, "01-01-2025"), (200, "VALOR DÓLAR"), (300, "900.00")])
    fila(p3, 215, [(200, "VALOR UF"), (300, "39,000.00")])
    fila(p3, 230, [(50, "TASADOR"), (120, "Prueba Tasador Sintetico")])

    doc.save(ruta)
    doc.close()


def probar_lectura_de_pdf_sintetico():
    print("\n[7] La lectura de un PDF de retasación anterior (sin IA, por posición)")
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "retasacion_sintetica.pdf"
        _generar_pdf_sintetico(ruta)
        datos, avisos = leer_retasacion_pdf(ruta)

    comprobar(datos.get("cliente", {}).get("nombre") == "CLIENTE SINTETICO SPA",
              "cliente.nombre", datos.get("cliente"))
    comprobar(datos.get("cliente", {}).get("rut") == "76543210-3", "cliente.rut")
    comprobar(datos.get("cliente", {}).get("propietario") == "CLIENTE SINTETICO SPA",
              "cliente.propietario")
    comprobar(datos.get("cliente", {}).get("rut_propietario") == "76543210-3",
              "cliente.rut_propietario -viene de la SEGUNDA ocurrencia de 'Rut:'")

    predio = datos.get("predio", {})
    comprobar(predio.get("nombre") == "Fundo de Prueba",
              "predio.nombre -viene de 'PREDIO DENOMINADO:', no de la tabla "
              "de la portada, que en un PDF real puede tener el nombre partido "
              "en varias líneas", predio)
    comprobar(predio.get("rol_avaluo") == "111-11",
              "predio.rol_avaluo -viene de la tabla de Inmuebles, con el margen "
              "de tolerancia que hizo falta agregar (el valor real arranca un "
              "poco antes que su propio encabezado)")
    comprobar(predio.get("comuna") == "Comuna Prueba",
              "predio.comuna -dos palabras, no se corta en la primera")
    comprobar(predio.get("superficie_legal_ha") == 25.5, "predio.superficie_legal_ha")
    insc = predio.get("inscripcion", {})
    comprobar(insc == {"fojas": "1500", "numero": "200", "anio": 2015, "conservador": "Prueba"},
              "predio.inscripcion completa -el nombre del predio, a la izquierda "
              "de Fs/Nº/Año, no se cuela como si fuera un número", insc)
    for lado, esperado in (("deslinde_norte", "Deslinde norte de prueba."),
                           ("deslinde_sur", "Deslinde sur de prueba."),
                           ("deslinde_oriente", "Deslinde oriente de prueba."),
                           ("deslinde_poniente", "Deslinde poniente de prueba.")):
        comprobar(predio.get(lado) == esperado,
                  "predio.{} -no se mezcla con la columna vecina".format(lado),
                  predio.get(lado))

    comprobar(datos.get("tasador", {}).get("nombre") == "Prueba Tasador Sintetico",
              "tasador.nombre")

    suelos = datos.get("suelos", [])
    comprobar(len(suelos) == 2, "suelos: dos filas -no la fila de encabezado ni ninguna vacía",
              suelos)
    if len(suelos) == 2:
        comprobar(suelos[0]["clase_uso"] == "Clase III R Plantado",
                  "suelos[0].clase_uso -el texto libre completo, no cortado antes del rol",
                  suelos[0])
        comprobar(suelos[0]["superficie_m2"] == 50000.0,
                  "suelos[0].superficie_m2 -convertida de 5.00 ha", suelos[0])
        comprobar(abs(suelos[0]["valor_unitario_uf_m2"] - 0.026738) < 0.0001,
                  "suelos[0].valor_unitario_uf_m2 -de (Valor UF total / ha) / 10000, "
                  "no del $/ha ni de un tipo de cambio supuesto", suelos[0])

    plant_sin = datos.get("plantaciones_sin_royalty", [])
    comprobar(len(plant_sin) == 2, "plantaciones_sin_royalty: dos filas", plant_sin)
    if len(plant_sin) == 2:
        comprobar(plant_sin[0]["especie"] == "PALTO" and plant_sin[0]["variedad"] == "Hass",
                  "plantaciones_sin_royalty[0] especie/variedad", plant_sin[0])
        comprobar(plant_sin[0]["densidad_plantas_ha"] == 667,
                  "plantaciones_sin_royalty[0].densidad_plantas_ha -de un marco '5X3' "
                  "(10000/(5*3)), no un texto literal", plant_sin[0])
        comprobar(
            plant_sin[1]["especie"] == "VID DE MESA" and plant_sin[1]["variedad"] == "Thompson",
            "plantaciones_sin_royalty[1] -especie de tres palabras SIN cortarse antes de "
            "tiempo: bug real encontrado esta misma sesion, un margen izquierdo adivinado "
            "en vez de anclado al encabezado clasificaba 'MESA' como si fuera la variedad",
            plant_sin[1])
        comprobar("precio" not in plant_sin[0] and "total" not in plant_sin[0],
                  "plantaciones_sin_royalty no trae el valor -eso lo calcula PLANT en el "
                  "motor nuevo, precargarlo repetiria un numero de la tasacion VIEJA",
                  plant_sin[0])

    aguas = datos.get("aguas", [])
    comprobar(len(aguas) == 2, "aguas: dos filas -no la de encabezado ni ninguna vacía", aguas)
    if len(aguas) == 2:
        comprobar(aguas[0]["derechos_cantidad"] == 3.0 and aguas[0]["tipo"] == "acciones",
                  "aguas[0] cantidad/tipo -fila simple, una sola línea", aguas[0])
        comprobar(aguas[0]["caudal_nominal_ls"] == 25.5, "aguas[0].caudal_nominal_ls", aguas[0])
        comprobar(aguas[0]["valor_unitario_clp"] == 2_000_000.0,
                  "aguas[0].valor_unitario_clp -el PRIMER '$' en orden de X, no el "
                  "'Total $' -los dos '\"$\"' de esta fila se pusieron con un jitter de "
                  "Y a propósito (el de Total, más a la derecha, con Y más chica): "
                  "bug real encontrado en el PDF de Hijuela, donde ordenar por (y, x) "
                  "en vez de por x devolvía el total en vez del precio unitario",
                  aguas[0])
        comprobar(aguas[0]["rol_avaluo"] == "111-11", "aguas[0].rol_avaluo", aguas[0])
        comprobar(aguas[0]["fuente"] == "Canal Prueba", "aguas[0].fuente", aguas[0])
        comprobar(aguas[1]["derechos_cantidad"] == 4.175,
                  "aguas[1].derechos_cantidad -'4,175' es 4.175, NO 4175: la coma es "
                  "el separador DECIMAL en esta columna (al revés que en las columnas "
                  "de dinero), verificado cruzando con la propia tabla de los PDF "
                  "reales", aguas[1])
        comprobar(aguas[1]["tipo"] == "Regadores",
                  "aguas[1].tipo -palabra completa, no mezclada con la descripción "
                  "que envuelve en dos líneas ANTES de la línea del ancla", aguas[1])
        comprobar(aguas[1]["caudal_nominal_ls"] == 62.6, "aguas[1].caudal_nominal_ls", aguas[1])

    comprobar(datos.get("aguas_observaciones") == "Se recomienda revisar aforo de prueba.",
              "aguas_observaciones -el bloque libre DESPUÉS de la tabla y ANTES de "
              "'Diferencias Tasación Anterior', sin tragarse esa sección siguiente",
              datos.get("aguas_observaciones"))

    hipotecables = datos.get("construcciones_hipotecables", [])
    prendables = datos.get("construcciones_prendables", [])
    comprobar(len(hipotecables) == 1 and len(prendables) == 1,
              "construcciones: una hipotecable y una prendable -la tabla es UNA "
              "sola con una columna 'Tipo de Garantía' que discrimina a cuál "
              "lista pertenece cada fila, no dos tablas separadas",
              (hipotecables, prendables))
    if len(hipotecables) == 1:
        h = hipotecables[0]
        comprobar(h["edificacion"] == "Bodega Prueba" and h["ubicacion_rol"] == "100-10",
                  "construcciones_hipotecables[0] -fila simple, una sola línea", h)
        comprobar(h["asegurable"] is True, "construcciones_hipotecables[0].asegurable", h)
        comprobar(abs(h["costo_reposicion_uf_m2"] - 2.13) < 0.001,
                  "construcciones_hipotecables[0].costo_reposicion_uf_m2 -derivado de "
                  "Total UF / Superficie (106.50/50.0), nunca de un tipo de cambio "
                  "supuesto", h)
    if len(prendables) == 1:
        p = prendables[0]
        comprobar(p["ubicacion_rol"] == "100-20",
                  "construcciones_prendables[0].ubicacion_rol -el rol partido en dos "
                  "líneas ('100-' / '20') CON otra columna (el año) metida en el "
                  "medio por Y: se empareja por columna de X, no por ser vecina "
                  "inmediata en la lista ordenada por lectura", p)
        comprobar(p["asegurable"] is False, "construcciones_prendables[0].asegurable", p)
        comprobar(abs(p["costo_reposicion_uf_m2"] - 2.65) < 0.001,
                  "construcciones_prendables[0].costo_reposicion_uf_m2", p)
    comprobar(not any(f.get("edificacion") == "Sin Garantia"
                      for f in hipotecables + prendables),
              "la fila SIN 'Tipo de Garantía' se descarta entera -no se inventa a "
              "cuál lista pertenece", hipotecables + prendables)
    comprobar(not any(f.get("superficie_m2") == 160.0 for f in hipotecables + prendables),
              "la fila de SUMA sin la palabra 'Total' -pasa en el PDF real de "
              "Ejemplo, no es un caso raro- no se lee como si fuera un bien: se "
              "descarta por no tener nombre de edificación, no por buscar un "
              "rótulo que en este caso no existe", hipotecables + prendables)

    instalaciones = datos.get("instalaciones", [])
    comprobar(len(instalaciones) == 1, "instalaciones: una fila", instalaciones)
    if len(instalaciones) == 1:
        i = instalaciones[0]
        comprobar(i["item"] == "Riego tecnificado" and i["cantidad"] == 40.0, "instalaciones[0]", i)
        comprobar(i["asegurable"] is False, "instalaciones[0].asegurable", i)
        comprobar(abs(i["costo_reposicion_unitario_uf"] - 1.3325) < 0.0001,
                  "instalaciones[0].costo_reposicion_unitario_uf -el '$' viene PEGADO "
                  "al número ('$2,000,000', un solo token, distinto de "
                  "construcciones/aguas) y el Total UF sin el rótulo 'UF' -solo el "
                  "número pegado al 'Total $', como en el PDF real de Ejemplo-", i)

    testigos_prop = datos.get("testigos_propiedades", [])
    comprobar(len(testigos_prop) == 2, "testigos_propiedades: dos filas", testigos_prop)
    if len(testigos_prop) == 2:
        t1, t2 = testigos_prop
        comprobar(t1["ubicacion"] == "Prueba Uno" and t1["superficie_m2"] == 220000.0,
                  "testigos_propiedades[0] -superficie convertida de 22.0 ha", t1)
        comprobar(t1.get("valor_total_uf") is None
                  and t1.get("valor_total_clp_sin_convertir") == 550000000.0,
                  "testigos_propiedades[0]: sin tasa UF declarada en la página, el valor "
                  "en pesos queda aparte -nunca un UF inventado", t1)
        comprobar(t2.get("valor_total_clp_sin_convertir") == 1789607350.0,
                  "testigos_propiedades[1].valor_total_clp_sin_convertir -el número MAS "
                  "LARGO (13 caracteres) arranca más a la izquierda que el de la fila "
                  "anterior porque la columna está alineada a la derecha; si el margen de "
                  "búsqueda no le daba cabida, se perdía y caía a la columna vecina", t2)
        comprobar(t1.get("comuna") == "Rancagua" and t2.get("comuna") == "Peumo",
                  "testigos_propiedades comuna -sin contaminación de la columna vecina",
                  testigos_prop)

    testigos_agua = datos.get("testigos_derechos_agua", [])
    comprobar(len(testigos_agua) == 1, "testigos_derechos_agua: una fila", testigos_agua)
    if len(testigos_agua) == 1:
        ta = testigos_agua[0]
        comprobar(ta["caudal_nominal_ls"] == 5.0, "testigos_derechos_agua[0].caudal_nominal_ls", ta)
        comprobar(abs(ta["valor_unitario_uf_ls"] - 100.0) < 0.001,
                  "testigos_derechos_agua[0].valor_unitario_uf_ls -derivado de Total UF "
                  "(500) / Caudal (5.0), los dos ya extraídos del PDF: nunca de una tasa "
                  "externa", ta)
        comprobar(abs(ta["valor_total_m_pesos"] - 10000.0) < 0.01,
                  "testigos_derechos_agua[0].valor_total_m_pesos -Caudal × Valor unit. "
                  "($2.000.000) / 1000, la misma aritmética por el otro camino", ta)

    diferencias = datos.get("diferencias", {})
    for campo, esperado in (
        ("suelos", "Se mantiene criterio de prueba en suelos."),
        ("aguas", "Sin diferencias en aguas."),
        ("construcciones", "Sin diferencias en construcciones."),
        ("instalaciones", "Se agrega una instalación de prueba."),
        ("plantaciones", "Sin diferencias en plantaciones."),
    ):
        comprobar(diferencias.get(campo) == esperado,
                  "diferencias.{} -de la página correcta, no confundida con la "
                  "fila 'SUELOS'/'CONSTRUCCIONES'/etc. del Resumen de Valores"
                  .format(campo), diferencias.get(campo))

    comprobar(any("UF base 39,000.00" in a for a in avisos),
              "el aviso informa la tasación anterior -nunca se precarga "
              "fecha/UF/USD como si fueran los de la tasación NUEVA", avisos)


def probar_incompleta_se_genera_y_avisa():
    print("\n[8] Una retasación incompleta se genera igual y dice qué le falta "
          "(2026-10-02)")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)

    # -- sin nada: antes era un 422 y no salía ningún archivo
    retasacion, pendientes = leer_borrador({})
    secciones = {p.seccion for p in pendientes}
    comprobar({"Datos de la tasación", "Cliente", "Predio", "Tasador"} <= secciones,
              "sin datos, avisa cabecera, cliente, predio y tasador", secciones)
    comprobar(all(p.tipo == "falta" for p in pendientes),
              "todo eso es «falta», nada «aviso» -no vino nada mal escrito",
              pendientes)
    graves = [p for p in pendientes if "PORTADA!H" in p.detalle]
    comprobar(len(graves) == 3,
              "la fecha, la UF y el dólar van aparte, cada uno con lo que "
              "descalza y dónde escribirlo después", graves)
    with tempfile.TemporaryDirectory() as tmp:
        resultado = motor.generar(retasacion, Path(tmp) / "vacia.xlsx")
        comprobar(resultado.ok, "y el archivo se genera igual")

    # -- completa: nada se marca de más
    _, pendientes = leer_borrador(informe_de_prueba().model_dump(mode="json"))
    comprobar(not pendientes,
              "una retasación completa no tiene ningún pendiente", pendientes)

    # -- mal escrito: no se escribe, y se dice qué se escribió
    datos = informe_de_prueba().model_dump(mode="json")
    datos["cliente"]["rut"] = "12.345.678-0"
    datos["parametros"]["uf_base"] = 0
    retasacion, pendientes = leer_borrador(datos)
    rut = [p for p in pendientes if "RUT del cliente" in p.detalle]
    comprobar(len(rut) == 1 and rut[0].tipo == "aviso"
              and "12.345.678-0" in rut[0].detalle,
              "un RUT con el dígito verificador malo se avisa con lo que se "
              "escribió", rut)
    uf = [p for p in pendientes if "Valor de la UF" in p.detalle]
    comprobar(len(uf) == 1 and uf[0].tipo == "aviso" and "PORTADA!H8" in uf[0].detalle,
              "una UF en 0 dice que no vale -no que falta- y dice igual dónde "
              "escribirla", uf)
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "rut_malo.xlsx"
        motor.generar(retasacion, destino)
        portada = load_workbook(destino)[mapa.nombre_hoja("portada")]
        comprobar(portada["H14"].value in (None, "") and portada["H8"].value in (None, ""),
                  "ni el RUT malo ni la UF en 0 llegan al Excel",
                  (portada["H14"].value, portada["H8"].value))
        comprobar(portada["C14"].value == datos["cliente"]["nombre"],
                  "el resto del cliente sí se escribe")

    # -- tablas: fila vacía en medio, campo extra de un PDF, filas de más
    datos = informe_de_prueba().model_dump(mode="json")
    datos["construcciones_hipotecables"] += [
        {}, {"edificacion": "Bodega", "superficie_m2": 100}]
    datos["testigos_propiedades"][0]["valor_total_clp_sin_convertir"] = 1789607350
    datos["suelos"] = datos["suelos"] * 9
    datos["plantaciones_sin_royalty"].append({"variedad": "Lapins", "estructura": "No"})
    retasacion, pendientes = leer_borrador(datos)
    textos = [p.detalle for p in pendientes]
    comprobar("fila 3 (Bodega): falta Costo de reposición (UF/m²), Vida útil (años)"
              in textos,
              "la fila se nombra con su número EN LA PANTALLA -3, aunque la 2 "
              "vacía se descarte- y su edificación", textos)
    comprobar(len(retasacion.construcciones_hipotecables) == 2,
              "la fila vacía no llega al Excel -dejaría un renglón en blanco")
    comprobar(not any("sin_convertir" in t for t in textos)
              and len(retasacion.testigos_propiedades) == len(informe_de_prueba().testigos_propiedades),
              "el campo extra que agrega la lectura de un PDF se descarta sin "
              "rechazar la fila -antes hacía fallar la generación-")
    comprobar(any("quedó fuera la última" in t for t in textos)
              and len(retasacion.suelos) == 8,
              "una fila de suelos de más: se avisa y se escriben las 8 que caben",
              textos)
    with tempfile.TemporaryDirectory() as tmp:
        try:
            resultado = motor.generar(retasacion, Path(tmp) / "filas.xlsx")
            comprobar(resultado.ok, "una plantación sin especie no hace caer "
                                    "al motor de PLANT")
        except Exception as error:  # noqa: BLE001
            comprobar(False, "una plantación sin especie no hace caer al "
                             "motor de PLANT", repr(error))

    # -- lo que se lee de un archivo anterior se puede generar sin rechazos
    with tempfile.TemporaryDirectory() as tmp:
        destino = Path(tmp) / "anterior.xlsx"
        motor.generar(informe_de_prueba(), destino)
        leidos, _ = leer_retasacion_excel(destino, mapa)
        _, pendientes = leer_borrador(leidos)
        comprobar(not [p for p in pendientes if p.tipo == "aviso"],
                  "lo leído de un .xlsx anterior no trae nada «mal escrito»",
                  pendientes)
        pdf = Path(tmp) / "anterior.pdf"
        _generar_pdf_sintetico(pdf)
        leidos, _ = leer_retasacion_pdf(str(pdf))
        try:
            retasacion, pendientes = leer_borrador(leidos)
            comprobar(motor.generar(retasacion, Path(tmp) / "desde_pdf.xlsx").ok,
                      "lo leído de un PDF anterior se genera, con lo que falte "
                      "en la lista ({} pendientes)".format(len(pendientes)))
        except Exception as error:  # noqa: BLE001
            comprobar(False, "lo leído de un PDF anterior se genera", repr(error))


def _recalcular_antes_y_despues(destino: Path, mapa):
    """Recalcula con Excel real un libro generado SIN fecha, UF ni dolar, lee
    lo que los avisos dicen que se descalza, y despues escribe esos tres
    datos a mano en PORTADA -como haria el tasador- y vuelve a leer."""
    guion = r"""
$ErrorActionPreference = 'Stop'
$excel = $null
$libro = $null
function Leer($libro, $etapa) {{
    $errores = 0
    foreach ($hoja in $libro.Sheets) {{{contar_errores}
    }}
    Write-Output "ERRORES_$etapa=$errores"
    Write-Output "SUELOS_G46_$etapa=$($libro.Sheets.Item('{suelos}').Range('G46').Value2)"
    Write-Output "PLANT_R37_$etapa=$($libro.Sheets.Item('PLANT').Range('R37').Value2)"
    Write-Output "PLANT_AA37_$etapa=$($libro.Sheets.Item('PLANT').Range('AA37').Value2)"
}}
try {{
{arranque_impresora}
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $libro = $excel.Workbooks.Open('{ruta}', $false, $false)
    $excel.CalculateFullRebuild()
    Leer $libro 'ANTES'
    $portada = $libro.Sheets.Item('{portada}')
    $portada.Range('H7').Value2 = ([datetime]'2026-09-23').ToOADate()
    $portada.Range('H8').Value2 = 40844.79
    $portada.Range('H9').Value2 = 911.58
    $excel.CalculateFullRebuild()
    Leer $libro 'DESPUES'
    $libro.Close($false)
}} finally {{
    if ($excel -ne $null) {{ $excel.Quit() }}
{liberar_com}
{restaurar_impresora}
}}
""".format(
        ruta=str(destino.resolve()).replace("'", "''"),
        suelos=mapa.nombre_hoja("suelos").replace("'", "''"),
        portada=mapa.nombre_hoja("portada").replace("'", "''"),
        arranque_impresora=GUION_ARRANQUE_IMPRESORA_SEGURA,
        restaurar_impresora=GUION_RESTAURAR_IMPRESORA,
        liberar_com=GUION_LIBERAR_COM_EXCEL,
        contar_errores=CONTAR_ERRORES_PS.format(),
    )
    try:
        proceso = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
            capture_output=True, text=True, timeout=170)
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        return None, "no se pudo recalcular con Excel real: {}".format(error)
    if proceso.returncode != 0:
        return None, "Excel no está instalado o falló: {}".format(
            (proceso.stderr or proceso.stdout)[:200])
    valores = {}
    for linea in proceso.stdout.splitlines():
        if "=" in linea:
            clave, _, valor = linea.partition("=")
            valores[clave.strip()] = valor.strip()
    return valores, None


def _numero(valores, clave):
    try:
        return float(valores.get(clave) or 0)
    except ValueError:
        return float("nan")


def probar_lo_que_falta_se_completa_en_el_excel():
    print("\n[9] Sin fecha, UF ni dólar, pasa lo que dicen los avisos -y "
          "escribirlos después en PORTADA lo arregla todo (Excel real)")
    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    datos = informe_de_prueba().model_dump(mode="json")
    datos["parametros"].update(fecha_retasacion=None, uf_base=None, valor_usd=None)
    retasacion, _ = leer_borrador(datos)
    destino = ajustes.ruta_salida_pruebas / "test_retasacion_incompleta.xlsx"
    comprobar(motor.generar(retasacion, destino).ok, "el archivo se generó")

    valores, motivo = _recalcular_antes_y_despues(destino, mapa)
    if valores is None:
        print("  ({})".format(motivo))
        return

    comprobar(_numero(valores, "SUELOS_G46_ANTES") == 0,
              "sin UF, el valor de los suelos sale en $0 -lo que dice el aviso",
              valores.get("SUELOS_G46_ANTES"))
    comprobar(_numero(valores, "PLANT_AA37_ANTES") == 0,
              "sin dólar, la plantación sin royalty sale valorada en $0",
              valores.get("PLANT_AA37_ANTES"))
    comprobar(_numero(valores, "PLANT_R37_ANTES") < 0,
              "sin fecha, la edad de la plantación sale mal (negativa)",
              valores.get("PLANT_R37_ANTES"))
    comprobar(_numero(valores, "ERRORES_ANTES") > 0,
              "y las columnas en UF y dólares salen en error",
              valores.get("ERRORES_ANTES"))

    comprobar(_numero(valores, "ERRORES_DESPUES") == 0,
              "escritos después en PORTADA!H7/H8/H9: cero celdas en error",
              valores.get("ERRORES_DESPUES"))
    comprobar(_numero(valores, "SUELOS_G46_DESPUES") > 0
              and _numero(valores, "PLANT_AA37_DESPUES") > 0,
              "y suelos y plantación recuperan su valor -las fórmulas siguen "
              "vivas, no hay que generar de nuevo",
              (valores.get("SUELOS_G46_DESPUES"), valores.get("PLANT_AA37_DESPUES")))
    comprobar(_numero(valores, "PLANT_R37_DESPUES") == 2026 - 2018,
              "con la fecha, la edad es la correcta (8 años)",
              valores.get("PLANT_R37_DESPUES"))


def _imagen_de_hueco(ruta, hoja, forma):
    """Los bytes de la imagen que muestra un hueco del dibujo."""
    from services import dibujos_xlsx as dx
    partes = dx._leer_paquete(ruta)
    parte_hoja = dx._hojas(partes)[hoja]
    dibujo = dx._dibujo_de(partes, parte_hoja)[1]
    raiz = dx._parsear(partes[dibujo])
    encontrado = dx._buscar_forma(raiz, forma)
    if not encontrado:
        return None
    blip = encontrado[2].find("xdr:blipFill/a:blip", dx.NS)
    rid = blip.get("{%s}embed" % dx.NS["r"])
    rel = next(r for r in dx._rels(partes, dibujo) if r["Id"] == rid)
    return partes[dx._resolver(dibujo, rel["Target"])]


def _fila_del_ancla(ruta, hoja, forma):
    """La fila (1-based) donde arranca una forma, o None si la hoja no trae
    dibujo o la forma no esta -que una prueba lo informe, no que se caiga-."""
    from services import dibujos_xlsx as dx
    partes = dx._leer_paquete(ruta)
    dibujo = dx._dibujo_de(partes, dx._hojas(partes)[hoja])
    encontrado = dx._buscar_forma(dx._parsear(partes[dibujo[1]]), forma) if dibujo else None
    if not encontrado:
        return None
    return int(encontrado[0].find("xdr:from/xdr:row", dx.NS).text) + 1


def _formas_en_excel(destino: Path):
    """Cuantas formas (sin contar comentarios) ve el Excel real en cada hoja:
    si tuvo que "reparar" el archivo al abrirlo, se le pierden."""
    guion = r"""
$ErrorActionPreference = 'Stop'
$excel = $null
$libro = $null
try {{
{arranque_impresora}
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $libro = $excel.Workbooks.Open('{ruta}', $false, $true)
    foreach ($hoja in $libro.Worksheets) {{
        $n = 0
        foreach ($s in $hoja.Shapes) {{ if ($s.Type -ne 4) {{ $n += 1 }} }}
        Write-Output ("FORMAS|{{0}}|{{1}}" -f $hoja.Name, $n)
    }}
    $libro.Close($false)
}} finally {{
    if ($excel -ne $null) {{ $excel.Quit() }}
{liberar_com}
{restaurar_impresora}
}}
""".format(ruta=str(destino.resolve()).replace("'", "''"),
           arranque_impresora=GUION_ARRANQUE_IMPRESORA_SEGURA,
           restaurar_impresora=GUION_RESTAURAR_IMPRESORA,
           liberar_com=GUION_LIBERAR_COM_EXCEL)
    try:
        proceso = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
            capture_output=True, text=True, timeout=170, encoding="utf-8", errors="replace")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return None
    if proceso.returncode != 0:
        return None
    return {linea.split("|")[1]: int(linea.split("|")[2])
            for linea in proceso.stdout.splitlines() if linea.startswith("FORMAS|")}


def probar_formato_identico():
    print("\n[10] El Excel sale con el formato del banco completo: cuadros de "
          "texto, imágenes de muestra, referencias de mercado (2026-10-03)")
    from openpyxl.cell.rich_text import CellRichText
    from PIL import Image, ImageStat

    from schemas.retasacion import Instalacion
    from services import dibujos_xlsx as dx

    mapa = cargar_mapa(MAPPING)
    motor = MotorRabobank(PLANTILLA, mapa)
    cuadros = {k: v for k, v in mapa.cuadros_de_texto.items() if isinstance(v, dict)}
    huecos = {k: v for k, v in mapa.huecos_de_imagen.items() if isinstance(v, dict)}
    hoja = mapa.nombre_hoja

    # -- la plantilla
    formas = dx.nombres_de_formas(PLANTILLA)
    faltan = [(s["hoja"], s["forma"]) for s in list(cuadros.values()) + list(huecos.values())
              if s["forma"] not in [n for _, n in formas.get(hoja(s["hoja"]), [])]]
    comprobar(not faltan, "la plantilla trae los {} cuadros de texto y los {} huecos de "
                          "imagen que declara el mapping".format(len(cuadros), len(huecos)), faltan)
    textos = dx.leer_textos(PLANTILLA, [(hoja(s["hoja"]), s["forma"]) for s in cuadros.values()])
    comprobar(not textos, "los cuadros de texto de la plantilla están vacíos -ningún "
                          "texto del cliente de origen-", list(textos))
    grises = []
    for clave, spec in huecos.items():
        datos = _imagen_de_hueco(PLANTILLA, hoja(spec["hoja"]), spec["forma"])
        media = ImageStat.Stat(Image.open(io.BytesIO(datos)).convert("L")).mean[0]
        if not 200 < media < 245:
            grises.append((clave, round(media)))
    comprobar(not grises, "cada hueco de imagen trae la muestra gris, no una imagen del "
                          "cliente de origen", grises)
    wb = load_workbook(PLANTILLA)
    rotos = [(ws.title, n) for ws in wb.worksheets for n, d in ws.defined_names.items()
             if n.upper() in ("UF", "USD") and "#REF!" in (d.attr_text or "")]
    comprobar(not rotos, "ningún nombre local UF/USD roto tapa al de la portada", rotos)
    ref = wb[hoja("ref_propiedades")]
    comprobar(all(ref.row_dimensions[f].hidden for f in range(6, 226)),
              "la base de referencias viene oculta: no se cuela la selección de otro informe")

    # -- un informe con todo
    r = informe_de_prueba()
    r.observaciones_portada = "El bien tasado corresponde a la Parcela de Prueba."
    r.legal_observaciones = "1) Nombre legal: Parcela de Prueba.\n2) Hipoteca vigente."
    r.suelos_descripcion.descripcion = "Predio plano, con acceso pavimentado."
    r.construcciones_observaciones = "Casa patronal en buen estado."
    r.plantaciones_observaciones = "Viñedo en producción."
    r.diferencias.plantaciones = "Se agregan 2 ha de vid."
    r.testigos_propiedades_analisis = "I. TAMAÑO DE LA MUESTRA: 4 referencias."
    r.predio.deslinde_norte = "con camino público"
    r.predio.deslinde_sur = "con Parcela N°3"
    r.predio.hipoteca_rabofinance = True
    r.aguas[0].inscrito_cpa = False
    r.instalaciones = [Instalacion(item="Riego {}".format(i), estado_conservacion="Operativo",
                                   ubicacion_rol="1001-1", cantidad=10 + i, unidad="ha",
                                   costo_reposicion_unitario_uf=60, vida_util_total_anios=20)
                       for i in range(3)]
    r.plantaciones_sin_royalty[0].especie = "Vid vinífera"
    r.plantaciones_sin_royalty[0].castigo_pct = 0.7
    r.plantaciones_sin_royalty[0].opex_anios = 0
    destino = ajustes.ruta_salida_pruebas / "test_retasacion_formato.xlsx"
    resultado = motor.generar(r, destino)
    comprobar(not resultado.avisos, "sin avisos del motor -'Vid vinífera' calza con "
                                    "'Vid Vinifera' de PLANT-", resultado.avisos)

    esperados = {
        (hoja("portada"), "CuadroTexto 2"): r.observaciones_portada,
        (hoja("legal"), "CuadroTexto 1"): r.legal_observaciones,
        (hoja("suelos"), "CuadroTexto 1"): r.suelos_descripcion.descripcion,
        (hoja("constr"), "Shape 4"): r.construcciones_observaciones,
        (hoja("plantacion"), "Shape 5"): r.plantaciones_observaciones,
        (hoja("plantacion"), "Shape 6"): r.diferencias.plantaciones,
        (hoja("ref_propiedades"), "CuadroTexto 1"): r.testigos_propiedades_analisis,
    }
    leidos = dx.leer_textos(destino, esperados)
    distintos = [k for k, v in esperados.items() if leidos.get(k) != v]
    comprobar(not distintos, "los siete cuadros de texto llevan su texto, con sus saltos "
                             "de línea", distintos)
    comprobar(_fila_del_ancla(destino, hoja("constr"), "Shape 4") == 29,
              "las dos instalaciones de más corren el cuadro de observaciones de CONSTR "
              "dos filas (27 -> 29), junto con la tabla")

    wb = load_workbook(destino, rich_text=True)
    deslindes = wb[hoja("legal")]["C35"].value
    comprobar(isinstance(deslindes, CellRichText) and str(deslindes).startswith("Norte: con camino")
              and deslindes[0].font.b,
              "los deslindes van en LEGAL!C35, un renglón por punto cardinal y su nombre "
              "en negrita, como en el formato", str(deslindes))
    comprobar((wb[hoja("legal")]["J6"].value, wb[hoja("agua")]["G7"].value,
               wb[hoja("agua")]["F7"].value, wb[hoja("constr")]["L6"].value)
              == ("Sí", "SÍ", "NO", "Sí"),
              "cada sí/no con la ortografía de su lista desplegable (AGUA 'SÍ/NO', el "
              "resto 'Sí/No')")
    plant = wb["PLANT"]
    comprobar((plant["C37"].value, plant["T37"].value, plant["X37"].value)
              == ("Vid Vinifera", 0, 0.7),
              "PLANT usa la clave de su tabla de costos y los ajustes fijados a mano",
              (plant["C37"].value, plant["T37"].value, plant["X37"].value))
    comprobar(wb[hoja("plantacion")]["A5"].value == "Vid vinífera",
              "y PLANTACION muestra la especie como se escribió")
    constr = wb[hoja("constr")]
    comprobar((constr["E25"].value, wb[hoja("portada")]["F30"].value)
              == ("=SUM(E22:E24)", "=CONSTR!S25"),
              "con tres instalaciones, el total cubre las tres y la portada lo sigue",
              (constr["E25"].value, wb[hoja("portada")]["F30"].value))

    ref = wb[hoja("ref_propiedades")]
    # La base (6-225 en la plantilla) crece con las 4 insertadas: 6-229.
    visibles = [f for f in range(6, 230) if not ref.row_dimensions[f].hidden and ref["A{}".format(f)].value]
    comprobar(visibles == [224, 225, 226, 227],
              "las 4 referencias del informe quedan visibles, dentro de la base, y el "
              "resto de la base oculto", visibles)
    comprobar(ref["B233"].value == "=SUBTOTAL(102,D6:D228)"
              and ref["B234"].value == '=IFERROR(SKEW(M224:M227),"")',
              "la estadística las cubre: el conteo crece con la base y la asimetría "
              "apunta a ellas", (ref["B233"].value, ref["B234"].value))
    comprobar(ref["H2"].value == "=PORTADA!$H$8",
              "la UF de la hoja de referencias es la de la portada", ref["H2"].value)

    # -- el criterio de valor
    r2 = informe_de_prueba()
    r2.valorar_plantaciones = False
    r2.valorar_instalaciones = False
    with tempfile.TemporaryDirectory() as tmp:
        sin = Path(tmp) / "sin_plantaciones.xlsx"
        motor.generar(r2, sin)
        w2 = load_workbook(sin)
        comprobar((w2[hoja("resumen")]["C13"].value, w2[hoja("portada")]["F28"].value,
                   w2[hoja("resumen")]["C16"].value, w2[hoja("constr")]["S23"].value)
                  == (None, 0, None, None),
                  "sin considerar plantaciones ni instalaciones: el resumen en blanco y la "
                  "portada en 0, como el formato real de ese caso")
        leidos, _ = leer_retasacion_excel(sin, mapa)
        comprobar(leidos.get("valorar_plantaciones") is False
                  and leidos.get("valorar_instalaciones") is False,
                  "y el Paso 1 lo vuelve a leer así", leidos.get("valorar_plantaciones"))

    # -- la lectura de vuelta
    leidos, avisos = leer_retasacion_excel(destino, mapa)
    comprobar(leidos.get("observaciones_portada") == r.observaciones_portada
              and leidos.get("suelos_descripcion", {}).get("descripcion")
              == r.suelos_descripcion.descripcion
              and leidos.get("diferencias", {}).get("plantaciones") == r.diferencias.plantaciones,
              "el Paso 1 lee los cuadros de texto de un Excel anterior")
    comprobar(leidos.get("predio", {}).get("deslinde_sur") == "con Parcela N°3",
              "y los deslindes, separados por punto cardinal", leidos.get("predio"))
    plantas = leidos.get("plantaciones_sin_royalty", [{}])
    comprobar(plantas[0].get("castigo_pct") == 0.7 and plantas[0].get("opex_anios") == 0,
              "y los ajustes de PLANT fijados a mano (no las fórmulas)", plantas[0])
    comprobar(len(leidos.get("testigos_propiedades", [])) == 4,
              "y solo las referencias visibles -las del informe-, no la base entera",
              len(leidos.get("testigos_propiedades", [])))
    # Medidos en el navegador (2026-10-03): leia 1 de 3 instalaciones -solo la
    # capacidad de la plantilla- y desmarcaba "considerar plantaciones" en un
    # archivo recien generado, porque una formula que Excel no calculo todavia
    # vale None en la lectura calculada.
    comprobar(len(leidos.get("instalaciones", [])) == 3,
              "y las tres instalaciones, aunque la plantilla traiga una fila: lee hasta "
              "el rótulo de totales", len(leidos.get("instalaciones", [])))
    comprobar(leidos.get("valorar_plantaciones") is True
              and leidos.get("valorar_instalaciones") is True,
              "y el criterio de valor de un archivo que Excel todavía no abrió (fórmulas "
              "sin calcular)", (leidos.get("valorar_plantaciones"),
                                leidos.get("valorar_instalaciones")))

    # -- el Excel real
    excel = _formas_en_excel(destino)
    if excel is None:
        print("  (sin Excel: se omite la apertura real)")
        return
    # Contra la PLANTILLA, no contra el propio archivo: lo que se quiere saber
    # es que ninguna forma del formato se perdio en el camino (sin el
    # injerto, openpyxl se queda solo con las imagenes).
    esperado = {nombre: len(dx._parsear(partes).findall("*"))
                for nombre, partes in _dibujos_por_hoja(PLANTILLA).items()}
    comprobar(all(excel.get(n) == c for n, c in esperado.items()),
              "Excel abre el archivo sin repararlo y ve en cada hoja todas las formas "
              "de la plantilla", (excel, esperado))
    valores, motivo = _recalcular_con_excel(destino)
    if valores is not None:
        comprobar(int(valores.get("ERRORES", "-1")) == 0,
                  "recalculado con Excel: cero celdas en error, contadas celda por celda",
                  valores.get("ERRORES"))


def _dibujos_por_hoja(ruta):
    from services import dibujos_xlsx as dx
    partes = dx._leer_paquete(ruta)
    salida = {}
    for nombre, parte in dx._hojas(partes).items():
        dibujo = dx._dibujo_de(partes, parte)
        if dibujo:
            salida[nombre] = partes[dibujo[1]]
    return salida


def main():
    print("=" * 72)
    print("PRUEBA DEL MOTOR DE RETASACIONES RABOBANK")
    print("=" * 72)

    probar_plantilla_sin_rastro_de_cliente_real()
    probar_motor_genera_sin_errores()
    probar_especie_invalida_avisa_y_no_inventa_valor()
    probar_recalculo_con_excel_real()
    probar_lectura_recupera_lo_que_escribio_el_motor()
    probar_lectura_de_archivo_ajeno_no_inventa_nada()
    probar_lectura_de_pdf_sintetico()
    probar_incompleta_se_genera_y_avisa()
    probar_lo_que_falta_se_completa_en_el_excel()
    probar_formato_identico()

    print("\n" + "=" * 72)
    if fallos:
        print("{}/{} comprobaciones FALLARON:".format(len(fallos), comprobaciones))
        for f in fallos:
            print("  - {}".format(f))
        return 1
    print("{}/{} comprobaciones OK".format(comprobaciones, comprobaciones))
    return 0


if __name__ == "__main__":
    sys.exit(main())
